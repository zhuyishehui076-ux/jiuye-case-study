"""T7 导入中心闭环验收（真实 PostgreSQL）。

覆盖 SPEC §6.8 / §7 T7：模板、六类行级错误、日期错误、xlsx 回执、
确认入库、整批回滚零业务残留，以及写入口 RBAC。
"""

from __future__ import annotations

from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.models import AuditLog, College, ImportBatch, Major, Student
from app.models.enums import DegreeLevel, EmploymentStatus, Gender
from app.services.imports import (
    FIELD_LABELS,
    TEMPLATE_HEADERS,
    build_receipt_xlsx,
    build_template_xlsx,
    confirm_import_batch,
    list_templates,
    rollback_import_batch,
    upload_import_batch,
)
from tests.conftest import login


def _test_password(username: str) -> str:
    return f"{username}{username}!"


def _workbook(kind: str, rows: list[dict[str, object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([FIELD_LABELS[field] for field in TEMPLATE_HEADERS[kind]])
    for row in rows:
        sheet.append([row.get(field) for field in TEMPLATE_HEADERS[kind]])
    payload = BytesIO()
    workbook.save(payload)
    return payload.getvalue()


@pytest.fixture
def import_client(db):
    """给临时 PostgreSQL 测试库建立显式测试账号，不依赖任何运行态凭据。"""
    from fastapi.testclient import TestClient

    from app.core.config import settings
    from app.db.session import get_db
    from app.main import app
    from app.seed.users import seed_demo_users
    from app.services.installation import ensure_installation_profile

    ensure_installation_profile(
        db,
        expected_profile=settings.profile,
        expected_installation_id=settings.installation_id,
    )
    college = College(
        code="T7RBAC",
        name="T7 权限测试学院",
        created_by="t7-test",
    )
    db.add(college)
    db.flush()
    seed_demo_users(
        db,
            commit=False,
            passwords={
                username: _test_password(username)
                for username in ("admin", "office", "counselor", "leader")
            },
        counselor_college_id=college.id,
    )

    def override_db():
        yield db

    app.dependency_overrides[get_db] = override_db
    test_client = TestClient(app)
    try:
        yield test_client
    finally:
        test_client.close()
        app.dependency_overrides.pop(get_db, None)


def _catalog(db):
    suffix = uuid4().hex[:8]
    college = College(code=f"C{suffix}", name=f"T7 测试学院 {suffix}", created_by="t7-test")
    db.add(college)
    db.flush()
    major = Major(
        code=f"M{suffix}",
        name=f"T7 测试专业 {suffix}",
        level=DegreeLevel.BACHELOR.value,
        college_id=college.id,
        created_by="t7-test",
    )
    db.add(major)
    db.flush()
    return suffix, college, major


def _student_row(
    *,
    suffix: str,
    college_code: str,
    major_code: str,
    student_no: str | None = None,
    **overrides: object,
) -> dict[str, object]:
    row: dict[str, object] = {
        "student_no": student_no if student_no is not None else f"T7-{suffix}",
        "name": "T7 测试学生",
        "gender": Gender.MALE.value,
        "college_code": college_code,
        "major_code": major_code,
        "degree_level": DegreeLevel.BACHELOR.value,
        "cohort_year": 2026,
        "status": EmploymentStatus.UNSETTLED.value,
        "origin_place": "测试地区",
        "is_key_group": "",
    }
    row.update(overrides)
    return row


def test_all_five_templates_are_valid_xlsx() -> None:
    kinds = [row["kind"] for row in list_templates()]
    assert kinds == ["students", "assessments", "activity_attendance", "jobs", "outcomes"]

    for kind in kinds:
        workbook = load_workbook(
            BytesIO(build_template_xlsx(kind)),
            read_only=True,
            data_only=True,
        )
        sheet = workbook.active
        assert sheet.max_row >= 2
        assert sheet.max_column == len(TEMPLATE_HEADERS[kind])
        assert [cell.value for cell in sheet[1]] == [
            FIELD_LABELS[field] for field in TEMPLATE_HEADERS[kind]
        ]


def test_six_dirty_student_rows_are_all_reported_with_exact_row_numbers(db) -> None:
    suffix, college, major = _catalog(db)
    existing_no = f"T7-DUP-{suffix}"
    db.add(
        Student(
            student_no=existing_no,
            name="已有学生",
            gender=Gender.MALE.value,
            college_id=college.id,
            major_id=major.id,
            degree_level=DegreeLevel.BACHELOR.value,
            cohort_year=2026,
            status=EmploymentStatus.UNSETTLED.value,
            created_by="t7-test",
        )
    )
    db.flush()

    valid = {
        "suffix": suffix,
        "college_code": college.code,
        "major_code": major.code,
    }
    rows = [
        _student_row(**valid, student_no="", name="缺少学号"),
        _student_row(**valid, student_no=existing_no, name="重复学号"),
        _student_row(
            suffix=suffix,
            college_code="UNKNOWN",
            major_code=major.code,
            student_no=f"T7-C-{suffix}",
        ),
        _student_row(
            suffix=suffix,
            college_code=college.code,
            major_code="UNKNOWN",
            student_no=f"T7-M-{suffix}",
        ),
        _student_row(**valid, student_no=f"T7-E-{suffix}", gender="未知枚举"),
        _student_row(**valid, student_no=f"T7-Y-{suffix}", cohort_year="二零二六"),
    ]

    batch = upload_import_batch(
        db,
        kind="students",
        data=_workbook("students", rows),
        filename="six-dirty-students.xlsx",
        actor="t7-test",
    )

    assert batch.status == "failed"
    assert batch.row_count == 6
    assert batch.error_count == 6
    assert [item["row_number"] for item in batch.errors] == [2, 3, 4, 5, 6, 7]
    assert {item["code"] for item in batch.errors} == {
        "missing_required",
        "student_no_duplicate",
        "college_code_unknown",
        "major_code_unknown",
        "enum_out_of_range",
        "cohort_year_format",
    }
    assert all(item["field"] for item in batch.errors)

    receipt = load_workbook(
        BytesIO(build_receipt_xlsx(batch)),
        read_only=True,
        data_only=True,
    )["校验回执"]
    assert receipt.max_row == 7
    assert [cell.value for cell in receipt[1]] == [
        "行号",
        "问题类型",
        "修改说明",
        "需要修改的列",
        "当前内容",
    ]
    assert [receipt.cell(row=row, column=1).value for row in range(2, 8)] == [2, 3, 4, 5, 6, 7]


def test_invalid_date_is_a_row_level_error(db) -> None:
    batch = upload_import_batch(
        db,
        kind="jobs",
        data=_workbook(
            "jobs",
            [
                {
                    "company": "T7 测试企业",
                    "title": "测试岗位",
                    "raw_text": "只用于 T7 行级日期校验。",
                    "received_at": "2026-99-99",
                }
            ],
        ),
        filename="invalid-date.xlsx",
        actor="t7-test",
    )

    assert batch.status == "failed"
    assert batch.error_count == 1
    assert batch.errors[0] == {
        "row_number": 2,
        "code": "date_format",
        "reason": "日期请按“2026-03-01”的格式填写",
        "field": "接收日期",
        "value": None,
    }


def test_confirm_then_rollback_leaves_zero_student_residue_and_keeps_audit(db) -> None:
    suffix, college, major = _catalog(db)
    student_numbers = [f"T7-OK1-{suffix}", f"T7-OK2-{suffix}"]
    rows = [
        _student_row(
            suffix=suffix,
            college_code=college.code,
            major_code=major.code,
            student_no=student_no,
            name=f"测试学生 {index}",
        )
        for index, student_no in enumerate(student_numbers, start=1)
    ]

    batch = upload_import_batch(
        db,
        kind="students",
        data=_workbook("students", rows),
        filename="valid-students.xlsx",
        actor="t7-test",
    )
    batch_id = batch.id
    assert batch.status == "validated"
    assert batch.error_count == 0

    confirmed, inserted_count = confirm_import_batch(db, batch_id, actor="t7-test")
    assert confirmed.status == "imported"
    assert inserted_count == 2
    assert len(confirmed.inserted_refs["students"]) == 2
    assert db.scalar(
        select(func.count())
        .select_from(Student)
        .where(Student.student_no.in_(student_numbers))
    ) == 2

    rolled_back, rolled_back_count = rollback_import_batch(db, batch_id, actor="t7-test")
    assert rolled_back.status == "rolled_back"
    assert rolled_back_count == 2
    assert db.scalar(
        select(func.count())
        .select_from(Student)
        .where(Student.student_no.in_(student_numbers))
    ) == 0

    persisted_batch = db.get(ImportBatch, batch_id)
    assert persisted_batch is not None
    assert persisted_batch.status == "rolled_back"
    assert len(persisted_batch.inserted_refs["students"]) == 2
    actions = set(
        db.execute(
            select(AuditLog.action).where(
                AuditLog.object_type == "import_batch",
                AuditLog.object_id == str(batch_id),
            )
        ).scalars()
    )
    assert {"import.upload", "import.confirm", "import.rollback"} <= actions


@pytest.mark.parametrize(
    ("username", "password", "expected"),
    [
        ("admin", _test_password("admin"), 200),
        ("office", _test_password("office"), 200),
        ("counselor", _test_password("counselor"), 403),
        ("leader", _test_password("leader"), 403),
    ],
)
def test_import_template_access_respects_write_roles(
    import_client, username, password, expected
) -> None:
    token = login(import_client, username, password)
    response = import_client.get(
        "/api/v1/imports/templates",
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == expected, response.text
