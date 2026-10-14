"""T6 M2 一生一策帮扶台账测试（SPEC §6.2）。"""

from __future__ import annotations

from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook
from sqlalchemy import func, select

from tests.conftest import login


def test_red_students_auto_create_assist_cases(pg_engine):
    from app.db.session import SessionLocal
    from app.models import AssistCase, Student
    from tests.conftest import reset_case_demo_seed

    with SessionLocal() as s:
        counts = reset_case_demo_seed(s)
        risk_red = s.execute(
            select(func.count()).select_from(Student).where(Student.risk_level == "red")
        ).scalar_one()
        case_count = s.execute(select(func.count()).select_from(AssistCase)).scalar_one()
        assigned_count = s.execute(
            select(func.count()).select_from(AssistCase).where(AssistCase.assigned_to.is_not(None))
        ).scalar_one()

    assert counts["risk_red"] == risk_red
    assert case_count == risk_red
    assert assigned_count == risk_red


def test_assist_record_append_only_and_no_delete_route(client, pg_engine, demo_passwords):
    from app.db.session import SessionLocal
    from app.models import AssistCase, AssistRecord, AuditLog
    from tests.conftest import reset_case_demo_seed

    with SessionLocal() as s:
        reset_case_demo_seed(s)
        # 使用尚无记录的模拟档案，避免把故事种子的既有记录计入两次追加。
        case_id = s.execute(
            select(AssistCase.id)
            .where(~select(AssistRecord.id).where(AssistRecord.case_id == AssistCase.id).exists())
            .order_by(AssistCase.id)
            .limit(1)
        ).scalar_one()

    token = login(client, "office", demo_passwords['office'])
    headers = {"Authorization": f"Bearer {token}"}
    missing_key = client.post(
        f"/api/v1/assist/cases/{case_id}/records",
        headers=headers,
        json={"action_type": "约谈", "note": "缺少幂等键应被拒绝"},
    )
    assert missing_key.status_code == 422, missing_key.text
    first_headers = {**headers, "Idempotency-Key": str(uuid4())}
    first = client.post(
        f"/api/v1/assist/cases/{case_id}/records",
        headers=first_headers,
        json={"action_type": "约谈", "note": "<p>首次约谈，确认求职目标。</p>"},
    )
    assert first.status_code == 200, first.text
    replayed = client.post(
        f"/api/v1/assist/cases/{case_id}/records",
        headers=first_headers,
        json={"action_type": "约谈", "note": "<p>首次约谈，确认求职目标。</p>"},
    )
    assert replayed.status_code == 200, replayed.text
    assert replayed.json()["id"] == first.json()["id"]
    assert replayed.json()["replayed"] is True
    conflict = client.post(
        f"/api/v1/assist/cases/{case_id}/records",
        headers=first_headers,
        json={"action_type": "其他", "note": "同一幂等键不能改写记录"},
    )
    assert conflict.status_code == 409, conflict.text
    correction = client.post(
        f"/api/v1/assist/cases/{case_id}/records",
        headers={**headers, "Idempotency-Key": str(uuid4())},
        json={"action_type": "其他", "note": "<p>更正：目标企业补充为本地软件企业。</p>"},
    )
    assert correction.status_code == 200, correction.text

    with SessionLocal() as s:
        rows = s.execute(
            select(AssistRecord).where(AssistRecord.case_id == case_id).order_by(AssistRecord.id)
        ).scalars().all()
        audit_count = s.scalar(
            select(func.count()).select_from(AuditLog).where(
                AuditLog.action == "assist_record.append",
                AuditLog.object_id.in_([str(row.id) for row in rows]),
            )
        )

    assert len(rows) == 2
    assert audit_count == 2
    assert rows[0].note == "<p>首次约谈，确认求职目标。</p>"
    assert rows[1].note.startswith("<p>更正")
    from app.main import app

    assist_delete_routes = [
        route.path
        for route in app.routes
        if "/assist" in getattr(route, "path", "")
        and "DELETE" in (getattr(route, "methods", set()) or set())
    ]
    assert assist_delete_routes == []


def test_assist_exports_are_readable(client, pg_engine, demo_passwords):
    from docx import Document

    from app.db.session import SessionLocal
    from app.models import AssistCase
    from tests.conftest import reset_case_demo_seed

    with SessionLocal() as s:
        reset_case_demo_seed(s)
        case_id = s.execute(select(AssistCase.id).order_by(AssistCase.id).limit(1)).scalar_one()

    token = login(client, "office", demo_passwords['office'])
    headers = {"Authorization": f"Bearer {token}"}
    appended = client.post(
        f"/api/v1/assist/cases/{case_id}/records",
        headers={**headers, "Idempotency-Key": str(uuid4())},
        json={"action_type": "培训", "note": "参加简历训练营"},
    )
    assert appended.status_code == 200, appended.text

    xlsx = client.get("/api/v1/assist/export.xlsx", headers=headers)
    assert xlsx.status_code == 200, xlsx.text
    wb = load_workbook(BytesIO(xlsx.content))
    ws = wb.active
    headers_row = [cell.value for cell in ws[1]]
    for col in ("学号", "姓名", "学院", "预警原因", "帮扶状态", "动作汇总"):
        assert col in headers_row
    assert ws.max_row > 1

    docx = client.get("/api/v1/assist/export.docx", headers=headers)
    assert docx.status_code == 200, docx.text
    doc = Document(BytesIO(docx.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    for fragment in ("一生一策帮扶台账", "学生", "帮扶记录"):
        assert fragment in text


def test_counselor_assist_cases_limited_to_own_college(client, pg_engine, demo_passwords):
    from app.db.session import SessionLocal
    from tests.conftest import reset_case_demo_seed

    with SessionLocal() as s:
        reset_case_demo_seed(s)

    counselor_token = login(client, "counselor", demo_passwords['counselor'])
    office_token = login(client, "office", demo_passwords['office'])
    counselor_rows = client.get(
        "/api/v1/assist/cases",
        headers={"Authorization": f"Bearer {counselor_token}"},
    ).json()
    office_rows = client.get(
        "/api/v1/assist/cases",
        headers={"Authorization": f"Bearer {office_token}"},
    ).json()

    assert counselor_rows
    assert all(row["college_id"] == 1 for row in counselor_rows)
    assert len(office_rows) >= len(counselor_rows)
    assert any(row["college_id"] != 1 for row in office_rows)

    outside = next(row for row in office_rows if row["college_id"] != 1)
    denied = client.post(
        f"/api/v1/assist/cases/{outside['id']}/records",
        headers={
            "Authorization": f"Bearer {counselor_token}",
            "Idempotency-Key": str(uuid4()),
        },
        json={"action_type": "约谈", "note": "越权写入应被拒绝"},
    )
    assert denied.status_code == 403
