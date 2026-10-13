"""T7 X2 import center service."""

from __future__ import annotations

from datetime import UTC, date, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.core.audit import record_db
from app.core.config import Profile, settings
from app.models import (
    Activity,
    ActivityAttendance,
    Assessment,
    College,
    ImportBatch,
    Job,
    Major,
    Student,
)
from app.models.enums import (
    AssessmentType,
    AttendanceSource,
    DegreeLevel,
    EmploymentStatus,
    Gender,
    JobSource,
)

LEGACY_PRODUCTION_CORE_KINDS = frozenset({"students", "outcomes"})
LEGACY_CORE_PRODUCTION_DETAIL = (
    "正式档位的学生与去向只能通过“五类资料启用”流程处理，请新建数据启用记录。"
)
LEGACY_CORE_PRODUCTION_ROLLBACK_DETAIL = (
    "正式档位无法安全自动撤销升级前的学生或去向批次；请保留现状并通过备份与受控数据恢复流程处理。"
)
RUN_BOUND_BATCH_DETAIL = (
    "该批次属于五类资料启用记录，禁止单批次确认或撤销，请回到对应启用记录操作。"
)


class LegacyImportConflictError(ValueError):
    """A legacy mutation conflicts with the T16 atomic-import boundary."""


def ensure_legacy_upload_allowed(kind: str) -> None:
    """Prevent the T7 endpoint from creating new core-data batches in production."""
    if settings.profile is Profile.PRODUCTION and kind in LEGACY_PRODUCTION_CORE_KINDS:
        raise LegacyImportConflictError(LEGACY_CORE_PRODUCTION_DETAIL)


def _ensure_legacy_batch_mutation_allowed(
    batch: ImportBatch,
    *,
    operation: str,
) -> None:
    """Keep T16 children atomic and fail closed on legacy production core data."""
    if batch.run_id is not None:
        raise LegacyImportConflictError(RUN_BOUND_BATCH_DETAIL)
    if (
        settings.profile is Profile.PRODUCTION
        and batch.target_kind in LEGACY_PRODUCTION_CORE_KINDS
    ):
        if operation == "rollback":
            raise LegacyImportConflictError(LEGACY_CORE_PRODUCTION_ROLLBACK_DETAIL)
        raise LegacyImportConflictError(LEGACY_CORE_PRODUCTION_DETAIL)


TEMPLATE_HEADERS: dict[str, list[str]] = {
    "students": [
        "student_no",
        "name",
        "gender",
        "college_code",
        "major_code",
        "degree_level",
        "cohort_year",
        "status",
        "origin_place",
        "is_key_group",
    ],
    "assessments": [
        "student_no",
        "type",
        "term",
        "academic_basics",
        "learning_ability",
        "professional_confidence",
        "communication",
        "expression",
        "research_interest",
        "innovation_practice",
        "summary",
    ],
    "activity_attendance": ["student_no", "activity_name", "checked_at", "source"],
    "jobs": ["company", "title", "raw_text", "received_at"],
    "outcomes": ["student_no", "status", "occurred_at"],
}

FIELD_LABELS: dict[str, str] = {
    "student_no": "学号",
    "name": "姓名",
    "gender": "性别",
    "college_code": "学院代码",
    "major_code": "专业代码",
    "degree_level": "学历层次",
    "cohort_year": "毕业届别",
    "status": "就业去向",
    "origin_place": "生源地",
    "is_key_group": "重点群体",
    "type": "测评类型",
    "term": "学期",
    "academic_basics": "专业基础（1-5）",
    "learning_ability": "学习能力（1-5）",
    "professional_confidence": "专业信心（1-5）",
    "communication": "沟通能力（1-5）",
    "expression": "表达能力（1-5）",
    "research_interest": "科研兴趣（1-5）",
    "innovation_practice": "创新实践（1-5）",
    "summary": "结果摘要",
    "activity_name": "活动名称",
    "checked_at": "签到日期",
    "source": "记录来源",
    "company": "企业名称",
    "title": "岗位名称",
    "raw_text": "职位说明",
    "received_at": "接收日期",
    "occurred_at": "登记日期",
}

TEMPLATE_FILENAMES: dict[str, str] = {
    "students": "学生基础信息导入模板.xlsx",
    "assessments": "学生测评导入模板.xlsx",
    "activity_attendance": "活动签到导入模板.xlsx",
    "jobs": "岗位信息导入模板.xlsx",
    "outcomes": "就业去向导入模板.xlsx",
}

IMPORT_VALUE_ALIASES: dict[str, dict[str, str]] = {
    "type": {"性格类型测评": "MBTI"},
}

EXPORT_VALUE_LABELS: dict[str, dict[str, str]] = {
    "type": {"MBTI": "性格类型测评"},
}

ERROR_CODE_LABELS: dict[str, str] = {
    "missing_header": "缺少必填列",
    "missing_required": "必填内容为空",
    "student_no_duplicate": "学号重复",
    "student_no_unknown": "学号不存在",
    "college_code_unknown": "学院代码不存在",
    "major_code_unknown": "专业代码不存在",
    "enum_out_of_range": "内容不在可选范围内",
    "cohort_year_format": "毕业届别格式不正确",
    "date_format": "日期格式不正确",
    "score_format": "测评评分格式不正确",
}

REQUIRED_HEADERS: dict[str, list[str]] = {
    "students": [
        "student_no",
        "name",
        "gender",
        "college_code",
        "major_code",
        "degree_level",
        "cohort_year",
        "status",
    ],
    "assessments": TEMPLATE_HEADERS["assessments"],
    "activity_attendance": TEMPLATE_HEADERS["activity_attendance"],
    "jobs": ["company", "title", "raw_text"],
    "outcomes": ["student_no", "status", "occurred_at"],
}

TEMPLATE_TITLES: dict[str, str] = {
    "students": "生源表",
    "assessments": "测评表",
    "activity_attendance": "活动签到表",
    "jobs": "岗位表",
    "outcomes": "去向表",
}


def list_templates() -> list[dict[str, str]]:
    return [
        {"kind": kind, "title": TEMPLATE_TITLES[kind], "filename": TEMPLATE_FILENAMES[kind]}
        for kind in TEMPLATE_HEADERS
    ]


def build_template_xlsx(kind: str) -> bytes:
    _ensure_kind(kind)
    wb = Workbook()
    ws = wb.active
    ws.title = TEMPLATE_TITLES[kind]
    ws.append([FIELD_LABELS[header] for header in TEMPLATE_HEADERS[kind]])
    examples = {
        "students": [
            "20260001",
            "张三",
            Gender.MALE.value,
            "C01",
            "080901",
            DegreeLevel.BACHELOR.value,
            2026,
            EmploymentStatus.UNSETTLED.value,
            "南京",
            "",
        ],
        "assessments": [
            "20260001",
            AssessmentType.SELF_REPORT.value,
            "2025-2026-1",
            4,
            4,
            3,
            4,
            4,
            3,
            4,
            "方向基本清晰，建议增加目标岗位投递。",
        ],
        "activity_attendance": [
            "20260001",
            "春季双选会",
            "2026-03-01",
            AttendanceSource.SCAN.value,
        ],
        "jobs": ["启航科技", "算法工程师", "本科，计算机相关专业，南京", "2026-03-01"],
        "outcomes": ["20260001", EmploymentStatus.SIGNED.value, "2026-04-01"],
    }
    ws.append(examples[kind])
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="335CFF")
        cell.alignment = Alignment(horizontal="center")
    for column in ws.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 4, 36)
        ws.column_dimensions[column[0].column_letter].width = width
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _demo_workbook(
    kind: str,
    rows: list[dict[str, Any]],
    *,
    headers: list[str] | None = None,
) -> bytes:
    """构造仓库内置质量样本；仍走与正式导入相同的 Excel 读取器。"""
    selected_headers = headers or TEMPLATE_HEADERS[kind]
    wb = Workbook()
    ws = wb.active
    ws.title = TEMPLATE_TITLES[kind]
    ws.append([FIELD_LABELS[field] for field in selected_headers])
    for row in rows:
        ws.append([row.get(field) for field in selected_headers])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _demo_preflight_sources(corrected: bool) -> list[tuple[str, str, bytes, int]]:
    student_base = {
        "student_no": "PREF001",
        "name": "模拟学生",
        "gender": Gender.MALE.value,
        "college_code": "C01",
        "major_code": "080901",
        "degree_level": DegreeLevel.BACHELOR.value,
        "cohort_year": 2026,
        "status": EmploymentStatus.UNSETTLED.value,
        "origin_place": "模拟地区",
        "is_key_group": "",
    }
    assessment_base = {
        "student_no": "MOCKREF001",
        "type": AssessmentType.SELF_REPORT.value,
        "term": "2025-2026-1",
        "academic_basics": 4,
        "learning_ability": 4,
        "professional_confidence": 4,
        "communication": 4,
        "expression": 4,
        "research_interest": 4,
        "innovation_practice": 4,
        "summary": "内置模拟质量样本",
    }
    job_base = {
        "company": "模拟企业",
        "title": "模拟岗位",
        "raw_text": "仅用于内置质量检查的虚构岗位说明。",
        "received_at": "2026-03-01",
    }

    if corrected:
        student_rows = [
            {**student_base, "student_no": f"PREF00{index}", "name": f"模拟学生{index}"}
            for index in range(1, 7)
        ]
        assessment_rows = [{**assessment_base, "term": f"2025-2026-{index}"} for index in (1, 2)]
        job_rows = [job_base]
        header_rows = [{**job_base, "title": "模拟岗位二"}]
        header_fields = TEMPLATE_HEADERS["jobs"]
    else:
        student_rows = [
            {**student_base, "student_no": ""},
            {**student_base, "student_no": "MOCKREF001"},
            {**student_base, "student_no": "PREF003", "college_code": "C99"},
            {**student_base, "student_no": "PREF004", "major_code": "999999"},
            {**student_base, "student_no": "PREF005", "gender": "未知"},
            {**student_base, "student_no": "PREF006", "cohort_year": "二零二六"},
        ]
        assessment_rows = [
            {**assessment_base, "academic_basics": 9},
            {**assessment_base, "student_no": "UNKNOWN001"},
        ]
        job_rows = [{**job_base, "received_at": "2026-99-99"}]
        header_rows = [job_base]
        header_fields = ["title", "raw_text", "received_at"]

    return [
        ("students", "学生基础信息", _demo_workbook("students", student_rows), len(student_rows)),
        (
            "assessments",
            "学生测评",
            _demo_workbook("assessments", assessment_rows),
            len(assessment_rows),
        ),
        ("jobs", "岗位信息", _demo_workbook("jobs", job_rows), len(job_rows)),
        (
            "jobs",
            "岗位表头",
            _demo_workbook("jobs", header_rows, headers=header_fields),
            len(header_rows),
        ),
    ]


def validate_demo_preflight(db: Session, *, corrected: bool) -> dict[str, Any]:
    """用真实导入读取/校验管线复跑 10 行内置样本，不写入业务数据。"""
    nested = db.begin_nested()
    try:
        college = College(name="内置质量检查学院", code="C01", created_by="demo-preflight")
        db.add(college)
        db.flush()
        major = Major(
            code="080901",
            name="内置质量检查专业",
            level=DegreeLevel.BACHELOR.value,
            college_id=college.id,
            created_by="demo-preflight",
        )
        db.add(major)
        db.flush()
        db.add(
            Student(
                student_no="MOCKREF001",
                name="模拟参照学生",
                gender=Gender.MALE.value,
                college_id=college.id,
                major_id=major.id,
                degree_level=DegreeLevel.BACHELOR.value,
                cohort_year=2026,
                status=EmploymentStatus.UNSETTLED.value,
                created_by="demo-preflight",
            )
        )
        db.flush()

        source_rows = 0
        inserted_count = 0
        errors: list[dict[str, Any]] = []
        for kind, label, data, physical_rows in _demo_preflight_sources(corrected):
            source_rows += physical_rows
            rows, header_errors = _read_rows(kind, data)
            staged: list[dict[str, Any]] = []
            validation_errors = header_errors
            if not validation_errors:
                staged, validation_errors = _validate_rows(db, kind, rows)
            inserted_count += len(staged)
            for error in validation_errors:
                errors.append(
                    {
                        **error,
                        "field": f"{label} / {_field_label(str(error.get('field') or '内容'))}",
                    }
                )
        for index, error in enumerate(errors, start=2):
            error["row_number"] = index
        return {
            "source_rows": source_rows,
            "inserted_count": inserted_count,
            "updated_count": 0,
            "skipped_count": 0,
            "error_count": len(errors),
            "errors": errors,
        }
    finally:
        nested.rollback()
        db.expire_all()


def upload_import_batch(
    db: Session,
    *,
    kind: str,
    data: bytes,
    filename: str | None,
    actor: str,
) -> ImportBatch:
    _ensure_kind(kind)
    ensure_legacy_upload_allowed(kind)
    try:
        rows, header_errors = _read_rows(kind, data)
    except Exception as exc:
        raise ValueError("工作簿内容损坏或不是有效的 Excel 文件") from exc
    errors = header_errors
    staged_rows: list[dict[str, Any]] = []
    if not errors:
        staged_rows, errors = _validate_rows(db, kind, rows)
    batch = ImportBatch(
        target_kind=kind,
        original_filename=filename,
        status="validated" if not errors else "failed",
        row_count=len(rows),
        error_count=len(errors),
        errors=errors,
        staged_rows=staged_rows,
        inserted_refs={},
        created_by=actor,
    )
    db.add(batch)
    db.flush()
    record_db(
        db,
        actor=actor,
        action="import.upload",
        object_type="import_batch",
        object_id=batch.id,
        change_summary={
            "target_kind": kind,
            "row_count": batch.row_count,
            "error_count": batch.error_count,
        },
    )
    db.commit()
    db.refresh(batch)
    return batch


def build_receipt_xlsx(batch: ImportBatch) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "校验回执"
    ws.append(["行号", "问题类型", "修改说明", "需要修改的列", "当前内容"])
    if batch.errors:
        for err in batch.errors:
            ws.append(
                [
                    err.get("row_number"),
                    ERROR_CODE_LABELS.get(str(err.get("code")), "内容需要修改"),
                    err.get("reason"),
                    err.get("field"),
                    err.get("value"),
                ]
            )
    else:
        ws.append(["—", "检查通过", "文件内容可以导入", "—", "—"])
    ws.freeze_panes = "A2"
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                cell.value = f"'{cell.value}"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="335CFF")
    for key, width in {"A": 10, "B": 22, "C": 42, "D": 24, "E": 28}.items():
        ws.column_dimensions[key].width = width
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def confirm_import_batch(db: Session, batch_id: int, *, actor: str) -> tuple[ImportBatch, int]:
    batch = _get_batch(db, batch_id, lock=True)
    _ensure_legacy_batch_mutation_allowed(batch, operation="confirm")
    if batch.error_count:
        raise ValueError("batch has validation errors")
    if batch.status != "validated":
        raise ValueError(f"batch status {batch.status} cannot be confirmed")
    inserted_refs: dict[str, Any]
    if batch.target_kind == "students":
        inserted_refs = _confirm_students(db, batch.staged_rows, actor)
    elif batch.target_kind == "jobs":
        inserted_refs = _confirm_jobs(db, batch.staged_rows, actor)
    elif batch.target_kind == "assessments":
        inserted_refs = _confirm_assessments(db, batch.staged_rows, actor)
    elif batch.target_kind == "activity_attendance":
        inserted_refs = _confirm_activity_attendance(db, batch.staged_rows, actor)
    elif batch.target_kind == "outcomes":
        inserted_refs = _confirm_outcomes(db, batch.staged_rows)
    else:
        raise ValueError("unknown target kind")
    batch.status = "imported"
    batch.confirmed_at = datetime.now(UTC)
    batch.inserted_refs = inserted_refs
    record_db(
        db,
        actor=actor,
        action="import.confirm",
        object_type="import_batch",
        object_id=batch.id,
        change_summary={
            "target_kind": batch.target_kind,
            "inserted_refs": inserted_refs,
        },
    )
    db.commit()
    db.refresh(batch)
    return batch, _count_refs(inserted_refs)


def rollback_import_batch(db: Session, batch_id: int, *, actor: str) -> tuple[ImportBatch, int]:
    batch = _get_batch(db, batch_id, lock=True)
    _ensure_legacy_batch_mutation_allowed(batch, operation="rollback")
    if batch.status != "imported":
        raise ValueError(f"batch status {batch.status} cannot be rolled back")
    refs = batch.inserted_refs or {}
    count = _rollback_refs(db, refs)
    batch.status = "rolled_back"
    batch.rolled_back_at = datetime.now(UTC)
    record_db(
        db,
        actor=actor,
        action="import.rollback",
        object_type="import_batch",
        object_id=batch.id,
        change_summary={"target_kind": batch.target_kind, "rolled_back_count": count},
    )
    db.commit()
    db.refresh(batch)
    return batch, count


def _ensure_kind(kind: str) -> None:
    if kind not in TEMPLATE_HEADERS:
        raise ValueError("unknown import template")


def _get_batch(db: Session, batch_id: int, *, lock: bool = False) -> ImportBatch:
    statement = select(ImportBatch).where(ImportBatch.id == batch_id)
    if lock:
        statement = statement.with_for_update()
    batch = db.execute(statement).scalar_one_or_none()
    if batch is None:
        raise ValueError("batch not found")
    return batch


def _read_rows(kind: str, data: bytes) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    raw_headers = [str(c.value or "").strip() for c in next(ws.iter_rows(max_row=1))]
    label_to_field = {label: field for field, label in FIELD_LABELS.items()}
    normalized_headers = [label_to_field.get(header, header) for header in raw_headers]
    template_headers = TEMPLATE_HEADERS[kind]
    required = REQUIRED_HEADERS[kind]
    missing = [header for header in required if header not in normalized_headers]
    if missing:
        return [], [
            _error(
                1,
                "missing_header",
                f"缺少必填列：{'、'.join(_field_label(field) for field in missing)}",
                field="、".join(_field_label(field) for field in missing),
            )
        ]
    index = {header: normalized_headers.index(header) for header in normalized_headers}
    headers_to_read = [header for header in template_headers if header in index]
    rows: list[dict[str, Any]] = []
    for row_number, row in enumerate(ws.iter_rows(min_row=2), start=2):
        parsed = {
            header: _normalize_import_value(
                header,
                _clean(row[index[header]].value if index[header] < len(row) else None),
            )
            for header in headers_to_read
        }
        for header in template_headers:
            parsed.setdefault(header, None)
        if any(value not in (None, "") for value in parsed.values()):
            parsed["_row_number"] = row_number
            rows.append(parsed)
    return rows, []


def _validate_rows(
    db: Session,
    kind: str,
    rows: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    if kind == "students":
        return _validate_students(db, rows)
    if kind == "jobs":
        return _validate_jobs(rows)
    if kind == "assessments":
        return _validate_assessments(db, rows)
    if kind == "activity_attendance":
        return _validate_activity_attendance(db, rows)
    if kind == "outcomes":
        return _validate_outcomes(db, rows)
    raise ValueError("unknown target kind")


def _validate_students(
    db: Session,
    rows: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    colleges = {row.code: row for row in db.execute(select(College)).scalars() if row.code}
    majors = {(row.college_id, row.code): row for row in db.execute(select(Major)).scalars()}
    existing_numbers = set(db.execute(select(Student.student_no)).scalars())
    seen: set[str] = set()
    staged: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for row in rows:
        row_number = int(row["_row_number"])
        student_no = str(row.get("student_no") or "").strip()
        if not student_no:
            errors.append(
                _error(row_number, "missing_required", "学号不能为空", field="student_no")
            )
            continue
        if student_no in seen or student_no in existing_numbers:
            errors.append(
                _error(
                    row_number,
                    "student_no_duplicate",
                    "学号重复",
                    field="student_no",
                    value=student_no,
                )
            )
            continue
        seen.add(student_no)
        required_missing = [
            field
            for field in ("name", "gender", "college_code", "major_code", "degree_level", "status")
            if not row.get(field)
        ]
        if required_missing:
            errors.append(
                _error(
                    row_number,
                    "missing_required",
                    f"必填内容为空：{'、'.join(_field_label(field) for field in required_missing)}",
                    field=",".join(required_missing),
                )
            )
            continue
        college = colleges.get(str(row["college_code"]))
        if college is None:
            errors.append(
                _error(row_number, "college_code_unknown", "学院代码不在目录", field="college_code")
            )
            continue
        major = majors.get((college.id, str(row["major_code"])))
        if major is None:
            errors.append(
                _error(row_number, "major_code_unknown", "专业代码不在目录", field="major_code")
            )
            continue
        enum_error = _first_enum_error(
            row,
            {
                "gender": _values(Gender),
                "degree_level": _values(DegreeLevel),
                "status": _values(EmploymentStatus),
            },
        )
        if enum_error:
            field, value = enum_error
            errors.append(
                _error(row_number, "enum_out_of_range", "枚举越界", field=field, value=value)
            )
            continue
        try:
            cohort_year = int(row.get("cohort_year"))
            if cohort_year < 2000 or cohort_year > 2100:
                raise ValueError
        except (TypeError, ValueError):
            errors.append(
                _error(
                    row_number,
                    "cohort_year_format",
                    "届别必须是四位年份",
                    field="cohort_year",
                    value=row.get("cohort_year"),
                )
            )
            continue
        staged.append(
            {
                "student_no": student_no,
                "name": str(row["name"]),
                "gender": str(row["gender"]),
                "college_id": college.id,
                "major_id": major.id,
                "degree_level": str(row["degree_level"]),
                "cohort_year": cohort_year,
                "status": str(row["status"]),
                "origin_place": str(row.get("origin_place") or "") or None,
                "is_key_group": str(row.get("is_key_group") or "") or None,
            }
        )
    return staged, errors


def _validate_jobs(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    staged: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for row in rows:
        row_number = int(row["_row_number"])
        if not row.get("company") or not row.get("title") or not row.get("raw_text"):
            errors.append(
                _error(row_number, "missing_required", "企业名称、岗位名称、职位说明不能为空")
            )
            continue
        received_at = _parse_date(row.get("received_at"))
        if row.get("received_at") and received_at is None:
            errors.append(
                _error(
                    row_number, "date_format", "日期请按“2026-03-01”的格式填写", field="received_at"
                )
            )
            continue
        staged.append(
            {
                "company": str(row["company"]),
                "title": str(row["title"]),
                "raw_text": str(row["raw_text"]),
                "received_at": (received_at or datetime.now(UTC).date()).isoformat(),
            }
        )
    return staged, errors


def _validate_assessments(
    db: Session,
    rows: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    students = {row.student_no: row.id for row in db.execute(select(Student)).scalars()}
    staged: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for row in rows:
        row_number = int(row["_row_number"])
        if not row.get("student_no") or not row.get("type") or not row.get("term"):
            errors.append(_error(row_number, "missing_required", "学号、类型、学期不能为空"))
            continue
        if row.get("type") not in _values(AssessmentType):
            errors.append(_error(row_number, "enum_out_of_range", "枚举越界", field="type"))
            continue
        student_id = students.get(str(row["student_no"]))
        if student_id is None:
            errors.append(
                _error(row_number, "student_no_unknown", "学号不存在", field="student_no")
            )
            continue
        score_fields = (
            "academic_basics",
            "learning_ability",
            "professional_confidence",
            "communication",
            "expression",
            "research_interest",
            "innovation_practice",
        )
        answers: dict[str, float] = {}
        invalid_field: str | None = None
        for index, field in enumerate(score_fields, start=1):
            try:
                score = float(row.get(field))
                if score < 1 or score > 5:
                    raise ValueError
            except (TypeError, ValueError):
                invalid_field = field
                break
            answers[f"q{index}"] = score
        if invalid_field is not None:
            errors.append(
                _error(
                    row_number,
                    "score_format",
                    "每项测评评分请填写 1 到 5 之间的数字",
                    field=invalid_field,
                    value=row.get(invalid_field),
                )
            )
            continue
        staged.append(
            {
                "student_id": student_id,
                "type": str(row["type"]),
                "term": str(row["term"]),
                "result_json": {
                    "answers": answers,
                    "summary": str(row.get("summary") or "") or None,
                },
            }
        )
    return staged, errors


def _validate_activity_attendance(
    db: Session,
    rows: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    students = {row.student_no: row.id for row in db.execute(select(Student)).scalars()}
    staged: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for row in rows:
        row_number = int(row["_row_number"])
        checked_at = _parse_date(row.get("checked_at"))
        if not row.get("student_no") or not row.get("activity_name"):
            errors.append(_error(row_number, "missing_required", "学号、活动名称不能为空"))
            continue
        if row.get("checked_at") and checked_at is None:
            errors.append(
                _error(
                    row_number, "date_format", "日期请按“2026-03-01”的格式填写", field="checked_at"
                )
            )
            continue
        if row.get("source") not in _values(AttendanceSource):
            errors.append(_error(row_number, "enum_out_of_range", "枚举越界", field="source"))
            continue
        student_id = students.get(str(row["student_no"]))
        if student_id is None:
            errors.append(
                _error(row_number, "student_no_unknown", "学号不存在", field="student_no")
            )
            continue
        staged.append(
            {
                "student_id": student_id,
                "activity_name": str(row["activity_name"]),
                "source": str(row["source"]),
                "checked_at": (checked_at or datetime.now(UTC).date()).isoformat(),
            }
        )
    return staged, errors


def _validate_outcomes(
    db: Session,
    rows: list[dict[str, Any]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    students = {row.student_no: row.id for row in db.execute(select(Student)).scalars()}
    staged: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for row in rows:
        row_number = int(row["_row_number"])
        occurred_at = _parse_date(row.get("occurred_at"))
        if not row.get("student_no") or not row.get("status") or not row.get("occurred_at"):
            errors.append(_error(row_number, "missing_required", "学号、去向、登记日期不能为空"))
            continue
        if row.get("status") not in _values(EmploymentStatus):
            errors.append(_error(row_number, "enum_out_of_range", "枚举越界", field="status"))
            continue
        if row.get("occurred_at") and occurred_at is None:
            errors.append(
                _error(
                    row_number, "date_format", "日期请按“2026-03-01”的格式填写", field="occurred_at"
                )
            )
            continue
        student_id = students.get(str(row["student_no"]))
        if student_id is None:
            errors.append(
                _error(row_number, "student_no_unknown", "学号不存在", field="student_no")
            )
            continue
        staged.append(
            {
                "student_id": student_id,
                "student_no": str(row["student_no"]),
                "status": str(row["status"]),
                "occurred_at": occurred_at.isoformat() if occurred_at else None,
            }
        )
    return staged, errors


def _confirm_students(db: Session, rows: list[dict[str, Any]], actor: str) -> dict[str, Any]:
    from app.services.risk import RiskInput, evaluate_risk

    students = [Student(**row, created_by=actor) for row in rows]
    db.add_all(students)
    db.flush()
    today = datetime.now(UTC).date()
    for student in students:
        result = evaluate_risk(
            RiskInput(
                status=student.status,
                is_graduation_year=student.cohort_year == today.year,
                current_month=today.month,
                weeks_since_last_job_action=None,
                apply_count=0,
                college_avg_apply=0,
            )
        )
        student.risk_level = result.level.value
        student.risk_reason = "；".join(result.reasons) or None
    return {"students": [student.id for student in students]}


def _confirm_jobs(db: Session, rows: list[dict[str, Any]], actor: str) -> dict[str, Any]:
    jobs = [
        Job(
            company=row["company"],
            title=row["title"],
            raw_text=row["raw_text"],
            source=JobSource.IMPORT.value,
            received_at=datetime.fromisoformat(row["received_at"]).replace(tzinfo=UTC),
            created_by=actor,
        )
        for row in rows
    ]
    db.add_all(jobs)
    db.flush()
    return {"jobs": [job.id for job in jobs]}


def _confirm_assessments(db: Session, rows: list[dict[str, Any]], actor: str) -> dict[str, Any]:
    assessments = [Assessment(**row, created_by=actor) for row in rows]
    db.add_all(assessments)
    db.flush()
    return {"assessments": [row.id for row in assessments]}


def _confirm_activity_attendance(
    db: Session,
    rows: list[dict[str, Any]],
    actor: str,
) -> dict[str, Any]:
    attendance: list[ActivityAttendance] = []
    created_activity_ids: list[int] = []
    for row in rows:
        activity = (
            db.execute(
                select(Activity).where(Activity.name == row["activity_name"]).order_by(Activity.id)
            )
            .scalars()
            .first()
        )
        if activity is None:
            activity = Activity(
                name=row["activity_name"],
                held_at=datetime.fromisoformat(row["checked_at"]).replace(tzinfo=UTC),
                created_by=actor,
            )
            db.add(activity)
            db.flush()
            created_activity_ids.append(activity.id)
        attendance.append(
            ActivityAttendance(
                student_id=row["student_id"],
                activity_id=activity.id,
                source=row["source"],
                checked_at=datetime.fromisoformat(row["checked_at"]).replace(tzinfo=UTC),
                created_by=actor,
            )
        )
    db.add_all(attendance)
    db.flush()
    return {
        "activity_attendance": [row.id for row in attendance],
        "activities": created_activity_ids,
    }


def _confirm_outcomes(db: Session, rows: list[dict[str, Any]]) -> dict[str, Any]:
    changed: list[dict[str, Any]] = []
    for row in rows:
        student = db.get(Student, row["student_id"])
        if student is None:
            continue
        changed.append(
            {
                "student_id": student.id,
                "before": student.status,
                "before_occurred_at": (
                    student.status_occurred_at.isoformat()
                    if student.status_occurred_at is not None
                    else None
                ),
                "after": row["status"],
                "after_occurred_at": row["occurred_at"],
            }
        )
        student.status = row["status"]
        student.status_occurred_at = datetime.fromisoformat(row["occurred_at"]).replace(tzinfo=UTC)
    db.flush()
    return {"outcomes": changed}


def _rollback_refs(db: Session, refs: dict[str, Any]) -> int:
    count = 0
    for table_key, model in (
        ("activity_attendance", ActivityAttendance),
        ("assessments", Assessment),
        ("jobs", Job),
        ("students", Student),
        ("activities", Activity),
    ):
        ids = refs.get(table_key) or []
        if ids:
            result = db.execute(delete(model).where(model.id.in_(ids)))
            count += result.rowcount or len(ids)
    outcomes = refs.get("outcomes") or []
    for item in outcomes:
        student = db.get(Student, item["student_id"])
        if student is not None:
            student.status = item["before"]
            before_occurred_at = item.get("before_occurred_at")
            student.status_occurred_at = (
                datetime.fromisoformat(before_occurred_at) if before_occurred_at else None
            )
            count += 1
    db.flush()
    return count


def _count_refs(refs: dict[str, Any]) -> int:
    count = 0
    for key, value in refs.items():
        if key == "activities":
            continue
        count += len(value) if isinstance(value, list) else 0
    return count


def _error(
    row_number: int,
    code: str,
    reason: str,
    *,
    field: str | None = None,
    value: Any = None,
) -> dict[str, Any]:
    return {
        "row_number": row_number,
        "code": code,
        "reason": reason,
        "field": "、".join(_field_label(item) for item in field.split(",")) if field else None,
        "value": str(value) if value is not None else None,
    }


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        return value.strip()
    return value


def _values(enum_cls: Any) -> set[str]:
    return {item.value for item in enum_cls}


def _field_label(field: str) -> str:
    return FIELD_LABELS.get(field, field)


def _normalize_import_value(field: str, value: Any) -> Any:
    if not isinstance(value, str):
        return value
    return IMPORT_VALUE_ALIASES.get(field, {}).get(value, value)


def _first_enum_error(row: dict[str, Any], allowed: dict[str, set[str]]) -> tuple[str, Any] | None:
    for field, values in allowed.items():
        if row.get(field) not in values:
            return field, row.get(field)
    return None


def _parse_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value)
        except ValueError:
            return None
    return None
