"""T4 岗位入库、复核、匹配与直推服务。"""

from __future__ import annotations

import csv
import hashlib
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from typing import Any, Literal

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session, joinedload

from app.core.audit import record_db
from app.llm.parsers.jd import build_major_index, parse_jd
from app.llm.provider import Provider
from app.models import College, Job, JobParse, Major, Match, Student
from app.models.enums import JobSource
from app.services.matching import load_match_weights, score_student_match
from app.services.operation_receipts import claim_operation, finish_operation

REVIEW_FIELDS = {
    "degree_req",
    "major_codes",
    "direction_keywords",
    "location",
    "salary_min",
    "salary_max",
    "headcount",
    "deadline",
    "parse_confidence",
    "needs_review",
    "review_reason",
}


class MatchNotFoundError(ValueError):
    """目标学生没有该岗位的匹配事实。"""


class MatchNotEligibleError(ValueError):
    """匹配存在，但当前分数不允许推荐。"""


class MatchForbiddenError(PermissionError):
    """当前用户无权处理目标学生。"""


class MatchInputNotFoundError(LookupError):
    """岗位匹配所需的业务事实不存在。"""


class MatchJobNotFoundError(MatchInputNotFoundError):
    """目标岗位不存在。"""


class MatchJobParseNotFoundError(MatchInputNotFoundError):
    """目标岗位尚无解析事实。"""


class JobParseReviewConflictError(RuntimeError):
    """岗位解析已被较新的人工复核修改。"""

    def __init__(self, *, expected_revision: int, current_revision: int) -> None:
        self.expected_revision = expected_revision
        self.current_revision = current_revision
        super().__init__(
            "job parse review revision conflict: "
            f"expected {expected_revision}, current {current_revision}"
        )


def _to_date(value: str | date | None) -> date | None:
    if value is None or isinstance(value, date):
        return value
    try:
        return date.fromisoformat(value)
    except ValueError:
        return None


def _jsonable(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, date | datetime):
        return value.isoformat()
    if isinstance(value, list):
        return [_jsonable(v) for v in value]
    if isinstance(value, dict):
        return {k: _jsonable(v) for k, v in value.items()}
    return value


def create_job(
    db: Session,
    *,
    company: str,
    title: str,
    raw_text: str,
    actor: str,
    idempotency_key: str,
    source: str = JobSource.FORM.value,
) -> Job:
    receipt, replayed = claim_operation(
        db,
        namespace="job.create",
        idempotency_key=idempotency_key,
        actor=actor,
        payload={
            "company": company,
            "title": title,
            "raw_text": raw_text,
            "source": source,
        },
    )
    if replayed:
        job_id = receipt.result_json.get("job_id")
        job = db.get(Job, job_id) if isinstance(job_id, int) else None
        if job is None:
            raise ValueError("idempotent job receipt has no committed job")
        return job
    job = Job(
        company=company,
        title=title,
        raw_text=raw_text,
        source=source,
        received_at=datetime.now(UTC),
        created_by=actor,
    )
    db.add(job)
    db.flush()
    finish_operation(receipt, result={"job_id": job.id})
    record_db(
        db,
        actor=actor,
        action="job.create",
        object_type="job",
        object_id=job.id,
        correlation_id=f"JOB-CREATE-{idempotency_key}",
        change_summary={"after": {"company": company, "title": title, "source": source}},
    )
    db.commit()
    db.refresh(job)
    return job


def import_jobs_from_xlsx(
    db: Session,
    data: bytes,
    *,
    actor: str,
    idempotency_key: str,
) -> list[Job]:
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        workbook_rows = list(ws.iter_rows())
    except Exception as exc:
        raise ValueError("工作簿内容损坏或不是有效的 Excel 文件") from exc
    if not workbook_rows:
        raise ValueError("工作簿没有可读取的内容")
    headers = [str(c.value or "").strip() for c in workbook_rows[0]]
    index = {h: i for i, h in enumerate(headers)}
    required = ["company", "title", "raw_text"]
    missing = [h for h in required if h not in index]
    if missing:
        raise ValueError(f"xlsx 缺少列：{','.join(missing)}")

    jobs: list[Job] = []
    for row in workbook_rows[1:]:
        company = str(row[index["company"]].value or "").strip()
        title = str(row[index["title"]].value or "").strip()
        raw_text = str(row[index["raw_text"]].value or "").strip()
        if not company and not title and not raw_text:
            continue
        if not company or not title or not raw_text:
            raise ValueError("xlsx 存在 company/title/raw_text 不完整的行")
        jobs.append(
            Job(
                company=company,
                title=title,
                raw_text=raw_text,
                source=JobSource.IMPORT.value,
                received_at=datetime.now(UTC),
                created_by=actor,
            )
        )
    receipt, replayed = claim_operation(
        db,
        namespace="job.import_xlsx",
        idempotency_key=idempotency_key,
        actor=actor,
        payload={
            "file_sha256": hashlib.sha256(data).hexdigest(),
        },
    )
    if replayed:
        raw_ids = receipt.result_json.get("job_ids")
        if not isinstance(raw_ids, list) or not all(
            isinstance(job_id, int) for job_id in raw_ids
        ):
            raise ValueError("idempotent job import receipt has no committed jobs")
        stored = list(db.execute(select(Job).where(Job.id.in_(raw_ids))).scalars())
        by_id = {job.id: job for job in stored}
        if len(by_id) != len(raw_ids):
            raise ValueError("idempotent job import receipt references missing jobs")
        return [by_id[job_id] for job_id in raw_ids]
    db.add_all(jobs)
    db.flush()
    finish_operation(receipt, result={"job_ids": [job.id for job in jobs]})
    record_db(
        db,
        actor=actor,
        action="job.import_xlsx",
        object_type="job",
        correlation_id=f"JOB-IMPORT-{idempotency_key}",
        change_summary={"count": len(jobs)},
    )
    db.commit()
    for job in jobs:
        db.refresh(job)
    return jobs


def parse_job(
    db: Session,
    job_id: int,
    provider: Provider,
    *,
    actor: str,
    commit: bool = True,
    correlation_id: str | None = None,
) -> JobParse:
    job = db.execute(
        select(Job).where(Job.id == job_id).with_for_update()
    ).scalar_one_or_none()
    if job is None:
        raise ValueError("job not found")
    existing = db.execute(
        select(JobParse).where(JobParse.job_id == job_id)
    ).scalar_one_or_none()
    if existing is not None:
        return existing
    result = parse_jd(provider, job.raw_text, build_major_index(db))
    parse = JobParse(
        job_id=job.id,
        degree_req=result.degree_req,
        major_codes=result.major_codes,
        direction_keywords=result.direction_keywords,
        location=result.location,
        salary_min=result.salary_min,
        salary_max=result.salary_max,
        headcount=result.headcount,
        deadline=_to_date(result.deadline),
        parse_confidence=result.parse_confidence,
        llm_meta={"provider": provider.name, "prompt_ver": result.prompt_ver},
        needs_review=result.needs_review,
        review_reason=result.review_reason,
        created_by=actor,
    )
    db.add(parse)
    db.flush()
    record_db(
        db,
        actor=actor,
        action="job_parse.create",
        object_type="job_parse",
        object_id=parse.id,
        correlation_id=correlation_id,
        change_summary={"job_id": job.id, "needs_review": parse.needs_review},
    )
    if commit:
        db.commit()
        db.refresh(parse)
    return parse


def review_job_parse(
    db: Session,
    parse_id: int,
    *,
    actor: str,
    updates: dict[str, Any],
    expected_revision: int,
    commit: bool = True,
    correlation_id: str | None = None,
) -> JobParse:
    parse = db.execute(
        select(JobParse).where(JobParse.id == parse_id).with_for_update()
    ).scalar_one_or_none()
    if parse is None:
        raise ValueError("job parse not found")
    if parse.review_revision != expected_revision:
        raise JobParseReviewConflictError(
            expected_revision=expected_revision,
            current_revision=parse.review_revision,
        )
    before = {field: _jsonable(getattr(parse, field)) for field in REVIEW_FIELDS}
    before["review_revision"] = parse.review_revision
    for field, value in updates.items():
        if field in REVIEW_FIELDS:
            setattr(parse, field, value)
    parse.review_revision += 1
    after = {field: _jsonable(getattr(parse, field)) for field in REVIEW_FIELDS}
    after["review_revision"] = parse.review_revision
    record_db(
        db,
        actor=actor,
        action="job_parse.review",
        object_type="job_parse",
        object_id=parse.id,
        correlation_id=correlation_id,
        change_summary={"before": before, "after": after},
    )
    if commit:
        db.commit()
        db.refresh(parse)
    return parse


def list_successful_parsed_jobs(db: Session, *, limit: int = 300) -> list[dict]:
    rows = db.execute(
        select(JobParse, Job)
        .join(Job, JobParse.job_id == Job.id)
        .where(JobParse.needs_review.is_(False))
        .where(~Job.company.contains("??"), ~Job.title.contains("??"))
        .order_by(Job.received_at.desc(), Job.id.desc())
        .limit(limit)
    ).all()
    major_names = dict(db.execute(select(Major.code, Major.name)).all())
    match_counts = {
        job_id: (int(recommended), int(pushed))
        for job_id, recommended, pushed in db.execute(
            select(
                Match.job_id,
                func.count(Match.id).filter(Match.score > 0),
                func.count(Match.id).filter(Match.pushed_at.is_not(None)),
            ).group_by(Match.job_id)
        ).all()
    }
    result = [
        {
            "parse_id": parse.id,
            "job_id": job.id,
            "company": job.company,
            "title": job.title,
            "label": f"{job.company} · {job.title}",
            "degree_req": parse.degree_req,
            "major_names": [major_names.get(code, code) for code in parse.major_codes],
            "location": parse.location,
            "headcount": parse.headcount,
            "recommended_count": match_counts.get(job.id, (0, 0))[0],
            "pushed_count": match_counts.get(job.id, (0, 0))[1],
            "parse_confidence": float(parse.parse_confidence)
            if parse.parse_confidence is not None
            else None,
            "received_at": job.received_at,
        }
        for parse, job in rows
    ]
    return result


def calculate_matches_for_job(
    db: Session,
    job_id: int,
    *,
    actor: str,
    limit: int | None = None,
    correlation_id: str | None = None,
) -> list[Match]:
    locked_job = db.execute(
        select(Job).where(Job.id == job_id).with_for_update()
    ).scalar_one_or_none()
    if locked_job is None:
        raise MatchJobNotFoundError("job not found")
    parse = db.execute(
        select(JobParse).where(JobParse.job_id == job_id)
    ).scalar_one_or_none()
    if parse is None:
        raise MatchJobParseNotFoundError("job parse not found")
    weights = load_match_weights(db)
    students = db.execute(
        select(Student)
        .options(
            joinedload(Student.major),
            joinedload(Student.direction),
            joinedload(Student.college),
        )
        .order_by(Student.id)
    ).scalars().all()
    scored = []
    for student in students:
        result = score_student_match(parse, student, weights)
        scored.append((student, result))
    scored.sort(key=lambda item: item[1].score, reverse=True)
    if limit:
        scored = scored[:limit]

    # Recalculation updates the recommendation projection, but pushed/feedback
    # rows are durable business facts used by the student journey.  Replacing
    # the whole job match set would erase those facts and their source IDs.
    existing = list(
        db.execute(
            select(Match).where(
                Match.job_id == job_id,
                Match.student_id.is_not(None),
            ).order_by(Match.id).with_for_update()
        ).scalars()
    )
    existing_by_student = {
        int(match.student_id): match
        for match in existing
        if match.student_id is not None
    }
    selected_student_ids: set[int] = set()
    matches: list[Match] = []
    created_count = 0
    updated_count = 0
    for student, result in scored:
        selected_student_ids.add(student.id)
        match = existing_by_student.get(student.id)
        if match is None:
            match = Match(
                job_id=job_id,
                student_id=student.id,
                score=result.score,
                reason_text=result.reason_text,
                created_by=actor,
            )
            db.add(match)
            created_count += 1
        else:
            match.score = result.score
            match.reason_text = result.reason_text
            updated_count += 1
        matches.append(match)

    removed_count = 0
    preserved_fact_count = 0
    for stale in existing:
        if stale.student_id in selected_student_ids:
            continue
        if stale.pushed_at is not None or stale.feedback is not None:
            # A no-longer-top recommendation remains available as immutable
            # journey evidence.  Its last evaluated score/reason are retained.
            preserved_fact_count += 1
            matches.append(stale)
        else:
            db.delete(stale)
            removed_count += 1
    record_db(
        db,
        actor=actor,
        action="job_match.calculate",
        object_type="job",
        object_id=job_id,
        correlation_id=correlation_id,
        change_summary={
            "selected_count": len(scored),
            "created_count": created_count,
            "updated_count": updated_count,
            "removed_unpushed_count": removed_count,
            "preserved_business_fact_count": preserved_fact_count,
            "weights": weights.__dict__,
        },
    )
    db.flush()
    return matches


def list_matches_for_job(
    db: Session,
    job_id: int,
    *,
    college_id: int | None = None,
    positive_only: bool = False,
) -> list[dict]:
    stmt = (
        select(Match, Student, College)
        .join(Student, Match.student_id == Student.id)
        .join(College, Student.college_id == College.id)
        .where(Match.job_id == job_id)
        .order_by(Match.score.desc(), Match.id)
    )
    if college_id is not None:
        stmt = stmt.where(Student.college_id == college_id)
    if positive_only:
        stmt = stmt.where(Match.score > 0)
    rows = db.execute(stmt).all()
    return [
        {
            "id": match.id,
            "job_id": match.job_id,
            "student_id": student.id,
            "student_no": student.student_no,
            "student_name": student.name,
            "college_id": student.college_id,
            "college_name": college.name,
            "major_id": student.major_id,
            "score": float(match.score),
            "reason_text": match.reason_text,
            "pushed_at": match.pushed_at,
        }
        for match, student, college in rows
    ]


def confirm_push_matches(
    db: Session,
    job_id: int,
    *,
    actor: str,
    college_id: int | None = None,
    student_id: int | None = None,
) -> dict[str, int | str]:
    locked_job_id = db.execute(
        select(Job.id).where(Job.id == job_id).with_for_update()
    ).scalar_one_or_none()
    if locked_job_id is None:
        raise MatchNotFoundError("没有找到该岗位。")
    stmt = (
        select(Match, Student)
        .join(Student, Match.student_id == Student.id)
        .where(Match.job_id == job_id)
        .order_by(Match.id)
        .with_for_update(of=Match)
    )
    if student_id is not None:
        stmt = stmt.where(Match.student_id == student_id)
    rows = db.execute(stmt).all()

    if not rows:
        raise MatchNotFoundError("没有找到这名学生与该岗位的匹配记录。")

    if student_id is not None and college_id is not None:
        target_student = rows[0][1]
        if target_student.college_id != college_id:
            raise MatchForbiddenError("无权向其他学院学生发送岗位。")

    scoped_rows = [
        (match, student)
        for match, student in rows
        if college_id is None or student.college_id == college_id
    ]
    if not scoped_rows:
        raise MatchForbiddenError("当前账号无权处理这些学生。")

    # A previously pushed row is already a durable fact, even if a later
    # recommendation recalculation lowers its projection score to zero.
    # Recognize that replay before applying eligibility to never-pushed rows.
    eligible = [
        match
        for match, _student in scoped_rows
        if match.pushed_at is not None or float(match.score) > 0
    ]
    if not eligible:
        raise MatchNotEligibleError("当前匹配度为 0，不能发送；请调整岗位条件后重新匹配。")

    pending = [match for match in eligible if match.pushed_at is None]
    if not pending:
        return {
            "pushed": 0,
            "already_pushed": len(eligible),
            "status": "already_pushed",
            "message": "该岗位此前已发送，无需重复操作。",
        }

    now = datetime.now(UTC)
    for row in pending:
        row.pushed_at = now
    record_db(
        db,
        actor=actor,
        action="job_match.push",
        object_type="job",
        object_id=job_id,
        change_summary={
            "count": len(pending),
            "already_pushed": len(eligible) - len(pending),
            "college_id": college_id,
            "student_id": student_id,
        },
    )
    db.commit()
    return {
        "pushed": len(pending),
        "already_pushed": len(eligible) - len(pending),
        "status": "pushed",
        "message": "岗位发送成功。",
    }


def export_matches_csv(
    rows: list[dict], *, header_language: Literal["zh", "en"] = "zh"
) -> str:
    headers = {
        "zh": ["学号", "姓名", "学院编号", "专业编号", "匹配度", "推荐理由"],
        "en": ["student_no", "student_name", "college_id", "major_id", "score", "reason_text"],
    }
    if header_language not in headers:
        raise ValueError("Unsupported CSV header language")
    out = StringIO()
    out.write("\ufeff")
    writer = csv.writer(out)
    writer.writerow(headers[header_language])
    for row in rows:
        writer.writerow([
            row["student_no"],
            row["student_name"],
            row["college_id"],
            row["major_id"],
            row["score"],
            row["reason_text"],
        ])
    return out.getvalue()
