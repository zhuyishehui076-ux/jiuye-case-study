"""T6 M2 一生一策帮扶台账服务。"""

from __future__ import annotations

import html
import re
from datetime import UTC, datetime
from io import BytesIO
from typing import Any

from docx import Document
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.audit import record_db
from app.core.rbac import Role as RbacRole
from app.models import AssistCase, AssistRecord, College, Student, User
from app.models.enums import AssistActionType, AssistStatus, RiskLevel
from app.models.user import Role as UserRole
from app.services.assist_idempotency import (
    assist_record_request_hash,
    validate_assist_idempotency_key,
)
from app.services.profiles import derive_journey_stage
from app.services.student_feedback import match_feedback_by_student


class AssistRecordIdempotencyConflict(ValueError):
    """The key already belongs to a different semantic request."""

ASSIST_XLSX_HEADERS = [
    "学号",
    "姓名",
    "学院",
    "专业",
    "预警原因",
    "帮扶状态",
    "就业去向",
    "结果日期",
    "闭环阶段",
    "动作汇总",
]

_TAG_RE = re.compile(r"<[^>]+>")


def _plain_text(value: str | None) -> str:
    if not value:
        return ""
    return html.unescape(_TAG_RE.sub("", value)).strip()


def _user_assignments(db: Session) -> tuple[dict[int, int], int | None]:
    counselors = db.execute(
        select(User)
        .join(UserRole, User.role_id == UserRole.id)
        .where(UserRole.name == RbacRole.COUNSELOR.value)
        .order_by(User.id)
    ).scalars().all()
    by_college = {
        user.college_id: user.id
        for user in counselors
        if user.college_id is not None
    }
    fallback = counselors[0].id if counselors else None
    if fallback is None:
        fallback = db.execute(
            select(User.id)
            .join(UserRole, User.role_id == UserRole.id)
            .where(UserRole.name.in_([RbacRole.EMPLOYMENT_OFFICE.value, RbacRole.ADMIN.value]))
            .order_by(User.id)
        ).scalar_one_or_none()
    return by_college, fallback


def ensure_assist_cases_for_red_students(
    db: Session,
    *,
    actor: str,
    opened_at: datetime | None = None,
) -> list[AssistCase]:
    """为所有 red 学生补建一生一策档案；已存在档案不重复建。"""
    existing_student_ids = set(
        db.execute(select(AssistCase.student_id)).scalars().all()
    )
    by_college, fallback_user_id = _user_assignments(db)
    students = db.execute(
        select(Student)
        .where(Student.risk_level == RiskLevel.RED.value)
        .order_by(Student.id)
    ).scalars().all()
    created: list[AssistCase] = []
    now = opened_at or datetime.now(UTC)
    for student in students:
        if student.id in existing_student_ids:
            continue
        assigned_to = by_college.get(student.college_id, fallback_user_id)
        case = AssistCase(
            student_id=student.id,
            opened_at=now,
            opened_reason=student.risk_reason or "系统识别为优先跟进学生，已自动建档",
            assigned_to=assigned_to,
            status=AssistStatus.OPEN.value,
            created_by=actor,
        )
        db.add(case)
        db.flush()
        record_db(
            db,
            actor=actor,
            action="assist_case.auto_create",
            object_type="assist_case",
            object_id=case.id,
            change_summary={
                "student_id": student.id,
                "risk_level": student.risk_level,
                "assigned_to": assigned_to,
            },
        )
        created.append(case)
    db.flush()
    return created


def _record_summary(records: list[AssistRecord]) -> str:
    if not records:
        return ""
    return "；".join(
        f"{r.occurred_at.date().isoformat()} {r.action_type} {_plain_text(r.note)}"
        for r in records
    )


def list_assist_cases(
    db: Session,
    *,
    college_id: int | None = None,
) -> list[dict[str, Any]]:
    stmt = (
        select(AssistCase, Student, College)
        .join(Student, AssistCase.student_id == Student.id)
        .join(College, Student.college_id == College.id)
        .order_by(AssistCase.id)
    )
    if college_id is not None:
        stmt = stmt.where(Student.college_id == college_id)
    rows = db.execute(stmt).all()
    case_ids = [case.id for case, _student, _college in rows]
    records_by_case: dict[int, list[AssistRecord]] = {case_id: [] for case_id in case_ids}
    if case_ids:
        records = db.execute(
            select(AssistRecord)
            .where(AssistRecord.case_id.in_(case_ids))
            .order_by(AssistRecord.occurred_at, AssistRecord.id)
        ).scalars().all()
        for record in records:
            records_by_case.setdefault(record.case_id, []).append(record)
    feedback_by_student = match_feedback_by_student(
        db,
        [student.id for _case, student, _college in rows],
    )
    result: list[dict[str, Any]] = []
    for case, student, college in rows:
        case_records = records_by_case.get(case.id, [])
        match_feedback = feedback_by_student.get(student.id, [])
        journey_stage, journey_stage_index = derive_journey_stage(
            student.status,
            len(case_records),
            match_feedback,
        )
        result.append({
            "id": case.id,
            "student_id": student.id,
            "student_no": student.student_no,
            "student_name": student.name,
            "college_id": student.college_id,
            "college_name": college.name,
            "major_id": student.major_id,
            "risk_reason": student.risk_reason,
            "opened_reason": case.opened_reason,
            "assigned_to": case.assigned_to,
            "status": case.status,
            "opened_at": case.opened_at,
            "record_count": len(case_records),
            "action_summary": _record_summary(case_records),
            "employment_status": student.status,
            "status_occurred_at": student.status_occurred_at,
            "journey_stage": journey_stage,
            "journey_stage_index": journey_stage_index,
            "match_feedback": match_feedback,
        })
    return result


def get_case_for_write(
    db: Session,
    case_id: int,
    *,
    college_id: int | None = None,
) -> AssistCase:
    stmt = (
        select(AssistCase)
        .join(Student, AssistCase.student_id == Student.id)
        .where(AssistCase.id == case_id)
    )
    if college_id is not None:
        stmt = stmt.where(Student.college_id == college_id)
    case = db.execute(stmt).scalar_one_or_none()
    if case is None:
        raise ValueError("assist case not found or forbidden")
    return case


def append_assist_record(
    db: Session,
    case_id: int,
    *,
    action_type: str,
    note: str | None,
    actor: str,
    idempotency_key: str,
    actor_id: int | None = None,
    actor_role: str | None = None,
    college_id: int | None = None,
) -> tuple[AssistRecord, bool]:
    # Scope authorization intentionally precedes every key lookup. A caller
    # cannot use a known key to discover or replay another college's record.
    case = get_case_for_write(db, case_id, college_id=college_id)
    if action_type not in {item.value for item in AssistActionType}:
        raise ValueError("invalid assist action type")
    idempotency_key = validate_assist_idempotency_key(idempotency_key)
    request_hash = assist_record_request_hash(
        case_id=case.id,
        action_type=action_type,
        note=note,
        actor=actor,
    )
    existing = db.scalar(
        select(AssistRecord).where(
            AssistRecord.idempotency_key == idempotency_key
        )
    )
    if existing is not None:
        if existing.request_hash != request_hash:
            raise AssistRecordIdempotencyConflict(
                "Idempotency-Key was already used for a different assist record request"
            )
        return existing, True

    record = AssistRecord(
        case_id=case.id,
        action_type=action_type,
        note=note,
        actor=actor,
        occurred_at=datetime.now(UTC),
        idempotency_key=idempotency_key,
        request_hash=request_hash,
        created_by=actor,
    )
    db.add(record)
    try:
        db.flush()
    except IntegrityError:
        # Two requests can pass the optimistic lookup together. PostgreSQL's
        # unique constraint elects one writer; the loser rolls back, repeats
        # scope authorization, then resolves the committed winner as a replay.
        db.rollback()
        get_case_for_write(db, case_id, college_id=college_id)
        existing = db.scalar(
            select(AssistRecord).where(
                AssistRecord.idempotency_key == idempotency_key
            )
        )
        if existing is None:
            raise
        if existing.request_hash != request_hash:
            raise AssistRecordIdempotencyConflict(
                "Idempotency-Key was already used for a different assist record request"
            ) from None
        return existing, True
    record_db(
        db,
        actor=actor,
        actor_id=actor_id,
        actor_role=actor_role,
        source_channel="management.assist",
        correlation_id=f"ASSIST-RECORD-{idempotency_key}",
        action="assist_record.append",
        object_type="assist_record",
        object_id=record.id,
        change_summary={"case_id": case.id, "action_type": action_type},
    )
    db.commit()
    db.refresh(record)
    return record, False


def export_assist_xlsx(rows: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "一生一策帮扶台账"
    ws.append(ASSIST_XLSX_HEADERS)
    for row in rows:
        ws.append(
            [
                row["student_no"],
                row["student_name"],
                row["college_name"],
                row["major_id"],
                row["risk_reason"] or row["opened_reason"] or "",
                row["status"],
                row["employment_status"],
                row["status_occurred_at"].date().isoformat()
                if row["status_occurred_at"]
                else "",
                row["journey_stage"],
                row["action_summary"],
            ]
        )
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_assist_docx(rows: list[dict[str, Any]]) -> bytes:
    doc = Document()
    doc.add_heading("一生一策帮扶台账", level=1)
    for idx, row in enumerate(rows, start=1):
        if idx > 1:
            doc.add_page_break()
        doc.add_heading(f"学生：{row['student_name']}（{row['student_no']}）", level=2)
        doc.add_paragraph(f"学院：{row['college_name']}")
        doc.add_paragraph(f"帮扶状态：{row['status']}")
        doc.add_paragraph(f"闭环阶段：{row['journey_stage']}")
        doc.add_paragraph(f"就业去向：{row['employment_status']}")
        doc.add_paragraph(
            "结果日期："
            + (row["status_occurred_at"].date().isoformat() if row["status_occurred_at"] else "尚未确认")
        )
        doc.add_paragraph(f"建档原因：{row['risk_reason'] or row['opened_reason'] or ''}")
        doc.add_heading("帮扶记录", level=3)
        doc.add_paragraph(row["action_summary"] or "暂无记录")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
