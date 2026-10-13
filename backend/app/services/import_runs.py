"""T16 production data onboarding: five-file preflight, correction, confirm and rollback.

The T7 single-file importer remains available for historical contracts.  This module is
the production 0->1 path and deliberately keeps every business write behind one
PostgreSQL transaction owned by an ``ImportRun``.
"""

from __future__ import annotations

import json
import re
from collections import Counter
from dataclasses import dataclass
from dataclasses import field as dataclass_field
from datetime import UTC, date, datetime
from hashlib import sha256
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import MetaData, Table, delete, func, inspect, or_, select, text, tuple_
from sqlalchemy.exc import IntegrityError, OperationalError
from sqlalchemy.orm import Session, aliased, defer

from app.core.audit import record_db
from app.core.config import ConfigError, Profile, settings
from app.core.crypto import decrypt_field, encrypt_field
from app.models import (
    AcademicClass,
    AssistCase,
    AssistRecord,
    AuditLog,
    College,
    ImportBatch,
    ImportRun,
    ImportTargetClaim,
    JobAction,
    Major,
    Student,
)
from app.models.enums import DegreeLevel, EmploymentStatus, Gender, JobActionType, KeyGroup
from app.services.dashboard import refresh_dashboard_views
from app.services.operation_receipts import claim_operation, finish_operation
from app.services.thresholds import load_risk_thresholds

T16_TEMPLATE_VERSION = "t16.v1"
T16_KIND_ORDER = ("colleges", "majors", "classes", "students", "outcomes")
T16_RUN_LIST_LIMIT = 100
MAX_IMPORT_ROWS = 20_000
MAX_IMPORT_EXTRA_COLUMNS = 8
INLINE_IMPORT_ISSUE_LIMIT = 100
DEPENDENT_FACT_SNAPSHOT_VERSION = 1
INLINE_IMPORT_EVIDENCE_LIMIT = 20
T16_MANAGEMENT_SOURCE = "management_api"
# Leave protocol headroom below PostgreSQL's 65,535 bind-parameter ceiling.
# Tuple predicates divide this budget by their arity before issuing a query.
T16_QUERY_BIND_PARAMETER_BUDGET = 50_000
# Very large row-valued IN expressions can exhaust PostgreSQL's parser stack
# before reaching the bind ceiling.  Bound item count independently as well.
T16_QUERY_MAX_ITEMS_PER_STATEMENT = 2_000
T16_CONFIRM_ADVISORY_MUTEX_KEY = "t16:confirm:global:v1"

IMPORT_EVIDENCE_LABELS = {
    "import.run.create": "创建启用记录",
    "import.run.upload": "上传并检查资料",
    "import.run.upload_idempotent": "重复资料已安全复用",
    "import.run.correct": "修正问题数据",
    "import.run.preflight": "完成跨表预检",
    "import.run.confirm": "确认导入并启用",
    "import.run.confirm_idempotent": "重复确认已安全复用",
    "import.run.confirm_blocked": "确认被安全阻止",
    "import.run.rollback": "撤销本次启用",
    "import.run.rollback_idempotent": "重复撤销已安全复用",
    "import.run.rollback_blocked": "撤销被安全阻止",
    "import.run.receipt_export": "导出对账回执",
    "import.run.batch_receipt_export": "导出单类资料检查结果",
}

T16_HEADERS: dict[str, list[str]] = {
    "colleges": ["college_code", "college_name"],
    "majors": ["college_code", "major_code", "major_name", "degree_level"],
    "classes": [
        "college_code",
        "major_code",
        "degree_level",
        "cohort_year",
        "class_code",
        "class_name",
    ],
    "students": [
        "student_no",
        "name",
        "gender",
        "college_code",
        "major_code",
        "degree_level",
        "cohort_year",
        "class_code",
        "origin_place",
        "contact",
        "is_key_group",
    ],
    "outcomes": ["student_no", "status", "occurred_at"],
}

# Every field except cohort/year and occurrence date carries text semantics.
# Excel may coerce any of them (not only identifiers) into numbers, booleans or
# dates before upload, irreversibly changing the operator's original text.
# Templates therefore format every text-semantic column as text and parsing
# accepts only raw strings instead of guessing a representation.
T16_TEXT_FIELDS: dict[str, frozenset[str]] = {
    "colleges": frozenset({"college_code", "college_name"}),
    "majors": frozenset(
        {"college_code", "major_code", "major_name", "degree_level"}
    ),
    "classes": frozenset(
        {"college_code", "major_code", "degree_level", "class_code", "class_name"}
    ),
    "students": frozenset(
        {
            "student_no",
            "name",
            "gender",
            "college_code",
            "major_code",
            "degree_level",
            "class_code",
            "origin_place",
            "contact",
            "is_key_group",
        }
    ),
    "outcomes": frozenset({"student_no", "status"}),
}

T16_LABELS: dict[str, str] = {
    "college_code": "学院代码",
    "college_name": "学院名称",
    "major_code": "专业代码",
    "major_name": "专业名称",
    "degree_level": "学历层次",
    "cohort_year": "毕业届别",
    "class_code": "班级代码",
    "class_name": "班级名称",
    "student_no": "学号",
    "name": "姓名",
    "gender": "性别",
    "origin_place": "生源地",
    "contact": "联系方式",
    "is_key_group": "重点群体",
    "status": "毕业去向",
    "occurred_at": "登记日期",
}

T16_TITLES = {
    "colleges": "学院",
    "majors": "专业",
    "classes": "班级",
    "students": "学生",
    "outcomes": "毕业去向",
}

T16_FILENAMES = {
    kind: f"{index:02d}-{T16_TITLES[kind]}数据导入模板.xlsx"
    for index, kind in enumerate(T16_KIND_ORDER, start=1)
}

T16_EXAMPLES: dict[str, list[Any]] = {
    "colleges": ["C01", "材料科学与工程学院"],
    "majors": ["C01", "080401", "材料科学与工程", DegreeLevel.BACHELOR.value],
    "classes": ["C01", "080401", DegreeLevel.BACHELOR.value, 2026, "CL2601", "材料 2601 班"],
    "students": [
        "20260001",
        "张同学",
        Gender.MALE.value,
        "C01",
        "080401",
        DegreeLevel.BACHELOR.value,
        2026,
        "CL2601",
        "湖北武汉",
        "SYNTHETIC-CONTACT-001",
        "",
    ],
    "outcomes": ["20260001", EmploymentStatus.SIGNED.value, "2026-04-01"],
}

IDENTIFIER_PATTERNS = {
    "college_code": re.compile(r"^[A-Za-z0-9_-]{1,32}$"),
    "major_code": re.compile(r"^[A-Za-z0-9_.-]{1,16}$"),
    "class_code": re.compile(r"^[A-Za-z0-9_-]{1,32}$"),
    "student_no": re.compile(r"^[A-Za-z0-9_-]{1,32}$"),
}

SENSITIVE_COLUMNS = {"student_no", "name", "origin_place", "contact", "is_key_group"}
_ENCRYPTED_STAGING_KEY = "__t16_encrypted_v1__"


class ImportConflictError(ValueError):
    """A safe, user-actionable 409 rather than an internal server error."""


@dataclass
class _ImportLookupScope:
    """Natural keys that one import run is allowed to inspect."""

    college_codes: set[str] = dataclass_field(default_factory=set)
    college_names: set[str] = dataclass_field(default_factory=set)
    major_keys: set[tuple[str, str, str]] = dataclass_field(default_factory=set)
    class_keys: set[tuple[str, str, str, int, str]] = dataclass_field(default_factory=set)
    student_numbers: set[str] = dataclass_field(default_factory=set)
    # Existing organization ownership is required only when validating a
    # student master row.  Outcome-only lookups never need those parent rows.
    student_scope_numbers: set[str] = dataclass_field(default_factory=set)


@dataclass(frozen=True)
class _ImportLookupState:
    colleges: dict[str, College]
    college_name_owners: dict[str, str | None]
    college_ids: dict[int, str | None]
    majors: dict[tuple[str, str, str], Major]
    major_ids: dict[int, tuple[str, str, str]]
    classes: dict[tuple[str, str, str, int, str], AcademicClass]
    class_ids: dict[int, tuple[str, str, str, int, str]]
    students: dict[str, Student]


def _require_production_import_profile() -> None:
    if settings.profile is not Profile.PRODUCTION:
        raise ConfigError("正式数据启用服务仅允许在 APP_PROFILE=production 中执行。")


def list_t16_templates() -> list[dict[str, Any]]:
    return [
        {
            "kind": kind,
            "title": T16_TITLES[kind],
            "filename": T16_FILENAMES[kind],
            "schema_version": T16_TEMPLATE_VERSION,
            "order": index,
            "required": True,
        }
        for index, kind in enumerate(T16_KIND_ORDER, start=1)
    ]


def build_t16_template_xlsx(kind: str) -> bytes:
    _ensure_t16_kind(kind)
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "数据"
    headers = T16_HEADERS[kind]
    data_sheet.append([T16_LABELS[field] for field in headers])
    data_sheet.freeze_panes = "A2"
    data_sheet.auto_filter.ref = data_sheet.dimensions
    for cell in data_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="335CFF")
        cell.alignment = Alignment(horizontal="center")
    for column in data_sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 4, 36)
        data_sheet.column_dimensions[column[0].column_letter].width = max(width, 14)
    for column_index, field in enumerate(headers, start=1):
        if field in T16_TEXT_FIELDS[kind]:
            data_sheet.column_dimensions[get_column_letter(column_index)].number_format = "@"

    # Production templates must be safe when a non-technical operator uploads
    # the freshly downloaded file without editing it.  Keep the import surface
    # empty and move the valid-looking sample away from the parsed sheet.
    example_sheet = workbook.create_sheet("示例（请勿导入）")
    example_sheet.append([T16_LABELS[field] for field in headers])
    example_sheet.append(T16_EXAMPLES[kind])
    example_sheet.freeze_panes = "A2"
    for cell in example_sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="64748B")
        cell.alignment = Alignment(horizontal="center")
    for column in example_sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 4, 36)
        example_sheet.column_dimensions[column[0].column_letter].width = max(width, 14)
    for column_index, field in enumerate(headers, start=1):
        if field in T16_TEXT_FIELDS[kind]:
            column_letter = get_column_letter(column_index)
            example_sheet.column_dimensions[column_letter].number_format = "@"
            example_sheet.cell(row=2, column=column_index).number_format = "@"

    guidance = workbook.create_sheet("填写说明")
    guidance.append(["填写顺序", "数据类型", "关键要求"])
    guidance_rows = {
        "colleges": "请在空白“数据”页填写；学院代码在校内保持唯一；示例页仅供参考。",
        "majors": "请在空白“数据”页填写；学院须已存在；示例页仅供参考。",
        "classes": "请在空白“数据”页填写；学院、专业、层次和届别须一致；示例页仅供参考。",
        "students": "请在空白“数据”页填写；班级须属于同一学院、专业、层次和届别；示例页仅供参考。",
        "outcomes": "请在空白“数据”页填写；学生须已存在，日期使用 YYYY-MM-DD；示例页仅供参考。",
    }
    text_guidance = (
        " 模板中除毕业届别、登记日期外的字段均已设为文本，粘贴时请保持文本格式；"
        "如 Excel 已自动转成数字、日期或布尔值，请从权威来源重新导出，"
        "不要根据当前显示值猜测、补零或改写。"
    )
    guidance.append(
        [T16_KIND_ORDER.index(kind) + 1, T16_TITLES[kind], guidance_rows[kind] + text_guidance]
    )
    for cell in guidance[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="335CFF")
    guidance.column_dimensions["A"].width = 12
    guidance.column_dimensions["B"].width = 18
    guidance.column_dimensions["C"].width = 72

    metadata = workbook.create_sheet("__meta__")
    metadata.append(["template_kind", kind])
    metadata.append(["schema_version", T16_TEMPLATE_VERSION])
    metadata.sheet_state = "hidden"
    workbook.active = 0
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def create_import_run(
    db: Session,
    *,
    cohort_year: int,
    business_note: str | None,
    actor: str,
    actor_id: int | None,
    actor_role: str,
    source_channel: str,
    idempotency_key: str,
) -> ImportRun:
    _require_production_import_profile()
    if cohort_year < 2000 or cohort_year > 2100:
        raise ValueError("毕业届别必须是 2000 到 2100 之间的四位年份")
    note = _bounded_text(business_note, 500, field="业务说明", required=False)
    receipt, replayed = claim_operation(
        db,
        namespace="import.run.create",
        idempotency_key=idempotency_key,
        actor=actor,
        payload={
            "cohort_year": cohort_year,
            "business_note": note,
        },
    )
    if replayed:
        run_id = receipt.result_json.get("run_id")
        run = db.get(ImportRun, run_id) if isinstance(run_id, int) else None
        if run is None:
            raise ValueError("idempotent import receipt has no committed run")
        return run
    run = ImportRun(
        cohort_year=cohort_year,
        business_note=note,
        status="draft",
        template_version=T16_TEMPLATE_VERSION,
        revision=1,
        result_summary=_empty_summary(),
        created_by=actor,
    )
    db.add(run)
    db.flush()
    finish_operation(receipt, result={"run_id": run.id})
    _record_run_audit(
        db,
        run,
        actor=actor,
        actor_id=actor_id,
        actor_role=actor_role,
        source_channel=source_channel,
        action="import.run.create",
        change_summary={
            "cohort_year": cohort_year,
            "reason": note or "创建首批正式数据启用会话",
            "before_status": None,
            "after_status": run.status,
            "before_revision": 0,
            "after_revision": run.revision,
        },
    )
    db.commit()
    db.refresh(run)
    return run


def list_import_runs(db: Session) -> list[ImportRun]:
    statement = (
        select(ImportRun)
        .order_by(ImportRun.id.desc())
        .limit(T16_RUN_LIST_LIMIT)
    )
    return list(db.execute(statement).scalars())


def get_import_run(db: Session, run_id: int, *, lock: bool = False) -> ImportRun:
    statement = select(ImportRun).where(ImportRun.id == run_id)
    if lock:
        statement = statement.with_for_update()
    run = db.execute(statement).scalar_one_or_none()
    if run is None:
        raise ValueError("启用记录不存在")
    return run


def _run_correlation_id(run: ImportRun) -> str:
    return f"QH-{run.cohort_year}-{run.id:04d}"


def _record_run_audit(
    db: Session,
    run: ImportRun,
    *,
    actor: str,
    actor_id: int | None,
    actor_role: str,
    source_channel: str,
    action: str,
    change_summary: dict[str, Any],
) -> AuditLog:
    """Persist a stable actor and trace envelope around one T16 event."""

    return record_db(
        db,
        actor=actor,
        actor_id=actor_id,
        actor_role=actor_role,
        source_channel=source_channel,
        correlation_id=_run_correlation_id(run),
        action=action,
        object_type="import_run",
        object_id=run.id,
        change_summary=change_summary,
    )


def active_batches(
    db: Session,
    run_id: int,
    *,
    summary_only: bool = False,
) -> list[ImportBatch]:
    statement = (
        select(ImportBatch)
        .where(ImportBatch.run_id == run_id, ImportBatch.status != "superseded")
        .order_by(ImportBatch.id)
    )
    if summary_only:
        statement = statement.options(
            defer(ImportBatch.errors),
            defer(ImportBatch.source_rows),
            defer(ImportBatch.staged_rows),
            defer(ImportBatch.inserted_refs),
            defer(ImportBatch.mapping_snapshot),
        )
    rows = list(db.execute(statement).scalars())
    latest: dict[str, ImportBatch] = {}
    for row in rows:
        latest[row.target_kind] = row
    return sorted(latest.values(), key=lambda row: T16_KIND_ORDER.index(row.target_kind))


def list_import_run_issues(
    db: Session,
    run_id: int,
    *,
    kind: str | None = None,
    code: str | None = None,
    page: int = 1,
    page_size: int = 20,
) -> dict[str, Any]:
    if kind is not None:
        _ensure_t16_kind(kind)
    page = max(int(page), 1)
    page_size = min(max(int(page_size), 1), 100)
    filters = ["ib.run_id = :run_id", "ib.status <> 'superseded'"]
    params: dict[str, Any] = {"run_id": run_id}
    if kind is not None:
        filters.append("ib.target_kind = :kind")
        params["kind"] = kind
    if code is not None:
        filters.append("issue.value->>'code' = :code")
        params["code"] = code[:64]
    where_clause = " AND ".join(filters)
    base = f"""
        FROM import_batches AS ib
        CROSS JOIN LATERAL jsonb_array_elements(ib.errors)
            WITH ORDINALITY AS issue(value, ordinality)
        WHERE {where_clause}
    """
    total = int(db.scalar(text(f"SELECT count(*) {base}"), params) or 0)
    rows = (
        db.execute(
            text(
                f"""
            SELECT ib.id AS batch_id, ib.target_kind AS kind, issue.value AS payload
            {base}
            ORDER BY ib.id, issue.ordinality
            LIMIT :limit OFFSET :offset
            """
            ),
            {**params, "limit": page_size, "offset": (page - 1) * page_size},
        )
        .mappings()
        .all()
    )
    code_filters = ["ib.run_id = :run_id", "ib.status <> 'superseded'"]
    code_params: dict[str, Any] = {"run_id": run_id}
    if kind is not None:
        code_filters.append("ib.target_kind = :kind")
        code_params["kind"] = kind
    codes = [
        str(value)
        for value in db.scalars(
            text(
                f"""
                SELECT DISTINCT issue.value->>'code'
                FROM import_batches AS ib
                CROSS JOIN LATERAL jsonb_array_elements(ib.errors) AS issue(value)
                WHERE {" AND ".join(code_filters)}
                ORDER BY issue.value->>'code'
                """
            ),
            code_params,
        )
        if value
    ]
    items = [
        {
            **dict(row["payload"]),
            "batch_id": int(row["batch_id"]),
            "kind": str(row["kind"]),
        }
        for row in rows
    ]
    return {
        "items": items,
        "total": total,
        "page": page,
        "page_size": page_size,
        "codes": codes,
    }


def list_import_run_batch_history(
    db: Session,
    run_id: int,
    *,
    page: int = 1,
    page_size: int = 20,
) -> dict[str, Any]:
    if db.scalar(select(ImportRun.id).where(ImportRun.id == run_id)) is None:
        raise ValueError("启用记录不存在")
    page = max(int(page), 1)
    page_size = min(max(int(page_size), 1), 100)
    total = int(
        db.scalar(select(func.count()).select_from(ImportBatch).where(ImportBatch.run_id == run_id))
        or 0
    )
    rows = list(
        db.execute(
            select(ImportBatch)
            .options(
                defer(ImportBatch.errors),
                defer(ImportBatch.source_rows),
                defer(ImportBatch.staged_rows),
                defer(ImportBatch.inserted_refs),
                defer(ImportBatch.mapping_snapshot),
                defer(ImportBatch.result_summary),
            )
            .where(ImportBatch.run_id == run_id)
            .order_by(ImportBatch.id.desc())
            .limit(page_size)
            .offset((page - 1) * page_size)
        ).scalars()
    )
    return {
        "items": [
            {
                "batch_id": batch.id,
                "target_kind": batch.target_kind,
                "title": T16_TITLES.get(batch.target_kind, batch.target_kind),
                "status": batch.status,
                "filename": batch.original_filename,
                "template_version": batch.template_version,
                "file_sha256": batch.source_sha256,
                "row_count": batch.row_count,
                "error_count": batch.error_count,
                "correction_count": batch.correction_count,
                "supersedes_batch_id": batch.supersedes_batch_id,
                "created_by": batch.created_by,
                "created_at": batch.created_at,
                "confirmed_by": batch.confirmed_by,
                "confirmed_at": batch.confirmed_at,
                "rolled_back_by": batch.rolled_back_by,
                "rolled_back_at": batch.rolled_back_at,
            }
            for batch in rows
        ],
        "total": total,
        "page": page,
        "page_size": page_size,
    }


def list_import_run_evidence(
    db: Session,
    run_id: int,
    *,
    page: int = 1,
    page_size: int = 20,
) -> dict[str, Any]:
    """Return one bounded, newest-first page of audit evidence for a run."""

    if db.scalar(select(ImportRun.id).where(ImportRun.id == run_id)) is None:
        raise ValueError("启用记录不存在")
    page = max(int(page), 1)
    page_size = min(max(int(page_size), 1), 100)
    filters = (
        AuditLog.object_type == "import_run",
        AuditLog.object_id == str(run_id),
    )
    total = int(db.scalar(select(func.count()).select_from(AuditLog).where(*filters)) or 0)
    rows = db.execute(
        select(
            AuditLog.id,
            AuditLog.action,
            AuditLog.actor_id,
            AuditLog.actor_name,
            AuditLog.actor_role,
            AuditLog.source_channel,
            AuditLog.correlation_id,
            AuditLog.at,
            AuditLog.change_summary,
        )
        .where(*filters)
        .order_by(AuditLog.id.desc())
        .limit(page_size)
        .offset((page - 1) * page_size)
    ).all()
    return {
        "items": [
            {
                "id": row.id,
                "event": row.action,
                "label": IMPORT_EVIDENCE_LABELS.get(row.action, row.action),
                "actor_id": row.actor_id,
                "actor": row.actor_name,
                "actor_role": row.actor_role,
                "source_channel": row.source_channel,
                "correlation_id": row.correlation_id,
                "batch_ids": _evidence_batch_ids(row.change_summary),
                "at": row.at,
                "detail": _evidence_detail(row.change_summary),
            }
            for row in rows
        ],
        "total": total,
        "page": page,
        "page_size": page_size,
    }


def import_run_payload(
    db: Session,
    run: ImportRun,
    *,
    lock: bool = True,
    include_evidence: bool = True,
) -> dict[str, Any]:
    # Every mutation also locks the run first. Holding the same row lock while
    # assembling batches, evidence and business counts prevents a torn detail
    # response across multiple READ COMMITTED statements.
    run = get_import_run(db, run.id, lock=lock)
    batches = active_batches(db, run.id, summary_only=True)
    by_kind = {batch.target_kind: batch for batch in batches}
    inline_issues = list_import_run_issues(
        db,
        run.id,
        page=1,
        page_size=INLINE_IMPORT_ISSUE_LIMIT,
    )["items"]
    errors_by_batch: dict[int, list[dict[str, Any]]] = {}
    for item in inline_issues:
        errors_by_batch.setdefault(int(item["batch_id"]), []).append(item)
    batch_payloads: list[dict[str, Any]] = []
    for batch in batches:
        summary = batch.result_summary or {}
        batch_errors = errors_by_batch.get(batch.id, [])
        batch_payloads.append(
            {
                "batch_id": batch.id,
                "target_kind": batch.target_kind,
                "title": T16_TITLES.get(batch.target_kind, batch.target_kind),
                "status": batch.status,
                "valid": batch.error_count == 0,
                "row_count": batch.row_count,
                "error_count": batch.error_count,
                "error_row_count": int(summary.get("errors", 0)),
                "errors": batch_errors,
                "filename": batch.original_filename,
                "template_version": batch.template_version,
                "file_sha256": batch.source_sha256,
                "inserted_count": int(summary.get("inserted", 0)),
                "updated_count": int(summary.get("updated", 0)),
                "unchanged_count": int(summary.get("unchanged", 0)),
                "correction_count": batch.correction_count,
            }
        )
    evidence_page = (
        list_import_run_evidence(
            db,
            run.id,
            page=1,
            page_size=INLINE_IMPORT_EVIDENCE_LIMIT,
        )
        if include_evidence
        else {"items": [], "total": 0}
    )
    return {
        "id": run.id,
        "status": run.status,
        "cohort_year": run.cohort_year,
        "business_note": run.business_note,
        "template_version": run.template_version,
        "revision": run.revision,
        "created_by": run.created_by,
        "created_at": run.created_at,
        "confirmed_by": run.confirmed_by,
        "confirmed_at": run.confirmed_at,
        "rolled_back_by": run.rolled_back_by,
        "rolled_back_at": run.rolled_back_at,
        "rollback_reason": run.rollback_reason,
        "missing_kinds": [kind for kind in T16_KIND_ORDER if kind not in by_kind],
        "batches": batch_payloads,
        "errors": inline_issues,
        "reconciliation": {
            **_normalized_summary(run.result_summary),
            "business_counts": _business_counts(db),
        },
        "evidence": evidence_page["items"],
        "evidence_total": evidence_page["total"],
    }


def parse_t16_workbook(
    kind: str, data: bytes
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, Any]]:
    _ensure_t16_kind(kind)
    try:
        workbook = load_workbook(BytesIO(data), read_only=False, data_only=False)
    except Exception as exc:
        raise ValueError("工作簿内容损坏或不是有效的 Excel 文件") from exc

    errors: list[dict[str, Any]] = []
    metadata: dict[str, Any] = {}
    if "__meta__" not in workbook.sheetnames:
        errors.append(
            _issue(
                1, "template_version_missing", "缺少模板版本信息", suggestion="请下载并使用当前模板"
            )
        )
    else:
        metadata_sheet = workbook["__meta__"]
        raw_metadata: dict[str, str] = {}
        for index in range(1, 10):
            key = str(metadata_sheet.cell(row=index, column=1).value or "").strip()
            if key not in {"template_kind", "schema_version"}:
                continue
            raw_metadata[key] = str(metadata_sheet.cell(row=index, column=2).value or "").strip()[
                :128
            ]
        kind_matches = raw_metadata.get("template_kind") == kind
        version_matches = raw_metadata.get("schema_version") == T16_TEMPLATE_VERSION
        # Persist only canonical allowlisted values or non-sensitive status
        # markers.  Operator-controlled metadata must never become an echo
        # channel for student identifiers or contact details.
        metadata["template_kind"] = kind if kind_matches else "__mismatch__"
        metadata["schema_version"] = T16_TEMPLATE_VERSION if version_matches else "__unsupported__"
        if not kind_matches:
            errors.append(
                _issue(
                    1,
                    "template_kind_mismatch",
                    "模板种类与当前资料不一致",
                    suggestion=f"请下载“{T16_TITLES[kind]}”模板",
                )
            )
        if not version_matches:
            errors.append(
                _issue(
                    1,
                    "template_version_unsupported",
                    "模板版本不受支持",
                    suggestion="请下载最新版模板后重新上传",
                )
            )

    if "数据" not in workbook.sheetnames:
        errors.append(
            _issue(
                1, "data_sheet_missing", "缺少“数据”工作表", suggestion="请勿删除或重命名模板工作表"
            )
        )
        return [], errors, metadata
    data_sheet = workbook["数据"]
    expected = T16_HEADERS[kind]
    if data_sheet.sheet_state != "visible":
        errors.append(
            _issue(
                1,
                "data_sheet_hidden",
                "“数据”工作表必须保持可见，系统不会导入操作者无法审阅的隐藏资料",
                suggestion="取消隐藏“数据”工作表后重新上传",
                editable=False,
            )
        )
        return [], errors, metadata
    if data_sheet.max_row - 1 > MAX_IMPORT_ROWS:
        errors.append(
            _issue(
                1,
                "row_limit_exceeded",
                f"单个模板最多允许 {MAX_IMPORT_ROWS} 行",
                suggestion="请拆分资料并联系管理员确认启用方案",
            )
        )
        return [], errors, metadata

    max_allowed_columns = len(expected) + MAX_IMPORT_EXTRA_COLUMNS
    if data_sheet.max_column > max_allowed_columns:
        errors.append(
            _issue(
                1,
                "column_limit_exceeded",
                f"模板最多允许 {max_allowed_columns} 列（含少量待识别的附加列）",
                suggestion="请删除模板有效区域之外的内容或重新下载当前模板",
            )
        )
        return [], errors, metadata

    for row_index, dimension in data_sheet.row_dimensions.items():
        if row_index > data_sheet.max_row or not (dimension.hidden or dimension.height == 0):
            continue
        cells = data_sheet[row_index][: data_sheet.max_column]
        if any(cell.value not in (None, "") for cell in cells):
            errors.append(
                _issue(
                    row_index,
                    "hidden_data_row",
                    f"第 {row_index} 行包含隐藏资料，无法确认操作者已经审阅",
                    suggestion="取消隐藏该行或删除其中内容后重新上传",
                    editable=False,
                )
            )

    hidden_columns: set[int] = set()
    for dimension in data_sheet.column_dimensions.values():
        if not (dimension.hidden or dimension.width == 0):
            continue
        start = int(dimension.min or 1)
        end = int(dimension.max or start)
        hidden_columns.update(range(start, min(end, data_sheet.max_column) + 1))
    for column_index in sorted(hidden_columns):
        first_value_row = next(
            (
                row_index
                for row_index in range(1, data_sheet.max_row + 1)
                if data_sheet.cell(row=row_index, column=column_index).value not in (None, "")
            ),
            None,
        )
        if first_value_row is None:
            continue
        errors.append(
            _issue(
                first_value_row,
                "hidden_data_column",
                f"第 {get_column_letter(column_index)} 列包含隐藏资料，无法确认操作者已经审阅",
                column=f"第{column_index}列",
                suggestion="取消隐藏该列或删除其中内容后重新上传",
                editable=False,
            )
        )

    if errors:
        return [], errors, metadata

    header_cells = next(data_sheet.iter_rows(min_row=1, max_row=1, max_col=data_sheet.max_column))
    raw_headers = [str(cell.value or "").strip()[:256] for cell in header_cells]
    label_to_field = {label: field for field, label in T16_LABELS.items()}
    normalized_headers = [label_to_field.get(header, header) for header in raw_headers]
    duplicates = {
        header for header, count in Counter(normalized_headers).items() if header and count > 1
    }
    for header in sorted(duplicates):
        positions = [
            index
            for index, candidate in enumerate(normalized_headers, start=1)
            if candidate == header
        ]
        is_expected = header in expected
        safe_column = header if is_expected else f"第{positions[0]}列"
        reason = (
            f"列“{_label(header)}”重复"
            if is_expected
            else f"第{'、'.join(str(index) for index in positions)}列表头重复且无法识别"
        )
        errors.append(
            _issue(
                1,
                "duplicate_header",
                reason,
                column=safe_column,
                suggestion="每个字段只保留一列",
            )
        )
    for header in expected:
        if header not in normalized_headers:
            errors.append(
                _issue(
                    1,
                    "missing_header",
                    f"缺少必填列“{_label(header)}”",
                    column=header,
                    suggestion="请使用当前模板并保留全部表头",
                )
            )
    for column_index, header in enumerate(normalized_headers, start=1):
        if header and header not in expected:
            errors.append(
                _issue(
                    1,
                    "unknown_header",
                    f"第{column_index}列表头无法识别",
                    column=f"第{column_index}列",
                    suggestion="删除额外列或改用当前模板中的标准列名",
                )
            )
    if errors:
        safe_headers = [
            (
                header
                if header in expected
                else f"__unknown_column_{column_index}__"
                if header
                else ""
            )
            for column_index, header in enumerate(normalized_headers, start=1)
        ]
        return [], errors, {**metadata, "headers": safe_headers}

    indexes = {header: normalized_headers.index(header) for header in expected}
    rows: list[dict[str, Any]] = []
    for row_number, cells in enumerate(
        data_sheet.iter_rows(
            min_row=2,
            max_row=data_sheet.max_row,
            max_col=data_sheet.max_column,
        ),
        start=2,
    ):
        parsed: dict[str, Any] = {"_row_number": row_number}
        for header in expected:
            cell = cells[indexes[header]] if indexes[header] < len(cells) else None
            if cell is not None and cell.data_type == "f":
                errors.append(
                    _issue(
                        row_number,
                        "formula_not_allowed",
                        "数据单元格不允许使用公式",
                        column=header,
                        suggestion="请粘贴为值后重新检查",
                    )
                )
                parsed[header] = None
            elif cell is not None and cell.data_type == "e":
                errors.append(
                    _issue(
                        row_number,
                        "excel_error_cell",
                        "Excel 单元格包含错误值，系统拒绝将错误标记当作正式数据",
                        column=header,
                        suggestion="请回到权威来源修复计算或导出错误后重新上传",
                        editable=False,
                    )
                )
                # Never persist an Excel error token such as ``#N/A`` as a
                # business string.  The authoritative source must be fixed.
                parsed[header] = None
            elif (
                cell is not None
                and header in T16_TEXT_FIELDS[kind]
                and cell.value is not None
                and not isinstance(cell.value, str)
            ):
                label = _label(header)
                errors.append(
                    _issue(
                        row_number,
                        "numeric_text_unsafe",
                        f"{label}在 Excel 中未以文本保存，原始文本语义、前导零或长数字精度可能已经丢失，系统拒绝猜测原值",
                        column=header,
                        suggestion=(
                            f"请回到权威来源，将“{label}”整列设为文本后重新粘贴或导出并重新上传；"
                            "不要根据当前显示值手工补零"
                        ),
                        editable=False,
                    )
                )
                # Never persist the ambiguous numeric value in source rows.
                # It may be PII and is not trustworthy enough for correction.
                parsed[header] = None
            else:
                parsed[header] = _clean_cell(cell.value if cell is not None else None)
        if any(parsed.get(header) not in (None, "") for header in expected):
            rows.append(parsed)
    if not rows and not errors:
        errors.append(
            _issue(2, "empty_file", "模板中没有可导入的数据", suggestion="请至少填写一行资料")
        )
    return rows, errors, {**metadata, "headers": normalized_headers}


def upload_run_batch(
    db: Session,
    *,
    run_id: int,
    kind: str,
    data: bytes,
    filename: str | None,
    revision: int,
    actor: str,
    actor_id: int | None,
    actor_role: str,
    source_channel: str,
) -> ImportRun:
    _require_production_import_profile()
    del filename  # T16 deliberately identifies uploads by digest, never by client filename.
    _ensure_t16_kind(kind)
    run = get_import_run(db, run_id, lock=True)
    before_status = run.status
    before_revision = run.revision
    if run.status not in {"draft", "needs_correction", "ready"}:
        raise ImportConflictError("当前启用记录已经确认或撤销，不能继续上传")
    digest = sha256(data).hexdigest()
    existing = db.execute(
        select(ImportBatch).where(
            ImportBatch.run_id == run.id,
            ImportBatch.target_kind == kind,
            ImportBatch.source_sha256 == digest,
        )
    ).scalar_one_or_none()
    if existing is not None:
        if existing.status == "superseded":
            raise ImportConflictError(
                "该文件曾被更新版本替代，请继续使用当前版本；如需恢复，请重新导出一份新文件后上传"
            )
        _record_run_audit(
            db,
            run,
            actor=actor,
            actor_id=actor_id,
            actor_role=actor_role,
            source_channel=source_channel,
            action="import.run.upload_idempotent",
            change_summary={
                "reason": "文件摘要与当前批次一致，复用既有检查结果",
                "batch_id": existing.id,
                "kind": kind,
                "source_sha256": digest,
                "before_status": before_status,
                "after_status": run.status,
                "before_revision": before_revision,
                "after_revision": run.revision,
            },
        )
        db.commit()
        db.refresh(run)
        return run
    if revision != run.revision:
        raise ImportConflictError("启用记录已被其他操作更新，请刷新后重新核对")

    rows, file_errors, mapping = parse_t16_workbook(kind, data)
    previous = (
        db.execute(
            select(ImportBatch)
            .where(
                ImportBatch.run_id == run.id,
                ImportBatch.target_kind == kind,
                ImportBatch.status != "superseded",
            )
            .order_by(ImportBatch.id.desc())
        )
        .scalars()
        .first()
    )
    if previous is not None:
        previous.status = "superseded"
        # Keep only irreversible lineage and aggregate evidence for replaced
        # workbooks; normalized rows may contain student PII.
        previous.source_rows = []
        previous.staged_rows = []
        previous.errors = []
        previous.original_filename = None
        # Header names and file-level errors are operator-controlled input too.
        # They may contain phone numbers or other PII, so a superseded batch
        # retains only the dedicated lineage columns and aggregate counters.
        previous.mapping_snapshot = {}
    batch = ImportBatch(
        run_id=run.id,
        target_kind=kind,
        # The client filename is untrusted metadata and may itself contain
        # names, phone numbers or student identifiers.  File identity is the
        # content digest; expose only a deterministic privacy-safe label.
        original_filename=f"{kind}-{digest[:12]}.xlsx",
        status="needs_correction" if file_errors else "uploaded",
        row_count=len(rows),
        error_count=len(file_errors),
        errors=file_errors,
        source_rows=_protect_source_rows(kind, rows),
        staged_rows=[],
        inserted_refs={},
        template_version=(
            T16_TEMPLATE_VERSION if mapping.get("schema_version") == T16_TEMPLATE_VERSION else None
        ),
        source_sha256=digest,
        mapping_snapshot={"file_errors": file_errors, **mapping},
        correction_count=0,
        result_summary=_empty_summary(),
        supersedes_batch_id=previous.id if previous is not None else None,
        created_by=actor,
    )
    db.add(batch)
    db.flush()
    _record_run_audit(
        db,
        run,
        actor=actor,
        actor_id=actor_id,
        actor_role=actor_role,
        source_channel=source_channel,
        action="import.run.upload",
        change_summary={
            "reason": "上传正式模板并自动执行跨表预检",
            "batch_id": batch.id,
            "kind": kind,
            "template_version": batch.template_version,
            "source_sha256": digest,
            "source_rows": len(rows),
            "file_errors": len(file_errors),
            "supersedes_batch_id": batch.supersedes_batch_id,
            "before_status": before_status,
            "after_status": run.status,
            "before_revision": before_revision,
            "after_revision": run.revision,
        },
    )
    return preflight_import_run(
        db,
        run.id,
        actor=actor,
        actor_id=actor_id,
        actor_role=actor_role,
        source_channel=source_channel,
        trigger="upload",
        commit=True,
        bump_revision=True,
    )


def correct_run_row(
    db: Session,
    *,
    run_id: int,
    batch_id: int,
    row_number: int,
    column: str,
    value: str | None,
    reason: str,
    revision: int,
    actor: str,
    actor_id: int | None,
    actor_role: str,
    source_channel: str,
) -> ImportRun:
    _require_production_import_profile()
    run = get_import_run(db, run_id, lock=True)
    before_status = run.status
    before_revision = run.revision
    if run.status not in {"draft", "needs_correction", "ready"}:
        raise ImportConflictError("当前启用记录已经确认或撤销，不能继续修正")
    if revision != run.revision:
        raise ImportConflictError("启用记录已被其他操作更新，请刷新后重新核对")
    batch = db.execute(
        select(ImportBatch).where(
            ImportBatch.id == batch_id,
            ImportBatch.run_id == run.id,
            ImportBatch.status != "superseded",
        )
    ).scalar_one_or_none()
    if batch is None:
        raise ValueError("导入批次不存在")
    field = {label: name for name, label in T16_LABELS.items()}.get(column, column)
    if field not in T16_HEADERS.get(batch.target_kind, []):
        raise ValueError("该列不能在线修正，请使用当前模板重新上传")
    if value is not None and not isinstance(value, str):
        raise ValueError("在线修正只接受单个文本或空值")
    issue_label = T16_LABELS[field]
    matching_issues = [
        issue
        for issue in (batch.errors or [])
        if int(issue.get("row_number") or -1) == row_number
        and issue.get("column") == issue_label
    ]
    if not matching_issues:
        raise ValueError("该单元格没有待修正问题，请刷新后重新核对")
    if any(issue.get("editable") is not True for issue in matching_issues):
        raise ValueError("该问题不允许在线修正，请从权威来源更正后重新上传")
    correction_reason = _bounded_text(reason, 500, field="修正原因", required=True)
    rows = _unprotect_source_rows(batch.target_kind, list(batch.source_rows or []))
    row = next((item for item in rows if int(item.get("_row_number", -1)) == row_number), None)
    if row is None:
        raise ValueError("需要修正的原始行不存在，请重新上传文件")
    before = row.get(field)
    row[field] = _clean_cell(value)
    batch.source_rows = _protect_source_rows(batch.target_kind, rows)
    batch.correction_count += 1
    _record_run_audit(
        db,
        run,
        actor=actor,
        actor_id=actor_id,
        actor_role=actor_role,
        source_channel=source_channel,
        action="import.run.correct",
        change_summary={
            "batch_id": batch.id,
            "kind": batch.target_kind,
            "row_number": row_number,
            "column": field,
            "before": _mask_value(field, before),
            "after": _mask_value(field, row.get(field)),
            "reason": correction_reason,
            "before_status": before_status,
            "after_status": run.status,
            "before_revision": before_revision,
            "after_revision": run.revision,
        },
    )
    return preflight_import_run(
        db,
        run.id,
        actor=actor,
        actor_id=actor_id,
        actor_role=actor_role,
        source_channel=source_channel,
        trigger="correction",
        commit=True,
        bump_revision=True,
    )


def preflight_import_run(
    db: Session,
    run_id: int,
    *,
    actor: str,
    actor_id: int | None,
    actor_role: str,
    source_channel: str,
    trigger: str,
    commit: bool,
    bump_revision: bool,
    deduplicate_unchanged: bool = False,
) -> ImportRun:
    _require_production_import_profile()
    run = get_import_run(db, run_id, lock=True)
    if run.status in {"imported", "rolled_back"}:
        _record_run_audit(
            db,
            run,
            actor=actor,
            actor_id=actor_id,
            actor_role=actor_role,
            source_channel=source_channel,
            action="import.run.preflight",
            change_summary={
                "reason": "显式重新校验；会话已处于最终状态，未改变业务事实",
                "trigger": trigger,
                "validation_status": run.status,
                "validation_noop": True,
                "validation_plan_sha256": _json_fingerprint({}),
                "missing_kinds": [],
                "before_status": run.status,
                "after_status": run.status,
                "before_revision": run.revision,
                "after_revision": run.revision,
                **_normalized_summary(run.result_summary),
            },
        )
        if commit:
            db.commit()
            db.refresh(run)
        else:
            db.flush()
        return run
    before_status = run.status
    before_revision = run.revision
    batches = active_batches(db, run.id)
    before_signature = _preflight_state_signature(run, batches)
    staged_before = {batch.id: batch.staged_rows for batch in batches}
    by_kind = {batch.target_kind: batch for batch in batches}
    _validate_run_batches(db, run, by_kind)
    missing = [kind for kind in T16_KIND_ORDER if kind not in by_kind]
    issue_count = sum(batch.error_count for batch in batches)
    if missing:
        run.status = "needs_correction" if issue_count else "draft"
    else:
        run.status = "ready" if issue_count == 0 else "needs_correction"
    run.result_summary = _aggregate_summaries(batches)
    validation_plan_sha256 = _json_fingerprint(
        {
            batch.target_kind: _staged_plan_signature(
                batch.target_kind,
                batch.staged_rows or [],
            )
            for batch in batches
        }
    )
    if deduplicate_unchanged and before_signature == _preflight_state_signature(run, batches):
        # Validation re-encrypts staged sensitive fields.  Restore the original
        # ciphertext when the operator-visible result is identical so a no-op
        # preflight is also a true database no-op.
        for batch in batches:
            batch.staged_rows = staged_before[batch.id]
        _record_run_audit(
            db,
            run,
            actor=actor,
            actor_id=actor_id,
            actor_role=actor_role,
            source_channel=source_channel,
            action="import.run.preflight",
            change_summary={
                "reason": "显式重新校验，业务状态与上一版一致",
                "trigger": trigger,
                "validation_status": run.status,
                "validation_noop": True,
                "validation_plan_sha256": validation_plan_sha256,
                "missing_kinds": missing,
                "before_status": before_status,
                "after_status": run.status,
                "before_revision": before_revision,
                "after_revision": run.revision,
                **_normalized_summary(run.result_summary),
            },
        )
        if commit:
            db.commit()
            db.refresh(run)
        else:
            db.flush()
        return run
    if bump_revision:
        run.revision += 1
    _record_run_audit(
        db,
        run,
        actor=actor,
        actor_id=actor_id,
        actor_role=actor_role,
        source_channel=source_channel,
        action="import.run.preflight",
        change_summary={
            "reason": "完成跨表完整性、一致性与目标影响校验",
            "trigger": trigger,
            "validation_status": run.status,
            "validation_noop": False,
            "validation_plan_sha256": validation_plan_sha256,
            "status": run.status,
            "missing_kinds": missing,
            "before_status": before_status,
            "after_status": run.status,
            "before_revision": before_revision,
            "after_revision": run.revision,
            **_normalized_summary(run.result_summary),
        },
    )
    if commit:
        db.commit()
        db.refresh(run)
    else:
        db.flush()
    return run


def _preflight_state_signature(
    run: ImportRun,
    batches: list[ImportBatch],
) -> dict[str, Any]:
    """Stable, PII-free signature for suppressing only true no-op preflights."""

    return {
        "status": run.status,
        "summary": _normalized_summary(run.result_summary),
        "batches": [
            {
                "id": batch.id,
                "kind": batch.target_kind,
                "status": batch.status,
                "error_count": batch.error_count,
                "errors": batch.errors or [],
                "summary": _normalized_summary(batch.result_summary),
                "plan": _staged_plan_signature(batch.target_kind, batch.staged_rows or []),
            }
            for batch in batches
        ],
    }


def _staged_plan_signature(
    kind: str,
    staged_rows: list[dict[str, Any]],
) -> list[dict[str, str | None]]:
    """Hash a staged plan without persisting natural keys or plaintext PII."""

    signature: list[dict[str, str | None]] = []
    for item in _unprotect_staged_rows(kind, list(staged_rows)):
        data = item.get("data")
        if not isinstance(data, dict):
            continue
        natural_key = _import_lock_key(kind, data)
        signature.append(
            {
                "key_sha256": sha256(natural_key.encode("utf-8")).hexdigest(),
                "action": str(item.get("action") or ""),
                "expected_before_sha256": (
                    str(item["expected_before_sha256"])
                    if item.get("expected_before_sha256") is not None
                    else None
                ),
                "desired_sha256": _dict_fingerprint(data),
            }
        )
    return sorted(signature, key=lambda item: str(item["key_sha256"]))


def _batch_change_evidence(
    batches: dict[str, ImportBatch],
    approved_plan: dict[str, list[dict[str, str | None]]] | None,
) -> list[dict[str, Any]]:
    """Return bounded, PII-free proof of the objects changed by each batch."""

    evidence: list[dict[str, Any]] = []
    for kind in T16_KIND_ORDER:
        batch = batches[kind]
        refs = dict(batch.inserted_refs or {})
        inserted_ids = sorted(int(value) for value in (refs.get("inserted_ids") or []))
        inserted_fingerprints = dict(refs.get("inserted_fingerprints") or {})
        updates = [dict(item) for item in (refs.get("updates") or [])]
        update_ids = sorted(int(item["id"]) for item in updates)
        before_proof = [
            {
                "id": int(item["id"]),
                "state_sha256": (
                    sha256(str(item["before_encrypted"]).encode("utf-8")).hexdigest()
                    if item.get("before_encrypted")
                    else _json_fingerprint(item.get("before") or {})
                ),
            }
            for item in updates
        ]
        after_proof = [
            {
                "id": int(item["id"]),
                "state_sha256": str(
                    item.get("after_fingerprint")
                    or _json_fingerprint(item.get("after") or {})
                ),
            }
            for item in updates
        ]
        evidence.append(
            {
                "batch_id": batch.id,
                "kind": kind,
                "source_sha256": batch.source_sha256,
                "plan_sha256": (
                    _json_fingerprint(approved_plan.get(kind, []))
                    if approved_plan is not None
                    else None
                ),
                "inserted_count": len(inserted_ids),
                "inserted_ids_sha256": _json_fingerprint(inserted_ids),
                "inserted_after_sha256": _json_fingerprint(inserted_fingerprints),
                "updated_count": len(update_ids),
                "updated_ids_sha256": _json_fingerprint(update_ids),
                "updated_before_sha256": _json_fingerprint(before_proof),
                "updated_after_sha256": _json_fingerprint(after_proof),
            }
        )
    return evidence


def _validation_lookup_scope(
    rows_by_kind: dict[str, list[dict[str, Any]]],
) -> _ImportLookupScope:
    """Extract only the database keys referenced by this run's source rows."""

    scope = _ImportLookupScope()
    for row in rows_by_kind.get("colleges", []):
        code = _text(row.get("college_code"))
        name = _text(row.get("college_name"))
        if code:
            scope.college_codes.add(code)
        if name:
            scope.college_names.add(name)

    for kind in ("majors", "classes", "students"):
        for row in rows_by_kind.get(kind, []):
            college_code = _text(row.get("college_code"))
            if college_code:
                scope.college_codes.add(college_code)

    for kind in ("majors", "classes", "students"):
        for row in rows_by_kind.get(kind, []):
            key = (
                _text(row.get("college_code")),
                _text(row.get("major_code")),
                _text(row.get("degree_level")),
            )
            if all(key):
                scope.major_keys.add(key)

    for kind in ("classes", "students"):
        for row in rows_by_kind.get(kind, []):
            cohort_year = _cohort(row.get("cohort_year"))
            key = (
                _text(row.get("college_code")),
                _text(row.get("major_code")),
                _text(row.get("degree_level")),
                cohort_year,
                _text(row.get("class_code")),
            )
            if cohort_year is not None and all(key):
                scope.class_keys.add((key[0], key[1], key[2], cohort_year, key[4]))

    for kind in ("students", "outcomes"):
        for row in rows_by_kind.get(kind, []):
            student_no = _text(row.get("student_no"))
            if student_no:
                scope.student_numbers.add(student_no)
                if kind == "students":
                    scope.student_scope_numbers.add(student_no)
    return scope


def _staged_lookup_scope(
    kind: str,
    staged: list[dict[str, Any]],
) -> _ImportLookupScope:
    """Extract normalized keys required while applying exactly one batch."""

    scope = _ImportLookupScope()
    for item in staged:
        data = item.get("data")
        if not isinstance(data, dict):
            continue
        if kind == "colleges":
            if data.get("code"):
                scope.college_codes.add(str(data["code"]))
            if data.get("name"):
                scope.college_names.add(str(data["name"]))
            continue
        if kind in {"majors", "classes", "students"} and data.get("college_code"):
            scope.college_codes.add(str(data["college_code"]))
        if kind in {"majors", "classes", "students"}:
            major_key = (
                str(data.get("college_code") or ""),
                str(data.get("code") if kind == "majors" else data.get("major_code") or ""),
                str(
                    data.get("level")
                    if kind in {"majors", "classes"}
                    else data.get("degree_level") or ""
                ),
            )
            if all(major_key):
                scope.major_keys.add(major_key)
        if kind in {"classes", "students"}:
            cohort_year = data.get("cohort_year")
            class_key = (
                str(data.get("college_code") or ""),
                str(data.get("major_code") or ""),
                str(data.get("level") if kind == "classes" else data.get("degree_level") or ""),
                cohort_year,
                str(data.get("code") if kind == "classes" else data.get("class_code") or ""),
            )
            if isinstance(cohort_year, int) and all(class_key):
                scope.class_keys.add(
                    (class_key[0], class_key[1], class_key[2], cohort_year, class_key[4])
                )
        if kind in {"students", "outcomes"} and data.get("student_no"):
            scope.student_numbers.add(str(data["student_no"]))
    return scope


def _postgres_bind_chunks(values: Any, *, tuple_width: int) -> list[list[Any]]:
    """Return globally sorted chunks below PostgreSQL's bind limit.

    SQLAlchemy expands ``IN`` and tuple-``IN`` values into one bind per scalar.
    The T16 ceiling applies independently to five files, so a combined lookup
    may contain more than 20k keys.  Chunk size therefore depends on tuple
    arity rather than only row count.
    """

    if tuple_width < 1:
        raise ValueError("tuple_width 必须大于 0")
    ordered = sorted(set(values))
    chunk_size = max(
        1,
        min(
            T16_QUERY_MAX_ITEMS_PER_STATEMENT,
            T16_QUERY_BIND_PARAMETER_BUDGET // tuple_width,
        ),
    )
    return [ordered[offset : offset + chunk_size] for offset in range(0, len(ordered), chunk_size)]


def _load_scoped_import_state(
    db: Session,
    scope: _ImportLookupScope,
) -> _ImportLookupState:
    """Load run-related rows without ever scanning an entire business table."""

    student_rows_by_id: dict[int, Student] = {}
    for chunk in _postgres_bind_chunks(scope.student_numbers, tuple_width=1):
        for row in db.execute(
            select(Student).where(Student.student_no.in_(chunk)).order_by(Student.id)
        ).scalars():
            student_rows_by_id[row.id] = row
    students = {row.student_no: row for row in student_rows_by_id.values()}

    scoped_existing_students = [
        students[student_no]
        for student_no in sorted(scope.student_scope_numbers)
        if student_no in students
    ]
    existing_college_ids = {row.college_id for row in scoped_existing_students}
    existing_major_ids = {row.major_id for row in scoped_existing_students}
    existing_class_ids = {
        row.class_id for row in scoped_existing_students if row.class_id is not None
    }

    college_rows_by_id: dict[int, College] = {}
    for column, values in (
        (College.code, scope.college_codes),
        (College.name, scope.college_names),
        (College.id, existing_college_ids),
    ):
        for chunk in _postgres_bind_chunks(values, tuple_width=1):
            for row in db.execute(
                select(College).where(column.in_(chunk)).order_by(College.id)
            ).scalars():
                college_rows_by_id[row.id] = row
    college_rows = list(college_rows_by_id.values())
    colleges = {row.code: row for row in college_rows if row.code}
    college_name_owners = {row.name: row.code for row in college_rows}
    college_ids = {row.id: row.code for row in colleges.values()}

    database_major_keys = {
        (college.id, code, level)
        for college_code, code, level in scope.major_keys
        if (college := colleges.get(college_code)) is not None
    }
    major_rows_by_id: dict[int, Major] = {}
    for chunk in _postgres_bind_chunks(database_major_keys, tuple_width=3):
        for row in db.execute(
            select(Major)
            .where(tuple_(Major.college_id, Major.code, Major.level).in_(chunk))
            .order_by(Major.id)
        ).scalars():
            major_rows_by_id[row.id] = row
    for chunk in _postgres_bind_chunks(existing_major_ids, tuple_width=1):
        for row in db.execute(
            select(Major).where(Major.id.in_(chunk)).order_by(Major.id)
        ).scalars():
            major_rows_by_id[row.id] = row
    major_rows = list(major_rows_by_id.values())
    majors = {
        (college_ids.get(row.college_id, "") or "", row.code, row.level): row
        for row in major_rows
        if college_ids.get(row.college_id)
    }
    major_ids = {row.id: key for key, row in majors.items()}

    database_class_keys = set()
    for college_code, major_code, level, cohort_year, class_code in scope.class_keys:
        college = colleges.get(college_code)
        major = majors.get((college_code, major_code, level))
        if college is not None and major is not None:
            database_class_keys.add((college.id, major.id, cohort_year, class_code))
    class_rows_by_id: dict[int, AcademicClass] = {}
    for chunk in _postgres_bind_chunks(database_class_keys, tuple_width=4):
        for row in db.execute(
            select(AcademicClass)
            .where(
                tuple_(
                    AcademicClass.college_id,
                    AcademicClass.major_id,
                    AcademicClass.cohort_year,
                    AcademicClass.code,
                ).in_(chunk)
            )
            .order_by(AcademicClass.id)
        ).scalars():
            class_rows_by_id[row.id] = row
    for chunk in _postgres_bind_chunks(existing_class_ids, tuple_width=1):
        for row in db.execute(
            select(AcademicClass).where(AcademicClass.id.in_(chunk)).order_by(AcademicClass.id)
        ).scalars():
            class_rows_by_id[row.id] = row
    class_rows = list(class_rows_by_id.values())
    classes = {
        (
            college_ids.get(row.college_id, "") or "",
            major_ids.get(row.major_id, ("", "", ""))[1],
            major_ids.get(row.major_id, ("", "", ""))[2],
            row.cohort_year,
            row.code,
        ): row
        for row in class_rows
        if college_ids.get(row.college_id) and row.major_id in major_ids
    }
    class_ids = {row.id: key for key, row in classes.items()}
    return _ImportLookupState(
        colleges=colleges,
        college_name_owners=college_name_owners,
        college_ids=college_ids,
        majors=majors,
        major_ids=major_ids,
        classes=classes,
        class_ids=class_ids,
        students=students,
    )


def _validate_run_batches(
    db: Session,
    run: ImportRun,
    batches: dict[str, ImportBatch],
) -> None:
    rows_by_kind = {
        kind: _unprotect_source_rows(kind, list(batch.source_rows or []))
        for kind, batch in batches.items()
    }
    state = _load_scoped_import_state(db, _validation_lookup_scope(rows_by_kind))
    colleges = state.colleges
    college_name_owners = state.college_name_owners
    college_ids = state.college_ids
    majors = state.majors
    major_ids = state.major_ids
    classes = state.classes
    class_ids = state.class_ids
    students = state.students

    virtual_colleges = set(colleges)
    virtual_majors = set(majors)
    virtual_classes = set(classes)
    virtual_students = set(students)
    uploaded_student_actions: dict[str, str] = {}

    for kind in T16_KIND_ORDER:
        batch = batches.get(kind)
        if batch is None:
            continue
        file_errors = list((batch.mapping_snapshot or {}).get("file_errors") or [])
        if file_errors:
            _set_batch_validation(batch, [], file_errors)
            continue
        rows = rows_by_kind[kind]
        if kind == "colleges":
            staged, errors = _validate_colleges(
                rows,
                colleges,
                college_name_owners,
                virtual_colleges,
            )
        elif kind == "majors":
            staged, errors = _validate_majors(rows, majors, virtual_colleges, virtual_majors)
        elif kind == "classes":
            staged, errors = _validate_classes(
                rows,
                classes,
                virtual_majors,
                virtual_classes,
                run.cohort_year,
            )
        elif kind == "students":
            staged, errors = _validate_students(
                rows,
                students,
                college_ids,
                major_ids,
                class_ids,
                virtual_classes,
                virtual_students,
                run.cohort_year,
            )
            uploaded_student_actions = {
                item["data"]["student_no"]: item["action"] for item in staged
            }
        else:
            staged, errors = _validate_outcomes(
                rows,
                students,
                virtual_students,
                uploaded_student_actions,
                run.cohort_year,
            )
        _set_batch_validation(batch, staged, errors)


def _set_batch_validation(
    batch: ImportBatch,
    staged: list[dict[str, Any]],
    errors: list[dict[str, Any]],
) -> None:
    inserted = sum(item.get("action") == "inserted" for item in staged)
    updated = sum(item.get("action") == "updated" for item in staged)
    unchanged = sum(item.get("action") == "unchanged" for item in staged)
    summary = {
        "source": batch.row_count,
        "inserted": inserted,
        "updated": updated,
        "unchanged": unchanged,
        # Structural workbook issues remain visible in ``batch.errors`` but are
        # not invented as source rows.  This keeps the four-way reconciliation
        # invariant exact for every batch, including an empty or malformed file.
        "errors": max(batch.row_count - inserted - updated - unchanged, 0),
    }
    batch.staged_rows = _protect_staged_rows(batch.target_kind, staged)
    batch.errors = errors
    batch.error_count = len(errors)
    batch.result_summary = summary
    batch.status = "ready" if not errors else "needs_correction"


def _target_before_fingerprint(
    kind: str,
    current: Any | None,
    data: dict[str, Any],
) -> str | None:
    """Fingerprint only target fields whose change can alter this import plan."""

    if current is None:
        return None
    fields: tuple[str, ...]
    if kind == "colleges":
        fields = ("code", "name")
    elif kind == "majors":
        fields = ("college_id", "code", "name", "level")
    elif kind == "classes":
        fields = ("college_id", "major_id", "cohort_year", "code", "name")
    elif kind == "students":
        student_fields = [
            "student_no",
            "name",
            "gender",
            "college_id",
            "major_id",
            "class_id",
            "degree_level",
            "cohort_year",
            "origin_place",
            "is_key_group",
        ]
        if data.get("contact") is not None:
            student_fields.append("contact_encrypted")
        fields = tuple(student_fields)
    elif kind == "outcomes":
        fields = ("student_no", "cohort_year", "status", "status_occurred_at")
    else:
        raise ValueError("未知资料类型")
    return _snapshot_fingerprint(current, fields)


def _staged_item(
    kind: str,
    *,
    current: Any | None,
    action: str,
    data: dict[str, Any],
    **extra: Any,
) -> dict[str, Any]:
    return {
        "action": action,
        "existing_id": current.id if current else None,
        "expected_before_sha256": _target_before_fingerprint(kind, current, data),
        "data": data,
        **extra,
    }


def _validate_colleges(
    rows: list[dict[str, Any]],
    existing: dict[str, College],
    existing_name_owners: dict[str, str | None],
    virtual: set[str],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    duplicates = Counter(str(row.get("college_code") or "") for row in rows)
    duplicate_names = Counter(_text(row.get("college_name")) for row in rows)
    staged: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for row in rows:
        start = len(errors)
        number = int(row["_row_number"])
        code = _text(row.get("college_code"))
        name = _text(row.get("college_name"))
        _require(errors, number, "college_code", code)
        _require(errors, number, "college_name", name)
        _identifier(errors, number, "college_code", code)
        _text_quality(errors, number, "college_name", name, 128)
        if code and duplicates[code] > 1:
            errors.append(
                _issue(
                    number,
                    "duplicate_in_file",
                    "同一学院代码在文件中出现多次",
                    column="college_code",
                    masked_value=code,
                    suggestion="每个学院只保留一行",
                )
            )
        if name and duplicate_names[name] > 1:
            errors.append(
                _issue(
                    number,
                    "duplicate_name_in_file",
                    "同一学院名称在文件中对应多行",
                    column="college_name",
                    masked_value=name,
                    suggestion="学院名称和学院代码都应保持唯一",
                )
            )
        if name in existing_name_owners and existing_name_owners[name] != code:
            errors.append(
                _issue(
                    number,
                    "natural_key_conflict",
                    "学院名称已被其他学院代码使用",
                    column="college_name",
                    masked_value=name,
                    suggestion="核对学院代码，避免覆盖已有学院",
                )
            )
        if len(errors) != start:
            continue
        current = existing.get(code)
        action = (
            "inserted" if current is None else ("unchanged" if current.name == name else "updated")
        )
        staged.append(
            _staged_item(
                "colleges",
                current=current,
                action=action,
                data={"code": code, "name": name},
            )
        )
        virtual.add(code)
    return staged, errors


def _validate_majors(
    rows: list[dict[str, Any]],
    existing: dict[tuple[str, str, str], Major],
    college_codes: set[str],
    virtual: set[tuple[str, str, str]],
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    keys = [
        (
            _text(row.get("college_code")),
            _text(row.get("major_code")),
            _text(row.get("degree_level")),
        )
        for row in rows
    ]
    duplicates = Counter(keys)
    staged: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for row, key in zip(rows, keys, strict=True):
        start = len(errors)
        number = int(row["_row_number"])
        college_code, code, level = key
        name = _text(row.get("major_name"))
        for field, value in (
            ("college_code", college_code),
            ("major_code", code),
            ("major_name", name),
            ("degree_level", level),
        ):
            _require(errors, number, field, value)
        _identifier(errors, number, "college_code", college_code)
        _identifier(errors, number, "major_code", code)
        _text_quality(errors, number, "major_name", name, 128)
        _enum(errors, number, "degree_level", level, _values(DegreeLevel))
        if college_code and college_code not in college_codes:
            errors.append(
                _issue(
                    number,
                    "parent_code_unknown",
                    "学院代码不存在",
                    column="college_code",
                    masked_value=college_code,
                    suggestion="先在本启用记录中上传学院资料",
                )
            )
        if all(key) and duplicates[key] > 1:
            errors.append(
                _issue(
                    number,
                    "duplicate_in_file",
                    "同一专业在文件中出现多次",
                    column="major_code",
                    masked_value=code,
                    suggestion="同一学院、专业代码和学历层次只保留一行",
                )
            )
        if len(errors) != start:
            continue
        current = existing.get(key)
        action = (
            "inserted" if current is None else ("unchanged" if current.name == name else "updated")
        )
        staged.append(
            _staged_item(
                "majors",
                current=current,
                action=action,
                data={"college_code": college_code, "code": code, "name": name, "level": level},
            )
        )
        virtual.add(key)
    return staged, errors


def _validate_classes(
    rows: list[dict[str, Any]],
    existing: dict[tuple[str, str, str, int, str], AcademicClass],
    major_keys: set[tuple[str, str, str]],
    virtual: set[tuple[str, str, str, int, str]],
    run_cohort: int,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    parsed_keys = [
        (
            _text(row.get("college_code")),
            _text(row.get("major_code")),
            _text(row.get("degree_level")),
            _cohort(row.get("cohort_year")),
            _text(row.get("class_code")),
        )
        for row in rows
    ]
    duplicates = Counter(parsed_keys)
    staged: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for row, key in zip(rows, parsed_keys, strict=True):
        start = len(errors)
        number = int(row["_row_number"])
        college_code, major_code, level, cohort_year, code = key
        name = _text(row.get("class_name"))
        for field, value in (
            ("college_code", college_code),
            ("major_code", major_code),
            ("degree_level", level),
            ("cohort_year", row.get("cohort_year")),
            ("class_code", code),
            ("class_name", name),
        ):
            _require(errors, number, field, value)
        _identifier(errors, number, "college_code", college_code)
        _identifier(errors, number, "major_code", major_code)
        _identifier(errors, number, "class_code", code)
        _text_quality(errors, number, "class_name", name, 128)
        _enum(errors, number, "degree_level", level, _values(DegreeLevel))
        if cohort_year is None:
            errors.append(
                _issue(
                    number,
                    "cohort_year_format",
                    "毕业届别必须是四位年份",
                    column="cohort_year",
                    masked_value=_text(row.get("cohort_year")),
                    suggestion="例如：2026",
                )
            )
        elif cohort_year != run_cohort:
            errors.append(
                _issue(
                    number,
                    "cohort_scope_mismatch",
                    "班级届别与本次启用届别不一致",
                    column="cohort_year",
                    masked_value=str(cohort_year),
                    suggestion=f"请填写 {run_cohort} 或另建启用记录",
                )
            )
        if (college_code, major_code, level) not in major_keys:
            errors.append(
                _issue(
                    number,
                    "parent_code_unknown",
                    "学院、专业代码和学历层次组合不存在",
                    column="major_code",
                    masked_value=major_code,
                    suggestion="先上传并修正专业资料",
                )
            )
        if all((college_code, major_code, level, cohort_year, code)) and duplicates[key] > 1:
            errors.append(
                _issue(
                    number,
                    "duplicate_in_file",
                    "同一班级在文件中出现多次",
                    column="class_code",
                    masked_value=code,
                    suggestion="每个班级只保留一行",
                )
            )
        if len(errors) != start or cohort_year is None:
            continue
        current = existing.get(key)
        action = (
            "inserted" if current is None else ("unchanged" if current.name == name else "updated")
        )
        staged.append(
            _staged_item(
                "classes",
                current=current,
                action=action,
                data={
                    "college_code": college_code,
                    "major_code": major_code,
                    "level": level,
                    "cohort_year": cohort_year,
                    "code": code,
                    "name": name,
                },
            )
        )
        virtual.add(key)
    return staged, errors


def _validate_students(
    rows: list[dict[str, Any]],
    existing: dict[str, Student],
    college_ids: dict[int, str | None],
    major_ids: dict[int, tuple[str, str, str]],
    class_ids: dict[int, tuple[str, str, str, int, str]],
    class_keys: set[tuple[str, str, str, int, str]],
    virtual_students: set[str],
    run_cohort: int,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    numbers = [_text(row.get("student_no")) for row in rows]
    duplicates = Counter(numbers)
    staged: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for row, student_no in zip(rows, numbers, strict=True):
        start = len(errors)
        number = int(row["_row_number"])
        desired = {
            "student_no": student_no,
            "name": _text(row.get("name")),
            "gender": _text(row.get("gender")),
            "college_code": _text(row.get("college_code")),
            "major_code": _text(row.get("major_code")),
            "degree_level": _text(row.get("degree_level")),
            "cohort_year": _cohort(row.get("cohort_year")),
            "class_code": _text(row.get("class_code")),
            "origin_place": _optional_text(row.get("origin_place")),
            "contact": _optional_text(row.get("contact")),
            "is_key_group": _optional_text(row.get("is_key_group")),
        }
        for field in (
            "student_no",
            "name",
            "gender",
            "college_code",
            "major_code",
            "degree_level",
            "cohort_year",
            "class_code",
        ):
            _require(errors, number, field, row.get(field))
        for field in ("student_no", "college_code", "major_code", "class_code"):
            _identifier(errors, number, field, _text(row.get(field)))
        _text_quality(errors, number, "name", desired["name"], 64)
        _text_quality(errors, number, "origin_place", desired["origin_place"], 64)
        _text_quality(errors, number, "contact", desired["contact"], 128)
        _enum(errors, number, "gender", desired["gender"], _values(Gender))
        _enum(errors, number, "degree_level", desired["degree_level"], _values(DegreeLevel))
        if desired["is_key_group"] and desired["is_key_group"] not in _values(KeyGroup):
            errors.append(
                _issue(
                    number,
                    "enum_out_of_range",
                    "重点群体不在允许范围内",
                    column="is_key_group",
                    masked_value="已填写",
                    suggestion=f"允许值：{'、'.join(sorted(_values(KeyGroup)))}；非重点群体请留空",
                )
            )
        cohort_year = desired["cohort_year"]
        if cohort_year is None:
            errors.append(
                _issue(
                    number,
                    "cohort_year_format",
                    "毕业届别必须是四位年份",
                    column="cohort_year",
                    masked_value=_text(row.get("cohort_year")),
                    suggestion="例如：2026",
                )
            )
        elif cohort_year != run_cohort:
            errors.append(
                _issue(
                    number,
                    "cohort_scope_mismatch",
                    "学生届别与本次启用届别不一致",
                    column="cohort_year",
                    masked_value=str(cohort_year),
                    suggestion=f"请填写 {run_cohort} 或另建启用记录",
                )
            )
        class_key = (
            desired["college_code"],
            desired["major_code"],
            desired["degree_level"],
            cohort_year or -1,
            desired["class_code"],
        )
        if class_key not in class_keys:
            errors.append(
                _issue(
                    number,
                    "class_scope_mismatch",
                    "班级与学院、专业、学历层次或届别不匹配",
                    column="class_code",
                    masked_value=desired["class_code"],
                    suggestion="核对班级资料中的完整归属组合",
                )
            )
        if student_no and duplicates[student_no] > 1:
            errors.append(
                _issue(
                    number,
                    "duplicate_in_file",
                    "同一学号在文件中出现多次",
                    column="student_no",
                    masked_value=student_no,
                    suggestion="每名学生只保留一行",
                )
            )
        if len(errors) != start or cohort_year is None:
            continue
        current = existing.get(student_no)
        if current is None:
            action = "inserted"
        else:
            current_major = major_ids.get(current.major_id, ("", "", current.degree_level))
            current_class = (
                class_ids.get(current.class_id, ("", "", "", current.cohort_year, ""))
                if current.class_id
                else ("", "", "", current.cohort_year, "")
            )
            current_external = {
                "student_no": current.student_no,
                "name": current.name,
                "gender": current.gender or "",
                "college_code": college_ids.get(current.college_id, "") or "",
                "major_code": current_major[1],
                "degree_level": current.degree_level,
                "cohort_year": current.cohort_year,
                "class_code": current_class[4],
                "origin_place": current.origin_place,
                "contact": (
                    decrypt_field(current.contact_encrypted)
                    if desired["contact"] is not None
                    else None
                ),
                "is_key_group": current.is_key_group,
            }
            immutable_conflicts = [
                field
                for field in (
                    "college_code",
                    "major_code",
                    "degree_level",
                    "cohort_year",
                )
                if current_external[field] != desired[field]
            ]
            if current.class_id and current_class != class_key:
                immutable_conflicts.append("class_code")
            if immutable_conflicts:
                field = immutable_conflicts[0]
                errors.append(
                    _issue(
                        number,
                        "student_scope_conflict",
                        "该学号已存在，学院、专业、学历层次、届别或既有班级归属不可由启用导入静默变更",
                        column=field,
                        masked_value=desired[field],
                        suggestion="请由就业中心核对既有主档，并通过受控学籍变更流程处理",
                    )
                )
                continue
            action = "unchanged" if current_external == desired else "updated"
        staged.append(
            _staged_item("students", current=current, action=action, data=desired)
        )
        virtual_students.add(student_no)
    return staged, errors


def _validate_outcomes(
    rows: list[dict[str, Any]],
    existing_students: dict[str, Student],
    virtual_students: set[str],
    uploaded_student_actions: dict[str, str],
    run_cohort: int,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    numbers = [_text(row.get("student_no")) for row in rows]
    duplicates = Counter(numbers)
    staged: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for row, student_no in zip(rows, numbers, strict=True):
        start = len(errors)
        number = int(row["_row_number"])
        status = _text(row.get("status"))
        occurred_at = _parse_date(row.get("occurred_at"))
        _require(errors, number, "student_no", student_no)
        _require(errors, number, "status", status)
        _require(errors, number, "occurred_at", row.get("occurred_at"))
        _identifier(errors, number, "student_no", student_no)
        _enum(errors, number, "status", status, _values(EmploymentStatus))
        if row.get("occurred_at") not in (None, "") and occurred_at is None:
            errors.append(
                _issue(
                    number,
                    "date_format_or_range",
                    "登记日期格式不正确",
                    column="occurred_at",
                    masked_value=_text(row.get("occurred_at")),
                    suggestion="请使用 YYYY-MM-DD，例如 2026-04-01",
                )
            )
        elif occurred_at is not None and occurred_at > datetime.now(UTC).date():
            errors.append(
                _issue(
                    number,
                    "date_in_future",
                    "登记日期不能晚于当前日期",
                    column="occurred_at",
                    masked_value=occurred_at.isoformat(),
                    suggestion="请填写已实际发生的登记日期",
                )
            )
        if student_no and student_no not in virtual_students:
            errors.append(
                _issue(
                    number,
                    "parent_code_unknown",
                    "学号在学生资料和系统中都不存在",
                    column="student_no",
                    masked_value=student_no,
                    suggestion="先上传并修正学生资料",
                )
            )
        current = existing_students.get(student_no)
        if current is not None and current.cohort_year != run_cohort:
            errors.append(
                _issue(
                    number,
                    "cohort_scope_mismatch",
                    "该学生不属于本次启用届别",
                    column="student_no",
                    masked_value=student_no,
                    suggestion=f"请仅填写 {run_cohort} 届学生，或另建对应届别启用记录",
                )
            )
        if student_no and duplicates[student_no] > 1:
            errors.append(
                _issue(
                    number,
                    "duplicate_in_file",
                    "同一学生存在多条毕业去向",
                    column="student_no",
                    masked_value=student_no,
                    suggestion="保留当前有效的一条去向记录",
                )
            )
        if len(errors) != start or occurred_at is None:
            continue
        current_status = current.status if current is not None else EmploymentStatus.UNSETTLED.value
        current_date = (
            current.status_occurred_at.date().isoformat()
            if current is not None and current.status_occurred_at
            else None
        )
        desired_date = occurred_at.isoformat()
        action = (
            "unchanged" if current_status == status and current_date == desired_date else "updated"
        )
        staged.append(
            _staged_item(
                "outcomes",
                current=current,
                action=action,
                data={"student_no": student_no, "status": status, "occurred_at": desired_date},
                student_action=uploaded_student_actions.get(student_no),
            )
        )
    return staged, errors


def _acquire_t16_confirm_mutex(db: Session) -> None:
    """Serialize all T16 confirmations with one transaction-scoped lock."""

    db.execute(
        text("SELECT pg_advisory_xact_lock(hashtextextended(:lock_key, 16))"),
        {"lock_key": T16_CONFIRM_ADVISORY_MUTEX_KEY},
    )


def confirm_import_run(
    db: Session,
    *,
    run_id: int,
    revision: int,
    actor: str,
    actor_id: int | None,
    actor_role: str,
    source_channel: str,
) -> ImportRun:
    _require_production_import_profile()
    try:
        # This is intentionally the first database operation in confirmation.
        # One fixed xact lock replaces up to ~100k per-natural-key advisory
        # locks at the five-file 20k ceiling.  PostgreSQL releases it on commit
        # or rollback, including every failure path below.
        _acquire_t16_confirm_mutex(db)
        run = get_import_run(db, run_id, lock=True)
        if run.status == "imported":
            _record_run_audit(
                db,
                run,
                actor=actor,
                actor_id=actor_id,
                actor_role=actor_role,
                source_channel=source_channel,
                action="import.run.confirm_idempotent",
                change_summary={
                    "reason": "该启用会话已经确认，复用既有业务结果",
                    "attempted_revision": revision,
                    "before_status": run.status,
                    "after_status": run.status,
                    "before_revision": run.revision,
                    "after_revision": run.revision,
                },
            )
            db.commit()
            db.refresh(run)
            return run
        if run.status == "rolled_back":
            raise ImportConflictError("该启用记录已经撤销，不能再次确认")
        if revision != run.revision:
            raise ImportConflictError("启用记录已被其他操作更新，请刷新后重新核对")
        before_status = run.status
        before_revision = run.revision
        batches = {batch.target_kind: batch for batch in active_batches(db, run.id)}
        expected_plan = {
            kind: _staged_plan_signature(kind, batch.staged_rows or [])
            for kind, batch in batches.items()
        }
        locked_target_ids = _lock_import_targets(db, batches)
        preflight_import_run(
            db,
            run.id,
            actor=actor,
            actor_id=actor_id,
            actor_role=actor_role,
            source_channel=source_channel,
            trigger="confirm_revalidation",
            commit=False,
            bump_revision=False,
        )
        if run.status != "ready":
            if any(
                issue.get("code") == "natural_key_conflict"
                for batch in batches.values()
                for issue in (batch.errors or [])
            ):
                raise ImportConflictError(
                    "natural_key_conflict：学院名称已被其他学院代码使用，请重新预检并修正"
                )
            raise ValueError("五类资料尚未齐全或仍有问题，不能确认导入")

        batches = {batch.target_kind: batch for batch in active_batches(db, run.id)}
        _assert_confirmation_targets_locked(batches, locked_target_ids)
        _assert_confirmation_plan_current(expected_plan, batches)
        affected_students: set[int] = set()
        pending_target_claims: list[tuple[int, str, str, int, str]] = []
        created_target_ids: dict[type[Any], set[int]] = {
            College: set(),
            Major: set(),
            AcademicClass: set(),
            Student: set(),
        }
        for kind in T16_KIND_ORDER:
            batch = batches[kind]
            refs, student_ids, target_claims = _apply_batch(
                db,
                kind,
                _unprotect_staged_rows(kind, list(batch.staged_rows or [])),
                actor,
                locked_target_ids=locked_target_ids,
                created_target_ids=created_target_ids,
            )
            for target_id, mode in target_claims:
                pending_target_claims.append(
                    (
                        int(batch.id),
                        _claim_resource_kind(kind),
                        _claim_field_scope(kind),
                        target_id,
                        mode,
                    )
                )
            batch.inserted_refs = refs
            batch.status = "imported"
            batch.confirmed_at = datetime.now(UTC)
            batch.confirmed_by = actor
            batch.source_rows = []
            batch.staged_rows = []
            affected_students.update(student_ids)

        _refresh_student_risk(db, affected_students)
        _capture_rollback_after_snapshots(db, batches)
        _capture_dependent_fact_snapshots(db, batches)
        _persist_target_claims(db, pending_target_claims)
        # The dashboard read path is materialized-view-only. Refresh inside the
        # same transaction so operators never commit enabled facts while the
        # leadership view still exposes the previous snapshot.
        refresh_dashboard_views(db, concurrently=False, commit=False)
        run.status = "imported"
        run.confirmed_by = actor
        run.confirmed_at = datetime.now(UTC)
        run.revision += 1
        run.result_summary = _aggregate_summaries(list(batches.values()))
        _record_run_audit(
            db,
            run,
            actor=actor,
            actor_id=actor_id,
            actor_role=actor_role,
            source_channel=source_channel,
            action="import.run.confirm",
            change_summary={
                "reason": "就业中心按当前预检版本显式确认导入并启用",
                "approval_method": "explicit_confirmation_phrase",
                "approved_revision": revision,
                "approved_plan_sha256": _json_fingerprint(expected_plan),
                "batch_ids": [batches[kind].id for kind in T16_KIND_ORDER],
                "batch_changes": _batch_change_evidence(batches, expected_plan),
                "validation_status": "ready",
                "before_status": before_status,
                "after_status": run.status,
                "before_revision": before_revision,
                "after_revision": run.revision,
                **_normalized_summary(run.result_summary),
            },
        )
        db.commit()
        db.refresh(run)
        return run
    except IntegrityError as exc:
        db.rollback()
        message = "确认期间正式数据已被其他操作更新，请刷新并重新预检"
        _record_run_blocked(
            db,
            run_id,
            actor=actor,
            actor_id=actor_id,
            actor_role=actor_role,
            source_channel=source_channel,
            action="import.run.confirm_blocked",
            reason=message,
            attempted_revision=revision,
        )
        raise ImportConflictError(message) from exc
    except OperationalError as exc:
        db.rollback()
        if _is_retryable_transaction_error(exc):
            message = "确认期间发生并发冲突，请刷新后重新预检"
            _record_run_blocked(
                db,
                run_id,
                actor=actor,
                actor_id=actor_id,
                actor_role=actor_role,
                source_channel=source_channel,
                action="import.run.confirm_blocked",
                reason=message,
                attempted_revision=revision,
            )
            raise ImportConflictError(message) from exc
        raise
    except (ImportConflictError, ValueError) as exc:
        db.rollback()
        _record_run_blocked(
            db,
            run_id,
            actor=actor,
            actor_id=actor_id,
            actor_role=actor_role,
            source_channel=source_channel,
            action="import.run.confirm_blocked",
            reason=str(exc),
            attempted_revision=revision,
        )
        raise
    except Exception:
        db.rollback()
        raise


def _assert_confirmation_plan_current(
    expected_plan: dict[str, list[dict[str, str | None]]],
    batches: dict[str, ImportBatch],
) -> None:
    """Reject a target change that the operator has not re-preflighted.

    A concurrent run may safely converge a missing or stale row to the exact
    desired value.  Any other before-state change requires a new preflight and
    an explicit second confirmation instead of silently widening the write.
    """

    current_plan = {
        kind: _staged_plan_signature(kind, batch.staged_rows or [])
        for kind, batch in batches.items()
    }
    if expected_plan.keys() != current_plan.keys():
        raise ImportConflictError("正式数据在预检后发生变化，请重新预检并再次确认")
    for kind in T16_KIND_ORDER:
        expected_items = {
            str(item["key_sha256"]): item for item in expected_plan.get(kind, [])
        }
        current_items = {
            str(item["key_sha256"]): item for item in current_plan.get(kind, [])
        }
        if expected_items.keys() != current_items.keys():
            raise ImportConflictError("正式数据在预检后发生变化，请重新预检并再次确认")
        for key_hash, expected in expected_items.items():
            current = current_items[key_hash]
            if expected["desired_sha256"] != current["desired_sha256"]:
                raise ImportConflictError("启用资料在确认前发生变化，请重新预检并再次确认")
            if expected["expected_before_sha256"] == current["expected_before_sha256"]:
                continue
            if current["action"] == "unchanged":
                # Another serialized run already wrote the exact desired fact.
                # Treat that as idempotent convergence, not a silent overwrite.
                continue
            raise ImportConflictError("正式数据在预检后发生变化，请重新预检并再次确认")


def _assert_confirmation_targets_locked(
    batches: dict[str, ImportBatch],
    locked_target_ids: dict[type[Any], set[int]],
) -> None:
    """Reject a row that appeared only after the target-lock snapshot."""

    for kind in T16_KIND_ORDER:
        batch = batches.get(kind)
        if batch is None:
            continue
        model = _model_for_kind(kind)
        for item in batch.staged_rows or []:
            existing_id = item.get("existing_id")
            if existing_id is None:
                continue
            if int(existing_id) not in locked_target_ids.get(model, set()):
                raise ImportConflictError("正式数据在目标锁定后发生变化，请重新预检并再次确认")


def _apply_batch(
    db: Session,
    kind: str,
    staged: list[dict[str, Any]],
    actor: str,
    *,
    locked_target_ids: dict[type[Any], set[int]],
    created_target_ids: dict[type[Any], set[int]],
) -> tuple[dict[str, Any], set[int], list[tuple[int, str]]]:
    refs: dict[str, Any] = {"inserted_ids": [], "updates": []}
    affected_students: set[int] = set()
    target_claims: list[tuple[int, str]] = []
    state = _load_scoped_import_state(db, _staged_lookup_scope(kind, staged))
    colleges = state.colleges
    majors = state.majors
    classes = state.classes
    students = state.students

    for item in sorted(staged, key=lambda row: _import_lock_key(kind, row["data"])):
        action = item["action"]
        data = item["data"]
        current = _resolved_import_target(kind, data, state)
        _assert_apply_target_current(
            kind,
            item,
            current,
            locked_target_ids=locked_target_ids,
            created_target_ids=created_target_ids,
        )
        if action == "unchanged":
            target_claims.append((int(current.id), action))
            continue
        if kind == "colleges":
            if current is None:
                current = College(code=data["code"], name=data["name"], created_by=actor)
                db.add(current)
                db.flush()
                colleges[data["code"]] = current
                refs["inserted_ids"].append(current.id)
                created_target_ids[College].add(int(current.id))
            else:
                _record_update(refs, current, {"name": data["name"]}, ("name",))
        elif kind == "majors":
            key = (data["college_code"], data["code"], data["level"])
            college = colleges.get(data["college_code"])
            _assert_apply_parent_allowed(
                college,
                College,
                locked_target_ids=locked_target_ids,
                created_target_ids=created_target_ids,
            )
            if current is None:
                current = Major(
                    college_id=college.id,
                    code=data["code"],
                    name=data["name"],
                    level=data["level"],
                    created_by=actor,
                )
                db.add(current)
                db.flush()
                majors[key] = current
                refs["inserted_ids"].append(current.id)
                created_target_ids[Major].add(int(current.id))
            else:
                _record_update(refs, current, {"name": data["name"]}, ("name",))
        elif kind == "classes":
            key = (
                data["college_code"],
                data["major_code"],
                data["level"],
                data["cohort_year"],
                data["code"],
            )
            college = colleges.get(data["college_code"])
            major = majors.get((data["college_code"], data["major_code"], data["level"]))
            _assert_apply_parent_allowed(
                college,
                College,
                locked_target_ids=locked_target_ids,
                created_target_ids=created_target_ids,
            )
            _assert_apply_parent_allowed(
                major,
                Major,
                locked_target_ids=locked_target_ids,
                created_target_ids=created_target_ids,
            )
            if current is None:
                current = AcademicClass(
                    college_id=college.id,
                    major_id=major.id,
                    cohort_year=data["cohort_year"],
                    code=data["code"],
                    name=data["name"],
                    created_by=actor,
                )
                db.add(current)
                db.flush()
                classes[key] = current
                refs["inserted_ids"].append(current.id)
                created_target_ids[AcademicClass].add(int(current.id))
            else:
                _record_update(refs, current, {"name": data["name"]}, ("name",))
        elif kind == "students":
            class_key = (
                data["college_code"],
                data["major_code"],
                data["degree_level"],
                data["cohort_year"],
                data["class_code"],
            )
            college = colleges.get(data["college_code"])
            major = majors.get((data["college_code"], data["major_code"], data["degree_level"]))
            academic_class = classes.get(class_key)
            _assert_apply_parent_allowed(
                college,
                College,
                locked_target_ids=locked_target_ids,
                created_target_ids=created_target_ids,
            )
            _assert_apply_parent_allowed(
                major,
                Major,
                locked_target_ids=locked_target_ids,
                created_target_ids=created_target_ids,
            )
            _assert_apply_parent_allowed(
                academic_class,
                AcademicClass,
                locked_target_ids=locked_target_ids,
                created_target_ids=created_target_ids,
            )
            desired = {
                "name": data["name"],
                "gender": data["gender"],
                "college_id": college.id,
                "major_id": major.id,
                "class_id": academic_class.id,
                "degree_level": data["degree_level"],
                "cohort_year": data["cohort_year"],
                "origin_place": data["origin_place"],
                "contact_encrypted": encrypt_field(data["contact"]),
                "is_key_group": data["is_key_group"],
            }
            if current is None:
                current = Student(
                    student_no=data["student_no"],
                    status=EmploymentStatus.UNSETTLED.value,
                    risk_level="green",
                    created_by=actor,
                    **desired,
                )
                db.add(current)
                db.flush()
                students[data["student_no"]] = current
                refs["inserted_ids"].append(current.id)
                created_target_ids[Student].add(int(current.id))
                affected_students.add(current.id)
            else:
                existing_desired = {
                    "name": desired["name"],
                    "gender": desired["gender"],
                    "origin_place": desired["origin_place"],
                    "is_key_group": desired["is_key_group"],
                }
                if data["contact"] is not None:
                    existing_desired["contact_encrypted"] = encrypt_field(data["contact"])
                if current.class_id is None:
                    existing_desired["class_id"] = desired["class_id"]
                _record_update(
                    refs,
                    current,
                    existing_desired,
                    # Confirmation recalculates risk after all five batches.
                    # Keep the original risk fields in the same encrypted
                    # rollback snapshot even when the student master row was
                    # the only direct change; otherwise rollback can leave a
                    # recomputed risk fact behind.
                    (*tuple(existing_desired), "risk_level", "risk_reason"),
                    sensitive=True,
                )
                affected_students.add(current.id)
        elif kind == "outcomes":
            desired = {
                "status": data["status"],
                "status_occurred_at": datetime.fromisoformat(data["occurred_at"]).replace(
                    tzinfo=UTC
                ),
            }
            _record_update(
                refs,
                current,
                desired,
                ("status", "status_occurred_at", "risk_level", "risk_reason"),
            )
            affected_students.add(current.id)
        else:
            raise ValueError("未知资料类型")
        target_claims.append((int(current.id), action))
    db.flush()
    return refs, affected_students, target_claims


def _resolved_import_target(
    kind: str,
    data: dict[str, Any],
    state: _ImportLookupState,
) -> Any | None:
    if kind == "colleges":
        return state.colleges.get(data["code"])
    if kind == "majors":
        return state.majors.get((data["college_code"], data["code"], data["level"]))
    if kind == "classes":
        return state.classes.get(
            (
                data["college_code"],
                data["major_code"],
                data["level"],
                data["cohort_year"],
                data["code"],
            )
        )
    if kind in {"students", "outcomes"}:
        return state.students.get(data["student_no"])
    raise ValueError("未知资料类型")


def _assert_apply_target_current(
    kind: str,
    item: dict[str, Any],
    current: Any | None,
    *,
    locked_target_ids: dict[type[Any], set[int]],
    created_target_ids: dict[type[Any], set[int]],
) -> None:
    """Prove a post-assertion re-read cannot adopt an unlocked target.

    Under READ COMMITTED, a row with a previously missing natural key may be
    committed after the final preflight assertion.  Existing targets must
    therefore match the staged ID exactly and belong to the row-lock set.
    The sole missing-ID exception is an outcome referencing the Student that
    this same transaction inserted in the immediately preceding student batch.
    """

    model = _model_for_kind(kind)
    expected_id_value = item.get("existing_id")
    expected_id = int(expected_id_value) if expected_id_value is not None else None
    if expected_id is not None:
        if (
            current is None
            or int(current.id) != expected_id
            or expected_id not in locked_target_ids.get(model, set())
        ):
            raise ImportConflictError(
                "正式数据在最终核对后发生变化，请重新预检并再次确认"
            )
        return

    if current is None:
        if kind == "outcomes":
            raise ImportConflictError(
                "正式数据在最终核对后发生变化，请重新预检并再次确认"
            )
        return

    current_id = int(current.id)
    if (
        kind == "outcomes"
        and item.get("student_action") == "inserted"
        and current_id in created_target_ids.get(Student, set())
    ):
        return
    raise ImportConflictError("正式数据在最终核对后发生变化，请重新预检并再次确认")


def _assert_apply_parent_allowed(
    current: Any | None,
    model: type[Any],
    *,
    locked_target_ids: dict[type[Any], set[int]],
    created_target_ids: dict[type[Any], set[int]],
) -> None:
    """Require every resolved parent to be locked or created by this confirm."""

    if current is None:
        raise ImportConflictError("正式数据在最终核对后发生变化，请重新预检并再次确认")
    current_id = int(current.id)
    if current_id in locked_target_ids.get(model, set()):
        return
    if current_id in created_target_ids.get(model, set()):
        return
    raise ImportConflictError("正式数据在最终核对后发生变化，请重新预检并再次确认")


def _lock_import_targets(
    db: Session,
    batches: dict[str, ImportBatch],
) -> dict[type[Any], set[int]]:
    """Resolve relevant natural keys to IDs, then row-lock in one hierarchy.

    The caller already holds the fixed T16 confirmation mutex, which safely
    serializes keys that do not exist yet.  Existing rows are resolved without
    locks first, then locked only by primary key in the fixed model order below.
    Each model's IDs are globally sorted before chunking, so chunk boundaries
    cannot reintroduce natural-key-dependent lock ordering.
    """

    scope = _ImportLookupScope()
    ids_by_model: dict[type[Any], set[int]] = {
        College: set(),
        Major: set(),
        AcademicClass: set(),
        Student: set(),
    }
    required_existing_ids: dict[type[Any], set[int]] = {
        College: set(),
        Major: set(),
        AcademicClass: set(),
        Student: set(),
    }
    for kind in T16_KIND_ORDER:
        batch = batches.get(kind)
        if batch is None:
            continue
        staged = [
            item
            for item in (batch.staged_rows or [])
            if isinstance(item.get("data"), dict)
        ]
        batch_scope = _staged_lookup_scope(kind, staged)
        for attribute in (
            "college_codes",
            "college_names",
            "major_keys",
            "class_keys",
            "student_numbers",
            "student_scope_numbers",
        ):
            getattr(scope, attribute).update(getattr(batch_scope, attribute))
        model = _model_for_kind(kind)
        staged_existing_ids = {
            int(item["existing_id"])
            for item in staged
            if item.get("existing_id") is not None and int(item["existing_id"]) > 0
        }
        ids_by_model[model].update(staged_existing_ids)
        required_existing_ids[model].update(staged_existing_ids)

    state = _load_scoped_import_state(db, scope)
    ids_by_model[College].update(row.id for row in state.colleges.values())
    ids_by_model[Major].update(row.id for row in state.majors.values())
    ids_by_model[AcademicClass].update(row.id for row in state.classes.values())
    ids_by_model[Student].update(row.id for row in state.students.values())

    locked_ids_by_model: dict[type[Any], set[int]] = {
        College: set(),
        Major: set(),
        AcademicClass: set(),
        Student: set(),
    }
    for model in (College, Major, AcademicClass, Student):
        for chunk in _postgres_bind_chunks(ids_by_model[model], tuple_width=1):
            locked_rows = list(
                db.execute(
                    select(model)
                    .where(model.id.in_(chunk))
                    .order_by(model.id)
                    .with_for_update()
                    # A row may have changed while this transaction waited for
                    # its lock.  Refresh identity-map objects with the latest
                    # READ COMMITTED version before final preflight.
                    .execution_options(populate_existing=True)
                ).scalars()
            )
            locked_ids_by_model[model].update(int(row.id) for row in locked_rows)
        if not required_existing_ids[model].issubset(locked_ids_by_model[model]):
            raise ImportConflictError("正式数据在预检后被删除，请重新预检并再次确认")
    return locked_ids_by_model


def _import_lock_key(kind: str, data: dict[str, Any]) -> str:
    if kind == "colleges":
        parts = (data.get("code"),)
    elif kind == "majors":
        parts = (data.get("college_code"), data.get("code"), data.get("level"))
    elif kind == "classes":
        parts = (
            data.get("college_code"),
            data.get("major_code"),
            data.get("level"),
            data.get("cohort_year"),
            data.get("code"),
        )
    elif kind in {"students", "outcomes"}:
        parts = (data.get("student_no"),)
    else:
        raise ValueError("未知资料类型")
    return f"t16:{kind if kind not in {'students', 'outcomes'} else 'student'}:" + ":".join(
        str(value or "") for value in parts
    )


def _is_retryable_transaction_error(exc: OperationalError) -> bool:
    original = getattr(exc, "orig", None)
    sqlstate = getattr(original, "sqlstate", None) or getattr(original, "pgcode", None)
    return sqlstate in {"40001", "40P01"}


def _record_update(
    refs: dict[str, Any],
    model: Any,
    desired: dict[str, Any],
    fields: tuple[str, ...],
    *,
    sensitive: bool = False,
) -> None:
    before = _snapshot(model, fields)
    for field, value in desired.items():
        setattr(model, field, value)
    after = _snapshot(model, fields)
    if sensitive:
        encrypted = encrypt_field(
            json.dumps(before, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        )
        if encrypted is None:
            raise ValueError("敏感回滚快照加密失败")
        refs["updates"].append(
            {
                "id": model.id,
                "fields": list(fields),
                "before_encrypted": encrypted,
                "after_fingerprint": _dict_fingerprint(after),
            }
        )
    else:
        refs["updates"].append({"id": model.id, "before": before, "after": after})


def _refresh_student_risk(db: Session, student_ids: set[int]) -> None:
    from app.services.risk import RiskInput, evaluate_risk

    if not student_ids:
        return
    students = list(
        db.execute(
            select(Student).where(Student.id.in_(sorted(student_ids))).order_by(Student.id)
        ).scalars()
    )
    if not students:
        return
    college_ids = sorted({student.college_id for student in students})
    action_rows = db.execute(
        select(
            JobAction.student_id,
            func.max(JobAction.occurred_at),
            func.count(JobAction.id).filter(JobAction.type == JobActionType.APPLY.value),
        )
        .where(JobAction.student_id.in_(sorted(student_ids)))
        .group_by(JobAction.student_id)
    ).all()
    action_by_student = {
        int(student_id): (last_action_at, int(apply_count or 0))
        for student_id, last_action_at, apply_count in action_rows
    }
    student_counts = {
        int(college_id): int(count)
        for college_id, count in db.execute(
            select(Student.college_id, func.count(Student.id))
            .where(Student.college_id.in_(college_ids))
            .group_by(Student.college_id)
        ).all()
    }
    college_apply_counts = {
        int(college_id): int(count)
        for college_id, count in db.execute(
            select(Student.college_id, func.count(JobAction.id))
            .join(JobAction, JobAction.student_id == Student.id)
            .where(
                Student.college_id.in_(college_ids),
                JobAction.type == JobActionType.APPLY.value,
            )
            .group_by(Student.college_id)
        ).all()
    }
    as_of = datetime.now(UTC)
    thresholds = load_risk_thresholds(db)
    for student in students:
        last_action_at, apply_count = action_by_student.get(student.id, (None, 0))
        weeks_since_last_action = (
            None if last_action_at is None else max((as_of - last_action_at).days // 7, 0)
        )
        college_student_count = student_counts.get(student.college_id, 0)
        result = evaluate_risk(
            RiskInput(
                status=student.status,
                is_graduation_year=student.cohort_year == as_of.year,
                current_month=as_of.month,
                weeks_since_last_job_action=weeks_since_last_action,
                apply_count=apply_count,
                college_avg_apply=(
                    college_apply_counts.get(student.college_id, 0) / college_student_count
                    if college_student_count
                    else 0.0
                ),
            ),
            thresholds,
        )
        student.risk_level = result.level.value
        student.risk_reason = "；".join(result.reasons) or None
    db.flush()


def _capture_rollback_after_snapshots(
    db: Session,
    batches: dict[str, ImportBatch],
) -> None:
    """Freeze post-confirm facts used to refuse a destructive stale rollback."""

    for kind, batch in batches.items():
        refs = dict(batch.inserted_refs or {})
        model = _model_for_kind(kind)
        inserted_fingerprints: dict[str, str] = {}
        for row_id in refs.get("inserted_ids") or []:
            row = db.get(model, int(row_id))
            if row is not None:
                inserted_fingerprints[str(row_id)] = _snapshot_fingerprint(
                    row,
                    _inserted_snapshot_fields(kind),
                )
        refs["inserted_fingerprints"] = inserted_fingerprints

        updates = list(refs.get("updates") or [])
        for update in updates:
            row = db.get(model, int(update["id"]))
            if row is not None:
                if update.get("before_encrypted"):
                    update["after_fingerprint"] = _snapshot_fingerprint(
                        row,
                        tuple(update["fields"]),
                    )
                else:
                    update["after"] = _snapshot(row, tuple(update["before"]))
        refs["updates"] = updates
        batch.inserted_refs = refs
    db.flush()


def rollback_import_run(
    db: Session,
    *,
    run_id: int,
    revision: int,
    reason: str,
    actor: str,
    actor_id: int | None,
    actor_role: str,
    source_channel: str,
) -> ImportRun:
    _require_production_import_profile()
    rollback_reason = _bounded_text(reason, 500, field="撤销原因", required=True)
    try:
        run = get_import_run(db, run_id, lock=True)
        if run.status == "rolled_back":
            _record_run_audit(
                db,
                run,
                actor=actor,
                actor_id=actor_id,
                actor_role=actor_role,
                source_channel=source_channel,
                action="import.run.rollback_idempotent",
                change_summary={
                    "reason": "该启用会话已经撤销，复用既有恢复结果",
                    "attempted_revision": revision,
                    "before_status": run.status,
                    "after_status": run.status,
                    "before_revision": run.revision,
                    "after_revision": run.revision,
                },
            )
            db.commit()
            db.refresh(run)
            return run
        if run.status != "imported":
            raise ImportConflictError("只有已启用的记录可以撤销")
        if revision != run.revision:
            raise ImportConflictError("启用记录已被其他操作更新，请刷新后重新核对")
        before_status = run.status
        before_revision = run.revision
        batches = {batch.target_kind: batch for batch in active_batches(db, run.id)}

        target_claims = _lock_rollback_targets(db, batches)
        _assert_no_later_import_claims(db, run.id, batches)
        _assert_no_later_student_facts(db, batches, run.confirmed_at)
        for kind in reversed(T16_KIND_ORDER):
            batch = batches[kind]
            _verify_batch_rollback(db, kind, batch.inserted_refs or {})
        for kind in reversed(T16_KIND_ORDER):
            batch = batches[kind]
            _rollback_batch_refs(db, kind, batch.inserted_refs or {})
            batch.status = "rolled_back"
            batch.rolled_back_at = datetime.now(UTC)
            batch.rolled_back_by = actor
        released_at = datetime.now(UTC)
        for claim in target_claims:
            claim.released_at = released_at
        db.flush()

        # Rollback and its dashboard projection must become visible atomically;
        # otherwise a removed student can remain on the leadership dashboard.
        refresh_dashboard_views(db, concurrently=False, commit=False)

        run.status = "rolled_back"
        run.rolled_back_by = actor
        run.rolled_back_at = datetime.now(UTC)
        run.rollback_reason = rollback_reason
        run.revision += 1
        _record_run_audit(
            db,
            run,
            actor=actor,
            actor_id=actor_id,
            actor_role=actor_role,
            source_channel=source_channel,
            action="import.run.rollback",
            change_summary={
                "reason": rollback_reason,
                "batch_ids": [batches[kind].id for kind in T16_KIND_ORDER],
                "batch_changes": _batch_change_evidence(
                    batches,
                    None,
                ),
                "before_status": before_status,
                "after_status": run.status,
                "before_revision": before_revision,
                "after_revision": run.revision,
                **_normalized_summary(run.result_summary),
            },
        )
        db.commit()
        db.refresh(run)
        return run
    except ImportConflictError as exc:
        db.rollback()
        _record_run_blocked(
            db,
            run_id,
            actor=actor,
            actor_id=actor_id,
            actor_role=actor_role,
            source_channel=source_channel,
            action="import.run.rollback_blocked",
            reason=str(exc),
            attempted_revision=revision,
        )
        raise
    except IntegrityError as exc:
        db.rollback()
        message = "已有后续业务引用本次导入的数据，不能直接撤销"
        _record_run_blocked(
            db,
            run_id,
            actor=actor,
            actor_id=actor_id,
            actor_role=actor_role,
            source_channel=source_channel,
            action="import.run.rollback_blocked",
            reason=message,
            attempted_revision=revision,
        )
        raise ImportConflictError(message) from exc
    except OperationalError as exc:
        db.rollback()
        if _is_retryable_transaction_error(exc):
            message = "撤销期间发生并发冲突，请刷新后重试"
            _record_run_blocked(
                db,
                run_id,
                actor=actor,
                actor_id=actor_id,
                actor_role=actor_role,
                source_channel=source_channel,
                action="import.run.rollback_blocked",
                reason=message,
                attempted_revision=revision,
            )
            raise ImportConflictError(message) from exc
        raise
    except ValueError:
        db.rollback()
        _record_run_blocked(
            db,
            run_id,
            actor=actor,
            actor_id=actor_id,
            actor_role=actor_role,
            source_channel=source_channel,
            action="import.run.rollback_blocked",
            reason="撤销证据无法安全校验或恢复，请联系系统管理员核验",
            attempted_revision=revision,
        )
        raise
    except Exception:
        db.rollback()
        _record_run_blocked(
            db,
            run_id,
            actor=actor,
            actor_id=actor_id,
            actor_role=actor_role,
            source_channel=source_channel,
            action="import.run.rollback_blocked",
            reason="撤销处理发生未预期错误，业务事实未改变，请联系系统管理员核验",
            attempted_revision=revision,
        )
        raise


def _lock_rollback_targets(
    db: Session,
    batches: dict[str, ImportBatch],
) -> list[ImportTargetClaim]:
    """Close the verify-then-restore race by locking every referenced fact row.

    Locks are acquired in the same fixed hierarchy order used by confirmation.
    PostgreSQL keeps them until the rollback transaction commits or aborts, so a
    concurrent business update cannot land between fingerprint verification and
    the restore/delete operation.
    """

    ids_by_model: dict[type[Any], set[int]] = {
        College: set(),
        Major: set(),
        AcademicClass: set(),
        Student: set(),
    }
    batch_ids = [batches[kind].id for kind in T16_KIND_ORDER if kind in batches]
    target_claims = list(
        db.execute(
            select(ImportTargetClaim)
            .where(
                ImportTargetClaim.batch_id.in_(batch_ids),
                ImportTargetClaim.released_at.is_(None),
            )
            .order_by(
                ImportTargetClaim.resource_kind,
                ImportTargetClaim.target_id,
                ImportTargetClaim.id,
            )
            .with_for_update()
        ).scalars()
    )
    claimed_batch_ids = {claim.batch_id for claim in target_claims}
    if any(batch.row_count > 0 and batch.id not in claimed_batch_ids for batch in batches.values()):
        raise ImportConflictError(
            "启用记录缺少目标依赖证据，不能安全撤销"
        )
    for kind in T16_KIND_ORDER:
        batch = batches.get(kind)
        if batch is None:
            raise ImportConflictError("启用记录的批次证据不完整，不能安全撤销")
        refs = batch.inserted_refs or {}
        row_ids = {int(value) for value in refs.get("inserted_ids") or []} | {
            int(update["id"])
            for update in refs.get("updates") or []
            if isinstance(update, dict) and update.get("id") is not None
        }
        ids_by_model[_model_for_kind(kind)].update(row_ids)

    for claim in target_claims:
        ids_by_model[_model_for_claim_resource(claim.resource_kind)].add(claim.target_id)

    for model in (College, Major, AcademicClass, Student):
        row_ids = ids_by_model[model]
        if not row_ids:
            continue
        list(
            db.execute(
                select(model)
                .where(model.id.in_(sorted(row_ids)))
                .order_by(model.id)
                .with_for_update()
            ).scalars()
        )
    return target_claims


def _persist_target_claims(
    db: Session,
    pending: list[tuple[int, str, str, int, str]],
) -> None:
    """Allocate a monotonic version within each locked business target."""

    if not pending:
        return
    keys = sorted({(resource_kind, target_id) for _, resource_kind, _, target_id, _ in pending})
    latest_versions: dict[tuple[str, int], int] = {}
    # Keep bind counts comfortably below PostgreSQL's protocol limit for the
    # 20k-row T16 ceiling, including runs whose student and outcome sets differ.
    for offset in range(0, len(keys), 10_000):
        chunk = keys[offset : offset + 10_000]
        rows = db.execute(
            select(
                ImportTargetClaim.resource_kind,
                ImportTargetClaim.target_id,
                func.max(ImportTargetClaim.target_version),
            )
            .where(
                tuple_(
                    ImportTargetClaim.resource_kind,
                    ImportTargetClaim.target_id,
                ).in_(chunk)
            )
            .group_by(
                ImportTargetClaim.resource_kind,
                ImportTargetClaim.target_id,
            )
        ).all()
        for resource_kind, target_id, latest in rows:
            latest_versions[(str(resource_kind), int(target_id))] = int(latest or 0)

    for batch_id, resource_kind, field_scope, target_id, mode in pending:
        key = (resource_kind, target_id)
        target_version = latest_versions.get(key, 0) + 1
        latest_versions[key] = target_version
        db.add(
            ImportTargetClaim(
                batch_id=batch_id,
                resource_kind=resource_kind,
                field_scope=field_scope,
                target_id=target_id,
                target_version=target_version,
                mode=mode,
            )
        )
    db.flush()


def _assert_no_later_import_claims(
    db: Session,
    run_id: int,
    batches: dict[str, ImportBatch],
) -> None:
    """Refuse to undo a fact still adopted by a causally later imported run.

    Target versions are allocated while the fixed confirmation mutex and the
    corresponding existing business row lock are held.  They express causal
    order per target, independent of primary-key sequences, connection caches
    or wall clocks.
    """

    batch_ids = [batches[kind].id for kind in T16_KIND_ORDER if kind in batches]
    if not batch_ids:
        return

    own_claim = aliased(ImportTargetClaim)
    dependent_claim = db.execute(
        select(ImportTargetClaim.id)
        .select_from(ImportTargetClaim)
        .join(
            own_claim,
            (own_claim.batch_id.in_(batch_ids))
            & (own_claim.released_at.is_(None))
            & (own_claim.mode.in_(("inserted", "updated")))
            & (ImportTargetClaim.resource_kind == own_claim.resource_kind)
            & (ImportTargetClaim.target_id == own_claim.target_id)
            & (ImportTargetClaim.target_version > own_claim.target_version)
            & (
                (own_claim.mode == "inserted")
                | (
                    (own_claim.mode == "updated")
                    & (
                        (own_claim.resource_kind == "student")
                        | (ImportTargetClaim.field_scope == own_claim.field_scope)
                    )
                )
            ),
        )
        .join(ImportBatch, ImportBatch.id == ImportTargetClaim.batch_id)
        .join(ImportRun, ImportRun.id == ImportBatch.run_id)
        .where(
            ImportBatch.run_id != run_id,
            ImportBatch.status == "imported",
            ImportRun.status == "imported",
            ImportTargetClaim.released_at.is_(None),
        )
        .limit(1)
    ).scalar_one_or_none()
    if dependent_claim is not None:
        raise ImportConflictError(
            "已有后续启用记录采用本次数据；请先撤销后续记录，再重试本次撤销"
        )


def _capture_dependent_fact_snapshots(
    db: Session,
    batches: dict[str, ImportBatch],
) -> None:
    """Freeze privacy-safe fingerprints for facts linked to updated students."""

    student_batch = batches.get("students")
    if student_batch is None:
        raise ImportConflictError("启用记录的学生批次证据不完整，不能安全确认")
    updated_ids = _updated_existing_student_ids(batches)
    snapshot = _dependent_fact_snapshot(db, updated_ids, lock=True)
    refs = dict(student_batch.inserted_refs or {})
    refs["dependent_facts_at_confirm"] = {
        "version": DEPENDENT_FACT_SNAPSHOT_VERSION,
        "student_ids": sorted(updated_ids),
        "tables": snapshot,
    }
    student_batch.inserted_refs = refs
    db.flush()


def _updated_existing_student_ids(batches: dict[str, ImportBatch]) -> set[int]:
    student_batch = batches.get("students")
    outcome_batch = batches.get("outcomes")
    student_refs = student_batch.inserted_refs or {} if student_batch else {}
    outcome_refs = outcome_batch.inserted_refs or {} if outcome_batch else {}
    inserted_ids = {int(value) for value in student_refs.get("inserted_ids") or []}
    return {
        int(update["id"])
        for refs in (student_refs, outcome_refs)
        for update in refs.get("updates") or []
        if isinstance(update, dict) and update.get("id") is not None
    } - inserted_ids


def _dependent_fact_snapshot(
    db: Session,
    student_ids: set[int],
    *,
    lock: bool,
) -> dict[str, list[dict[str, str]]]:
    """Return keyed row fingerprints for every direct or indirect student fact."""

    normalized_ids = sorted({int(value) for value in student_ids})
    if not normalized_ids:
        return {}

    bind = db.get_bind()
    inspector = inspect(bind)
    metadata = MetaData()
    snapshot: dict[str, list[dict[str, str]]] = {}
    for table_name in sorted(inspector.get_table_names()):
        student_columns = [
            fk["constrained_columns"][0]
            for fk in inspector.get_foreign_keys(table_name)
            if fk.get("referred_table") == "students"
            and fk.get("referred_columns") == ["id"]
            and len(fk.get("constrained_columns") or []) == 1
        ]
        if not student_columns:
            continue
        table = Table(table_name, metadata, autoload_with=bind)
        conditions = [
            table.c[column_name].in_(normalized_ids) for column_name in sorted(set(student_columns))
        ]
        order_columns = list(table.primary_key.columns) or [
            table.c[column_name] for column_name in sorted(set(student_columns))
        ]
        statement = select(table).where(or_(*conditions)).order_by(*order_columns)
        if lock:
            statement = statement.with_for_update()
        rows = list(db.execute(statement).mappings())
        fingerprints = _fact_row_fingerprints(table, rows)
        if fingerprints:
            snapshot[table_name] = fingerprints

    assist_statement = (
        select(AssistRecord.__table__)
        .join(AssistCase, AssistCase.id == AssistRecord.case_id)
        .where(AssistCase.student_id.in_(normalized_ids))
        .order_by(AssistRecord.id)
    )
    if lock:
        assist_statement = assist_statement.with_for_update(of=AssistRecord)
    assist_rows = list(db.execute(assist_statement).mappings())
    assist_fingerprints = _fact_row_fingerprints(AssistRecord.__table__, assist_rows)
    if assist_fingerprints:
        snapshot["assist_records_via_case"] = assist_fingerprints
    return snapshot


def _fact_row_fingerprints(
    table: Table,
    rows: list[Any],
) -> list[dict[str, str]]:
    primary_keys = list(table.primary_key.columns)
    fingerprints: list[dict[str, str]] = []
    for row in rows:
        payload = {column.name: _canonical_fact_value(row[column.name]) for column in table.columns}
        identity = (
            {column.name: _canonical_fact_value(row[column.name]) for column in primary_keys}
            if primary_keys
            else payload
        )
        fingerprints.append(
            {
                "identity_sha256": _dict_fingerprint(identity),
                "row_sha256": _dict_fingerprint(payload),
            }
        )
    return sorted(
        fingerprints,
        key=lambda item: (item["identity_sha256"], item["row_sha256"]),
    )


def _canonical_fact_value(value: Any) -> Any:
    if isinstance(value, datetime | date):
        return value.isoformat()
    if isinstance(value, bytes):
        return {"sha256": sha256(value).hexdigest(), "length": len(value)}
    if isinstance(value, dict):
        return {
            str(key): _canonical_fact_value(item)
            for key, item in sorted(value.items(), key=lambda pair: str(pair[0]))
        }
    if isinstance(value, list | tuple):
        return [_canonical_fact_value(item) for item in value]
    if value is None or isinstance(value, str | int | float | bool):
        return value
    return str(value)


def _snapshot_has_rows(snapshot: dict[str, list[dict[str, str]]]) -> bool:
    return any(rows for rows in snapshot.values())


def _assert_no_later_student_facts(
    db: Session,
    batches: dict[str, ImportBatch],
    confirmed_at: datetime | None,
) -> None:
    """Refuse rollback when dependent facts differ from confirmation."""

    if confirmed_at is None:
        raise ImportConflictError("启用记录缺少确认时间，不能安全撤销")
    student_batch = batches.get("students")
    outcome_batch = batches.get("outcomes")
    if student_batch is None or outcome_batch is None:
        raise ImportConflictError("启用记录的学生事实证据不完整，不能安全撤销")

    student_refs = student_batch.inserted_refs or {}
    inserted_ids = {int(value) for value in student_refs.get("inserted_ids") or []}
    updated_ids = _updated_existing_student_ids(batches)
    expected = student_refs.get("dependent_facts_at_confirm")
    if (
        not isinstance(expected, dict)
        or expected.get("version") != DEPENDENT_FACT_SNAPSHOT_VERSION
        or expected.get("student_ids") != sorted(updated_ids)
        or not isinstance(expected.get("tables"), dict)
    ):
        raise ImportConflictError("启用记录缺少可靠的后续事实快照，不能安全撤销")

    inserted_snapshot = _dependent_fact_snapshot(db, inserted_ids, lock=True)
    if _snapshot_has_rows(inserted_snapshot):
        raise ImportConflictError(
            "启用后已产生学生绑定、授权或业务事实，为保护后续记录，本次撤销已停止"
        )

    current_snapshot = _dependent_fact_snapshot(db, updated_ids, lock=True)
    if current_snapshot != expected["tables"]:
        raise ImportConflictError(
            "启用后的学生业务事实已经新增、删除或修改，为保护最新事实，本次撤销已停止"
        )


def _verify_batch_rollback(db: Session, kind: str, refs: dict[str, Any]) -> None:
    model = _model_for_kind(kind)
    for raw_id, expected in (refs.get("inserted_fingerprints") or {}).items():
        row = db.get(model, int(raw_id))
        if row is None or _snapshot_fingerprint(
            row,
            _inserted_snapshot_fields(kind),
        ) != str(expected):
            raise ImportConflictError("数据在启用后已被修改，为保护最新事实，本次撤销已停止")
    for update in refs.get("updates") or []:
        row = db.get(model, int(update["id"]))
        if update.get("before_encrypted"):
            matches = row is not None and _snapshot_fingerprint(
                row,
                tuple(update["fields"]),
            ) == str(update.get("after_fingerprint") or "")
        else:
            matches = (
                row is not None
                and _snapshot(
                    row,
                    tuple(update["after"]),
                )
                == update["after"]
            )
        if not matches:
            raise ImportConflictError("数据在启用后已被修改，为保护最新事实，本次撤销已停止")


def _rollback_batch_refs(db: Session, kind: str, refs: dict[str, Any]) -> None:
    model = _model_for_kind(kind)
    for update in reversed(refs.get("updates") or []):
        row = db.get(model, int(update["id"]))
        if row is not None:
            if update.get("before_encrypted"):
                _restore_snapshot(row, _decrypt_rollback_snapshot(update))
            else:
                _restore_snapshot(row, update["before"])
    inserted_ids = [int(value) for value in refs.get("inserted_ids") or []]
    if inserted_ids:
        db.execute(delete(model).where(model.id.in_(inserted_ids)))
    db.flush()


def _model_for_kind(kind: str) -> type[Any]:
    return {
        "colleges": College,
        "majors": Major,
        "classes": AcademicClass,
        "students": Student,
        "outcomes": Student,
    }[kind]


def _claim_resource_kind(kind: str) -> str:
    return {
        "colleges": "college",
        "majors": "major",
        "classes": "class",
        "students": "student",
        "outcomes": "student",
    }[kind]


def _claim_field_scope(kind: str) -> str:
    return {
        "colleges": "record",
        "majors": "record",
        "classes": "record",
        "students": "student_profile",
        "outcomes": "student_outcome",
    }[kind]


def _model_for_claim_resource(resource_kind: str) -> type[Any]:
    return {
        "college": College,
        "major": Major,
        "class": AcademicClass,
        "student": Student,
    }[resource_kind]


def _inserted_snapshot_fields(kind: str) -> tuple[str, ...]:
    return {
        "colleges": ("code", "name"),
        "majors": ("college_id", "code", "name", "level"),
        "classes": ("college_id", "major_id", "cohort_year", "code", "name"),
        "students": (
            "student_no",
            "name",
            "gender",
            "college_id",
            "major_id",
            "direction_id",
            "class_id",
            "degree_level",
            "cohort_year",
            "origin_place",
            "contact_encrypted",
            "status",
            "status_occurred_at",
            "risk_level",
            "risk_reason",
            "is_key_group",
        ),
        "outcomes": ("status", "status_occurred_at", "risk_level", "risk_reason"),
    }[kind]


def _record_run_blocked(
    db: Session,
    run_id: int,
    *,
    actor: str,
    actor_id: int | None,
    actor_role: str,
    source_channel: str,
    action: str,
    reason: str,
    attempted_revision: int,
) -> None:
    run = db.get(ImportRun, run_id)
    if run is None:
        return
    _record_run_audit(
        db,
        run,
        actor=actor,
        actor_id=actor_id,
        actor_role=actor_role,
        source_channel=source_channel,
        action=action,
        change_summary={
            "reason": reason,
            "attempted_revision": attempted_revision,
            "before_status": run.status,
            "after_status": run.status,
            "before_revision": run.revision,
            "after_revision": run.revision,
        },
    )
    db.commit()


def build_import_run_receipt(db: Session, run: ImportRun) -> bytes:
    """Render from the caller's already locked transaction.

    The HTTP boundary locks ``ImportRun``, records the export event, renders this
    workbook and commits only after rendering succeeds.  Keeping all four steps
    in one transaction makes the status/revision shown in the audit event match
    the downloaded workbook and prevents a false "exported" event on failure.
    """

    snapshot_run = get_import_run(db, run.id)
    payload = import_run_payload(
        db,
        snapshot_run,
        lock=False,
        include_evidence=False,
    )
    return _render_import_run_receipt(db, snapshot_run, payload)


def _render_import_run_receipt(
    db: Session,
    run: ImportRun,
    payload: dict[str, Any],
) -> bytes:
    workbook = Workbook(write_only=True)
    summary_sheet = workbook.create_sheet("对账")
    _prepare_receipt_sheet(summary_sheet, [22, 12, 12, 12, 12, 12, 14])
    _append_receipt_header(
        summary_sheet,
        ["资料类型", "源行", "新增", "更新", "无变化", "问题行", "状态"],
    )
    for batch in payload["batches"]:
        _append_receipt_row(
            summary_sheet,
            [
                batch["title"],
                batch["row_count"],
                batch["inserted_count"],
                batch["updated_count"],
                batch["unchanged_count"],
                batch["error_row_count"],
                batch["status"],
            ],
        )
    issue_sheet = workbook.create_sheet("问题")
    _prepare_receipt_sheet(issue_sheet, [18, 12, 10, 18, 42, 24, 42])
    _append_receipt_header(
        issue_sheet,
        ["资料类型", "工作表", "行号", "列", "问题", "当前内容", "修改建议"],
    )
    issue_statement = text(
        """
        SELECT ib.target_kind AS kind, issue.value AS payload
        FROM import_batches AS ib
        CROSS JOIN LATERAL jsonb_array_elements(ib.errors)
            WITH ORDINALITY AS issue(value, ordinality)
        WHERE ib.run_id = :run_id AND ib.status <> 'superseded'
        ORDER BY ib.id, issue.ordinality
        """
    ).execution_options(stream_results=True, yield_per=500)
    for row in db.execute(issue_statement, {"run_id": run.id}).mappings():
        item = dict(row["payload"])
        _append_receipt_row(
            issue_sheet,
            [
                T16_TITLES.get(str(row["kind"]), str(row["kind"] or "")),
                item.get("sheet"),
                item.get("row_number"),
                item.get("column") or item.get("field"),
                item.get("reason"),
                item.get("masked_value") or item.get("value"),
                item.get("suggestion"),
            ],
        )
    evidence_sheet = workbook.create_sheet("证据轨")
    _prepare_receipt_sheet(evidence_sheet, [24, 24, 16, 22, 18, 24, 72])
    _append_receipt_header(
        evidence_sheet,
        ["时间", "动作", "当时角色", "操作者", "操作来源", "追溯号", "说明"],
    )
    evidence_statement = text(
        """
        SELECT id, action, actor_name, actor_role, source_channel,
               correlation_id, at, change_summary
        FROM audit_logs
        WHERE object_type = 'import_run' AND object_id = :object_id
        ORDER BY id
        """
    ).execution_options(stream_results=True, yield_per=500)
    for item in db.execute(
        evidence_statement,
        {"object_id": str(run.id)},
    ).mappings():
        _append_receipt_row(
            evidence_sheet,
            [
                item["at"].isoformat(),
                IMPORT_EVIDENCE_LABELS.get(item["action"], item["action"]),
                item["actor_role"],
                item["actor_name"],
                item["source_channel"],
                item["correlation_id"],
                _evidence_detail(item["change_summary"]),
            ],
        )
    history_sheet = workbook.create_sheet("文件版本链")
    _prepare_receipt_sheet(history_sheet, [10, 16, 14, 24, 14, 68, 12, 12, 22, 24])
    _append_receipt_header(
        history_sheet,
        [
            "批次",
            "资料类型",
            "状态",
            "文件名",
            "模板版本",
            "SHA256",
            "替换批次",
            "数据行",
            "上传人",
            "上传时间",
        ],
    )
    history_statement = text(
        """
        SELECT id, target_kind, status, original_filename, template_version,
               source_sha256, supersedes_batch_id, row_count, created_by, created_at
        FROM import_batches
        WHERE run_id = :run_id
        ORDER BY id
        """
    ).execution_options(stream_results=True, yield_per=500)
    for item in db.execute(history_statement, {"run_id": run.id}).mappings():
        _append_receipt_row(
            history_sheet,
            [
                item["id"],
                T16_TITLES.get(str(item["target_kind"]), str(item["target_kind"])),
                item["status"],
                item["original_filename"],
                item["template_version"],
                item["source_sha256"],
                item["supersedes_batch_id"],
                item["row_count"],
                item["created_by"],
                item["created_at"].isoformat() if item["created_at"] else None,
            ],
        )
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _prepare_receipt_sheet(sheet: Any, widths: list[int]) -> None:
    sheet.freeze_panes = "A2"
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _append_receipt_header(sheet: Any, values: list[Any]) -> None:
    cells: list[WriteOnlyCell] = []
    for value in values:
        cell = WriteOnlyCell(sheet, value=_excel_safe(value))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="335CFF")
        cells.append(cell)
    sheet.append(cells)


def _append_receipt_row(sheet: Any, values: list[Any]) -> None:
    sheet.append([_excel_safe(value) for value in values])


def _ensure_t16_kind(kind: str) -> None:
    if kind not in T16_KIND_ORDER:
        raise ValueError("未知的正式启用资料类型")


def _issue(
    row_number: int,
    code: str,
    reason: str,
    *,
    column: str | None = None,
    masked_value: Any = None,
    suggestion: str | None = None,
    editable: bool | None = None,
) -> dict[str, Any]:
    is_editable = row_number >= 2 and code not in {
        "template_version_missing",
        "template_version_unsupported",
        "template_kind_mismatch",
        "data_sheet_missing",
        "missing_header",
        "duplicate_header",
        "unknown_header",
        "row_limit_exceeded",
        "formula_not_allowed",
        "data_sheet_hidden",
        "hidden_data_row",
        "hidden_data_column",
        "excel_error_cell",
        "numeric_text_unsafe",
    }
    label = _label(column) if column else None
    masked = _mask_value(column, masked_value) if masked_value is not None else None
    return {
        "sheet": "数据",
        "row_number": row_number,
        "column": label,
        "field": label,
        "code": code,
        "reason": reason,
        "masked_value": masked,
        "value": masked,
        "suggestion": suggestion or "请核对当前模板后重新检查",
        "editable": is_editable if editable is None else editable,
    }


def _require(
    errors: list[dict[str, Any]],
    row_number: int,
    field: str,
    value: Any,
) -> None:
    if value in (None, ""):
        errors.append(
            _issue(
                row_number,
                "missing_required",
                f"{_label(field)}不能为空",
                column=field,
                suggestion=f"请填写{_label(field)}",
            )
        )


def _identifier(
    errors: list[dict[str, Any]],
    row_number: int,
    field: str,
    value: str,
) -> None:
    pattern = IDENTIFIER_PATTERNS.get(field)
    if value and pattern is not None and pattern.fullmatch(value) is None:
        errors.append(
            _issue(
                row_number,
                "identifier_format",
                f"{_label(field)}格式不正确",
                column=field,
                masked_value=value,
                suggestion="仅使用字母、数字、短横线或下划线，并保持文本格式",
            )
        )


def _text_quality(
    errors: list[dict[str, Any]],
    row_number: int,
    field: str,
    value: str | None,
    max_length: int,
) -> None:
    if not value:
        return
    if len(value) > max_length or any(ord(char) < 32 and char not in "\t" for char in value):
        errors.append(
            _issue(
                row_number,
                "value_length_or_control_char",
                f"{_label(field)}过长或包含控制字符",
                column=field,
                masked_value=value,
                suggestion=f"请删除换行或控制字符，并控制在 {max_length} 个字符内",
            )
        )


def _enum(
    errors: list[dict[str, Any]],
    row_number: int,
    field: str,
    value: str,
    allowed: set[str],
) -> None:
    if value and value not in allowed:
        errors.append(
            _issue(
                row_number,
                "enum_out_of_range",
                f"{_label(field)}不在允许范围内",
                column=field,
                masked_value=value,
                suggestion=f"允许值：{'、'.join(sorted(allowed))}",
            )
        )


def _label(field: str | None) -> str:
    return T16_LABELS.get(field or "", field or "")


def _values(enum_cls: Any) -> set[str]:
    return {item.value for item in enum_cls}


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _optional_text(value: Any) -> str | None:
    cleaned = _text(value)
    return cleaned or None


def _cohort(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        parsed = value
    elif isinstance(value, float):
        if not value.is_integer():
            return None
        parsed = int(value)
    else:
        text_value = _text(value)
        if re.fullmatch(r"\d{4}", text_value) is None:
            return None
        parsed = int(text_value)
    return parsed if 2000 <= parsed <= 2100 else None


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(_text(value))
    except ValueError:
        return None


def _protect_source_rows(kind: str, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Encrypt contact values before draft rows are persisted in JSONB.

    The staged import tables are durable database state, so the final Student
    column encryption rule also applies before confirmation.  Other fields keep
    their original JSON types to preserve validation and online correction.
    """

    protected = [dict(row) for row in rows]
    if kind != "students":
        return protected
    for row in protected:
        value = row.get("contact")
        if value in (None, "") or _is_protected_staging_value(value):
            continue
        ciphertext = encrypt_field(json.dumps(value, ensure_ascii=False, separators=(",", ":")))
        if ciphertext is None:
            raise ValueError("联系方式暂存加密失败")
        row["contact"] = {_ENCRYPTED_STAGING_KEY: ciphertext}
    return protected


def _unprotect_source_rows(kind: str, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    restored = [dict(row) for row in rows]
    if kind != "students":
        return restored
    for row in restored:
        value = row.get("contact")
        if not _is_protected_staging_value(value):
            # Backward-compatible read for a draft created before this guard;
            # the next validation/correction rewrites it encrypted.
            continue
        row["contact"] = _decrypt_staging_value(value)
    return restored


def _protect_staged_rows(kind: str, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if kind != "students":
        return rows
    protected: list[dict[str, Any]] = []
    for item in rows:
        copied = dict(item)
        data = copied.get("data")
        if isinstance(data, dict):
            copied["data"] = _protect_source_rows(kind, [data])[0]
        protected.append(copied)
    return protected


def _unprotect_staged_rows(kind: str, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    if kind != "students":
        return rows
    restored: list[dict[str, Any]] = []
    for item in rows:
        copied = dict(item)
        data = copied.get("data")
        if isinstance(data, dict):
            copied["data"] = _unprotect_source_rows(kind, [data])[0]
        restored.append(copied)
    return restored


def _is_protected_staging_value(value: Any) -> bool:
    return isinstance(value, dict) and isinstance(value.get(_ENCRYPTED_STAGING_KEY), str)


def _decrypt_staging_value(value: dict[str, Any]) -> Any:
    plaintext = decrypt_field(str(value[_ENCRYPTED_STAGING_KEY]))
    if plaintext is None:
        raise ImportConflictError("联系方式暂存数据无法解密，请重新上传学生模板")
    try:
        return json.loads(plaintext)
    except json.JSONDecodeError as exc:
        raise ImportConflictError("联系方式暂存数据已损坏，请重新上传学生模板") from exc


def _clean_cell(value: Any) -> Any:
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def _bounded_text(
    value: str | None,
    max_length: int,
    *,
    field: str,
    required: bool,
) -> str | None:
    cleaned = (value or "").strip()
    if required and not cleaned:
        raise ValueError(f"{field}不能为空")
    if len(cleaned) > max_length:
        raise ValueError(f"{field}不能超过 {max_length} 个字符")
    if any(ord(char) < 32 and char not in "\t" for char in cleaned):
        raise ValueError(f"{field}包含不支持的控制字符")
    return cleaned or None


def _mask_value(field: str | None, value: Any) -> str | None:
    if value in (None, ""):
        return None
    text = _text(value)
    if field == "contact":
        if "@" in text:
            local, _, domain = text.partition("@")
            return f"{local[:1] or '*'}***@{domain}"
        return f"{text[:3]}****{text[-4:]}" if len(text) >= 7 else "****"
    if field in SENSITIVE_COLUMNS:
        if len(text) <= 1:
            return "*"
        return f"{text[0]}{'*' * min(max(len(text) - 1, 1), 4)}"
    return text[:128]


def _empty_summary() -> dict[str, int]:
    return {"source": 0, "inserted": 0, "updated": 0, "unchanged": 0, "errors": 0}


def _normalized_summary(value: dict[str, Any] | None) -> dict[str, int]:
    source = value or {}
    return {key: int(source.get(key, 0)) for key in _empty_summary()}


def _aggregate_summaries(batches: list[ImportBatch]) -> dict[str, int]:
    total = _empty_summary()
    for batch in batches:
        summary = _normalized_summary(batch.result_summary)
        for key in total:
            total[key] += summary[key]
    return total


def _business_counts(db: Session) -> dict[str, int]:
    return {
        "colleges": int(db.scalar(select(func.count()).select_from(College)) or 0),
        "majors": int(db.scalar(select(func.count()).select_from(Major)) or 0),
        "classes": int(db.scalar(select(func.count()).select_from(AcademicClass)) or 0),
        "students": int(db.scalar(select(func.count()).select_from(Student)) or 0),
        "outcomes": int(
            db.scalar(
                select(func.count())
                .select_from(Student)
                .where(Student.status != EmploymentStatus.UNSETTLED.value)
            )
            or 0
        ),
    }


def _snapshot(model: Any, fields: tuple[str, ...]) -> dict[str, Any]:
    return {field: _serialize(getattr(model, field)) for field in fields}


def _snapshot_fingerprint(model: Any, fields: tuple[str, ...]) -> str:
    return _dict_fingerprint(_snapshot(model, fields))


def _dict_fingerprint(value: dict[str, Any]) -> str:
    return _json_fingerprint(value)


def _json_fingerprint(value: Any) -> str:
    canonical = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return sha256(canonical.encode("utf-8")).hexdigest()


def _decrypt_rollback_snapshot(update: dict[str, Any]) -> dict[str, Any]:
    decrypted = decrypt_field(str(update.get("before_encrypted") or ""))
    if decrypted is None:
        raise ImportConflictError("敏感回滚快照无法解密，为保护原始事实，本次撤销已停止")
    try:
        payload = json.loads(decrypted)
    except json.JSONDecodeError as exc:
        raise ImportConflictError("敏感回滚快照已损坏，为保护原始事实，本次撤销已停止") from exc
    if not isinstance(payload, dict):
        raise ImportConflictError("敏感回滚快照格式异常，为保护原始事实，本次撤销已停止")
    return payload


def _serialize(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def _restore_snapshot(model: Any, snapshot: dict[str, Any]) -> None:
    datetime_fields = {"status_occurred_at"}
    for field, value in snapshot.items():
        if field in datetime_fields and isinstance(value, str):
            setattr(model, field, datetime.fromisoformat(value))
        else:
            setattr(model, field, value)


def _evidence_detail(summary: dict[str, Any] | None) -> str | None:
    if not summary:
        return None
    allowed = (
        "cohort_year",
        "batch_id",
        "batch_ids",
        "kind",
        "template_version",
        "source_sha256",
        "supersedes_batch_id",
        "source_rows",
        "file_errors",
        "row_number",
        "column",
        "before",
        "after",
        "reason",
        "trigger",
        "approval_method",
        "approved_revision",
        "approved_plan_sha256",
        "validation_status",
        "validation_noop",
        "validation_plan_sha256",
        "attempted_revision",
        "source",
        "inserted",
        "updated",
        "unchanged",
        "errors",
        "status",
        "revision",
        "before_status",
        "after_status",
        "before_revision",
        "after_revision",
    )
    parts = [f"{key}={summary[key]}" for key in allowed if key in summary]
    changes = summary.get("batch_changes")
    if isinstance(changes, list):
        compact = ",".join(
            (
                f"{item.get('kind')}#{item.get('batch_id')}:"
                f"+{item.get('inserted_count', 0)}/~{item.get('updated_count', 0)}"
            )
            for item in changes
            if isinstance(item, dict)
        )
        if compact:
            parts.append(f"batch_changes={compact}")
    return " · ".join(parts) or None


def _evidence_batch_ids(summary: dict[str, Any] | None) -> list[int]:
    if not isinstance(summary, dict):
        return []
    values: set[int] = set()
    if isinstance(summary.get("batch_id"), int):
        values.add(int(summary["batch_id"]))
    for value in summary.get("batch_ids") or []:
        if isinstance(value, int):
            values.add(value)
    for item in summary.get("batch_changes") or []:
        if isinstance(item, dict) and isinstance(item.get("batch_id"), int):
            values.add(int(item["batch_id"]))
    return sorted(values)


def _excel_safe(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value
