"""T8 M5 recruitment fair service."""

from __future__ import annotations

import hashlib
import hmac
import json
import secrets
from datetime import UTC, datetime
from io import BytesIO
from typing import Any

from docx import Document
from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from app.core.audit import record_db
from app.core.config import settings
from app.core.security import hash_password, verify_password
from app.models import (
    EnterpriseAccessToken,
    Fair,
    FairCheckin,
    FairCompany,
    FairEnterpriseCheckin,
    FairSnapshot,
    Job,
    JobAction,
    Match,
    OperationReceipt,
    Student,
)
from app.models.enums import FairStatus, JobActionType, JobSource

FAIR_COMPANY_SUBMITTED = "已报名"
FAIR_COMPANY_APPROVED = "已通过"
FAIR_COMPANY_REJECTED = "已拒绝"
FAIR_TRANSITIONS = {
    FairStatus.PLANNED.value: FairStatus.REGISTRATION.value,
    FairStatus.REGISTRATION.value: FairStatus.CHECKIN.value,
    FairStatus.CHECKIN.value: FairStatus.ENDED.value,
    FairStatus.ENDED.value: FairStatus.ARCHIVED.value,
}
FAIR_COMPANY_MUTABLE_STATES = {
    FairStatus.PLANNED.value,
    FairStatus.REGISTRATION.value,
}


class FairConflictError(RuntimeError):
    """The requested action conflicts with the fair lifecycle."""


class StudentOutsideCollegeError(PermissionError):
    """The counselor tried to check in a student from another college."""


def create_fair(
    db: Session,
    *,
    name: str,
    fair_type: str,
    held_at: datetime | None,
    location: str | None,
    actor: str,
    idempotency_key: str,
) -> Fair:
    normalized = {
        "name": name.strip(),
        "fair_type": fair_type.strip(),
        "held_at": held_at.isoformat() if held_at else None,
        "location": location.strip() if location and location.strip() else None,
    }
    request_hash = _operation_hash(normalized)
    _acquire_operation_lock(db, "fair.create", idempotency_key)
    existing = db.execute(
        select(Fair).where(Fair.idempotency_key == idempotency_key)
    ).scalar_one_or_none()
    if existing is not None:
        if existing.request_hash != request_hash:
            raise FairConflictError("该操作标识已用于另一份招聘会创建内容")
        db.commit()
        return existing
    fair = Fair(
        name=normalized["name"],
        fair_type=normalized["fair_type"],
        held_at=held_at,
        location=normalized["location"],
        status=FairStatus.PLANNED.value,
        idempotency_key=idempotency_key,
        request_hash=request_hash,
        created_by=actor,
    )
    db.add(fair)
    db.flush()
    record_db(
        db,
        actor=actor,
        action="fair.create",
        object_type="fair",
        object_id=fair.id,
        change_summary={"after": _fair_dict(fair)},
        correlation_id=f"FAIR-CREATE-{request_hash[:24]}",
    )
    db.commit()
    db.refresh(fair)
    return fair


def list_fairs(db: Session) -> list[Fair]:
    return list(db.execute(select(Fair).order_by(Fair.held_at.desc().nullslast(), Fair.id.desc())).scalars())


def update_fair(
    db: Session,
    fair_id: int,
    *,
    actor: str,
    updates: dict[str, Any],
) -> Fair:
    fair = _get_fair(db, fair_id, lock=True)
    _require_fair_status(
        fair,
        {FairStatus.PLANNED.value},
        action="编辑基本信息",
    )
    before = _fair_dict(fair)
    if "fair_type" in updates and updates["fair_type"] is None:
        raise ValueError("fair_type cannot be null")
    for field in ("name", "fair_type", "held_at", "location"):
        if field in updates:
            setattr(fair, field, updates[field])
    if before == _fair_dict(fair):
        db.commit()
        return fair
    record_db(
        db,
        actor=actor,
        action="fair.update",
        object_type="fair",
        object_id=fair.id,
        change_summary={"before": before, "after": _fair_dict(fair)},
    )
    db.commit()
    db.refresh(fair)
    return fair


def transition_fair(
    db: Session,
    fair_id: int,
    *,
    status: str,
    actor: str,
) -> Fair:
    fair = _get_fair(db, fair_id, lock=True)
    if status == fair.status:
        db.commit()
        return fair
    expected = FAIR_TRANSITIONS.get(fair.status)
    if expected != status:
        raise FairConflictError(f"场次当前为“{fair.status}”，下一状态只能是“{expected or '无'}”")
    before = fair.status
    now = datetime.now(UTC)
    fair.status = status
    if status == FairStatus.REGISTRATION.value:
        fair.registration_opens_at = now
    elif status == FairStatus.CHECKIN.value:
        fair.registration_closes_at = now
        fair.checkin_started_at = now
    elif status == FairStatus.ENDED.value:
        fair.checkin_ended_at = now
        freeze_snapshot(db, fair, actor=actor)
    elif status == FairStatus.ARCHIVED.value:
        fair.archived_at = now
    record_db(
        db,
        actor=actor,
        action="fair.status.transition",
        object_type="fair",
        object_id=fair.id,
        change_summary={"before": before, "after": status},
    )
    db.commit()
    db.refresh(fair)
    return fair


def delete_fair(db: Session, fair_id: int, *, actor: str) -> None:
    fair = _get_fair(db, fair_id, lock=True)
    _require_fair_status(
        fair,
        {FairStatus.PLANNED.value},
        action="删除场次",
    )
    record_db(
        db,
        actor=actor,
        action="fair.delete",
        object_type="fair",
        object_id=fair.id,
        change_summary={"before": _fair_dict(fair)},
    )
    db.delete(fair)
    db.commit()


def register_company(
    db: Session,
    fair_id: int,
    *,
    company: str,
    title: str,
    raw_text: str,
    booth: str | None,
    actor: str,
    idempotency_key: str,
    enterprise_organization_id: int | None = None,
) -> FairCompany:
    normalized = {
        "fair_id": fair_id,
        "company": company.strip(),
        "title": title.strip(),
        "raw_text": raw_text.strip(),
        "booth": booth.strip() if booth and booth.strip() else None,
        "enterprise_organization_id": enterprise_organization_id,
    }
    request_hash = _operation_hash(normalized)
    _acquire_operation_lock(db, "fair_company.register", idempotency_key)
    existing = db.execute(
        select(FairCompany).where(FairCompany.idempotency_key == idempotency_key)
    ).scalar_one_or_none()
    if existing is not None:
        if existing.request_hash != request_hash:
            raise FairConflictError("该操作标识已用于另一份企业报名内容")
        db.commit()
        return existing
    fair = _get_fair(db, fair_id, lock=True)
    _require_fair_status(fair, FAIR_COMPANY_MUTABLE_STATES, action="登记企业报名")
    job = Job(
        enterprise_organization_id=enterprise_organization_id,
        company=normalized["company"],
        title=normalized["title"],
        raw_text=normalized["raw_text"],
        source=JobSource.FORM.value,
        received_at=datetime.now(UTC),
        created_by=actor,
    )
    db.add(job)
    db.flush()
    row = FairCompany(
        fair_id=fair_id,
        enterprise_organization_id=enterprise_organization_id,
        company=normalized["company"],
        booth=normalized["booth"],
        status=FAIR_COMPANY_SUBMITTED,
        title=normalized["title"],
        raw_text=normalized["raw_text"],
        job_id=job.id,
        idempotency_key=idempotency_key,
        request_hash=request_hash,
        created_by=actor,
    )
    db.add(row)
    db.flush()
    record_db(
        db,
        actor=actor,
        action="fair_company.register",
        object_type="fair_company",
        object_id=row.id,
        change_summary={"fair_id": fair_id, "job_id": job.id, "status": row.status},
        correlation_id=f"FAIR-COMPANY-{request_hash[:24]}",
    )
    db.commit()
    db.refresh(row)
    return row


def list_companies(db: Session, fair_id: int) -> list[FairCompany]:
    _get_fair(db, fair_id)
    return list(
        db.execute(
            select(FairCompany)
            .where(FairCompany.fair_id == fair_id)
            .order_by(FairCompany.id)
        ).scalars()
    )


def review_company(
    db: Session,
    fair_id: int,
    company_id: int,
    *,
    status: str,
    booth: str | None,
    actor: str,
) -> FairCompany:
    if status not in {FAIR_COMPANY_APPROVED, FAIR_COMPANY_REJECTED}:
        raise ValueError("invalid review status")
    fair = _get_fair(db, fair_id, lock=True)
    _require_fair_status(fair, FAIR_COMPANY_MUTABLE_STATES, action="审核企业报名")
    row = db.execute(
        select(FairCompany)
        .where(FairCompany.id == company_id)
        .with_for_update()
    ).scalar_one_or_none()
    if row is None or row.fair_id != fair_id:
        raise ValueError("fair company not found")
    if row.status != FAIR_COMPANY_SUBMITTED:
        if row.status == status and (booth is None or row.booth == booth):
            db.commit()
            return row
        raise FairConflictError(f"企业报名已经处理为“{row.status}”，不能重复审核")
    before = _company_dict(row)
    row.status = status
    if booth is not None:
        row.booth = booth
    record_db(
        db,
        actor=actor,
        action="fair_company.review",
        object_type="fair_company",
        object_id=row.id,
        change_summary={"before": before, "after": _company_dict(row)},
    )
    db.commit()
    db.refresh(row)
    return row


def checkin_student(
    db: Session,
    fair_id: int,
    *,
    student_no: str,
    gate_id: str | None,
    actor: str,
    college_id: int | None = None,
    source: str = "校方登记",
    idempotency_key: str | None = None,
) -> FairCheckin:
    fair = _get_fair(db, fair_id, lock=True)
    student = db.execute(
        select(Student).where(Student.student_no == student_no)
    ).scalar_one_or_none()
    if student is None:
        raise ValueError("student not found")
    if college_id is not None and student.college_id != college_id:
        raise StudentOutsideCollegeError("只能为本学院学生办理签到。")
    if idempotency_key:
        keyed = db.execute(
            select(FairCheckin).where(FairCheckin.idempotency_key == idempotency_key)
        ).scalar_one_or_none()
        if keyed is not None:
            if keyed.fair_id != fair_id or keyed.student_id != student.id:
                raise FairConflictError("该操作标识已用于其他签到")
            db.commit()
            return keyed
    _require_checkin_open(fair)
    existing = db.execute(
        select(FairCheckin).where(
            FairCheckin.fair_id == fair_id,
            FairCheckin.student_id == student.id,
        )
    ).scalar_one_or_none()
    if existing is not None:
        db.commit()
        return existing
    row = FairCheckin(
        fair_id=fair_id,
        student_id=student.id,
        gate_id=gate_id,
        at=datetime.now(UTC),
        source=source,
        idempotency_key=idempotency_key,
        created_by=actor,
    )
    db.add(row)
    db.flush()
    record_db(
        db,
        actor=actor,
        action="fair_checkin.create",
        object_type="fair_checkin",
        object_id=row.id,
        change_summary={
            "fair_id": fair_id,
            "student_id": student.id,
            "gate_id": gate_id,
            "source": source,
        },
    )
    db.commit()
    db.refresh(row)
    return row


def list_student_fairs(db: Session, *, student_id: int) -> list[dict[str, Any]]:
    rows = db.execute(
        select(Fair, FairCheckin)
        .outerjoin(
            FairCheckin,
            (FairCheckin.fair_id == Fair.id) & (FairCheckin.student_id == student_id),
        )
        .where(Fair.status != FairStatus.PLANNED.value)
        .order_by(Fair.held_at.desc().nullslast(), Fair.id.desc())
    ).all()
    job_rows = db.execute(
        select(
            FairCompany.fair_id,
            Job.id,
            Match.id,
            Job.company,
            Job.title,
        )
        .join(Job, Job.id == FairCompany.job_id)
        .join(
            Match,
            (Match.job_id == Job.id)
            & (Match.student_id == student_id)
            & (Match.pushed_at.is_not(None)),
        )
        .where(FairCompany.status == FAIR_COMPANY_APPROVED)
        .order_by(FairCompany.fair_id, Job.id)
    ).all()
    jobs_by_fair: dict[int, list[dict[str, Any]]] = {}
    for fair_id, job_id, match_id, company, title in job_rows:
        jobs_by_fair.setdefault(fair_id, []).append(
            {
                "job_id": job_id,
                "match_id": match_id,
                "company": company,
                "title": title,
            }
        )
    return [
        {
            "fair_id": fair.id,
            "name": fair.name,
            "fair_type": fair.fair_type,
            "held_at": fair.held_at,
            "location": fair.location,
            "status": fair.status,
            "can_checkin": fair.status == FairStatus.CHECKIN.value and checkin is None,
            "checked_in": checkin is not None,
            "checked_in_at": checkin.at if checkin else None,
            "checkin_id": checkin.id if checkin else None,
            "jobs": jobs_by_fair.get(fair.id, []),
        }
        for fair, checkin in rows
    ]


def checkin_student_self(
    db: Session,
    fair_id: int,
    *,
    student_id: int,
    gate_id: str | None,
    idempotency_key: str,
    checkin_code: str | None = None,
    actor: str,
) -> tuple[FairCheckin, bool]:
    # Keep idempotency and per-student uniqueness deterministic without
    # serializing the whole fair.  The second identical request observes and
    # returns the committed winner instead of leaking a unique-constraint 500.
    _acquire_operation_lock(
        db,
        "fair.student_checkin.idempotency",
        idempotency_key,
    )
    _acquire_operation_lock(
        db,
        "fair.student_checkin.student",
        f"{fair_id}:{student_id}",
    )
    # Concurrent students may share the fair row while verifying the current
    # code.  FOR SHARE still blocks code rotation/status mutation (FOR UPDATE),
    # but avoids serializing every bcrypt verification on one exclusive lock.
    fair = _get_fair(db, fair_id, shared_lock=True)
    keyed = db.execute(
        select(FairCheckin).where(FairCheckin.idempotency_key == idempotency_key)
    ).scalar_one_or_none()
    if keyed is not None:
        if keyed.fair_id != fair_id or keyed.student_id != student_id:
            raise FairConflictError("该操作标识已用于其他签到")
        db.commit()
        return keyed, True
    _require_checkin_open(fair)
    if settings.student_production_auth_enabled:
        if not checkin_code or not _verify_student_checkin_code(fair, checkin_code):
            raise FairConflictError("场次码无效或已更新，请向现场工作人员核对")
    existing = db.execute(
        select(FairCheckin).where(
            FairCheckin.fair_id == fair_id,
            FairCheckin.student_id == student_id,
        )
    ).scalar_one_or_none()
    if existing is not None:
        db.commit()
        return existing, True
    student = db.get(Student, student_id)
    if student is None:
        raise ValueError("student not found")
    row = FairCheckin(
        fair_id=fair_id,
        student_id=student_id,
        gate_id=gate_id,
        at=datetime.now(UTC),
        source="学生自助",
        idempotency_key=idempotency_key,
        created_by=actor,
    )
    db.add(row)
    db.flush()
    record_db(
        db,
        actor=actor,
        action="fair_checkin.student_self",
        object_type="fair_checkin",
        object_id=row.id,
        change_summary={"fair_id": fair_id, "source": row.source},
    )
    db.commit()
    db.refresh(row)
    return row, False


def _student_checkin_code_digest(value: str) -> str:
    key = (settings.field_encryption_key or settings.jwt_secret).encode("utf-8")
    return hmac.new(
        key,
        f"fair-student-checkin:{value}".encode(),
        hashlib.sha256,
    ).hexdigest()


def _verify_student_checkin_code(fair: Fair, value: str) -> bool:
    normalized = value.strip()
    return bool(
        fair.student_checkin_code_digest
        and fair.student_checkin_code_hash
        and hmac.compare_digest(
            fair.student_checkin_code_digest,
            _student_checkin_code_digest(normalized),
        )
        and verify_password(normalized, fair.student_checkin_code_hash)
    )


def issue_student_checkin_code(
    db: Session,
    fair_id: int,
    *,
    reason: str,
    idempotency_key: str,
    expected_issued_at: datetime | None,
    actor: str,
    actor_id: int,
    actor_role: str,
) -> tuple[str, datetime]:
    """Rotate a fair code and return its plaintext exactly once."""
    request_hash = _operation_hash(
        {
            "fair_id": fair_id,
            "reason": reason.strip(),
            "expected_issued_at": (
                expected_issued_at.isoformat() if expected_issued_at else None
            ),
        }
    )
    _acquire_operation_lock(db, "fair.student_checkin_code", idempotency_key)
    receipt = db.execute(
        select(OperationReceipt).where(
            OperationReceipt.namespace == "fair.student_checkin_code",
            OperationReceipt.idempotency_key == idempotency_key,
        )
    ).scalar_one_or_none()
    if receipt is not None:
        if receipt.request_hash != request_hash:
            raise FairConflictError("该操作标识已用于另一份场次码签发内容")
        raise FairConflictError(
            "场次码已经签发，但一次性原文不会再次返回；请使用新的操作标识重新生成"
        )
    # Prepare the one-time code while holding only its request-key lock.  Fair
    # row exclusivity is reserved for the short status recheck and assignment,
    # so live student check-ins do not queue behind bcrypt.
    # 128 bits keeps a handoff code far beyond the 32-bit value originally
    # shipped in T18 while retaining an uppercase, QR/clipboard-safe format.
    plain_code = f"QH-{secrets.token_hex(16).upper()}"
    code_digest = _student_checkin_code_digest(plain_code)
    code_hash = hash_password(plain_code)
    fair = _get_fair(db, fair_id, lock=True)
    _require_checkin_open(fair)
    if fair.student_checkin_code_issued_at != expected_issued_at:
        raise FairConflictError(
            "场次码已由另一窗口更新，请刷新招聘会后重新生成并交接最新场次码"
        )
    issued_at = datetime.now(UTC)
    fair.student_checkin_code_digest = code_digest
    fair.student_checkin_code_hash = code_hash
    fair.student_checkin_code_issued_at = issued_at
    db.add(
        OperationReceipt(
            namespace="fair.student_checkin_code",
            idempotency_key=idempotency_key,
            request_hash=request_hash,
            result_json={
                "fair_id": fair.id,
                "issued_at": issued_at.isoformat(),
                "plaintext_recorded": False,
            },
            created_by=actor,
        )
    )
    record_db(
        db,
        actor=actor,
        actor_id=actor_id,
        actor_role=actor_role,
        source_channel="management",
        correlation_id=f"FAIR-CODE-{fair.id}-{issued_at.strftime('%Y%m%d%H%M%S')}",
        action="fair.student_checkin_code.issue",
        object_type="fair",
        object_id=fair.id,
        change_summary={
            "reason": reason.strip(),
            "issued_at": issued_at.isoformat(),
            "plaintext_recorded": False,
        },
    )
    db.commit()
    return plain_code, issued_at


def list_enterprise_fairs(
    db: Session,
    *,
    access_token_id: int | None = None,
    organization_id: int | None = None,
) -> list[dict[str, Any]]:
    if (access_token_id is None) == (organization_id is None):
        raise ValueError("exactly one enterprise owner is required")
    token = db.get(EnterpriseAccessToken, access_token_id) if access_token_id else None
    if access_token_id is not None and token is None:
        raise ValueError("enterprise token not found")
    count_subquery = (
        select(
            FairEnterpriseCheckin.fair_company_id,
            func.count(FairEnterpriseCheckin.id).label("checkin_count"),
        )
        .group_by(FairEnterpriseCheckin.fair_company_id)
        .subquery()
    )
    statement = (
        select(Fair, FairCompany, func.coalesce(count_subquery.c.checkin_count, 0))
        .join(FairCompany, FairCompany.fair_id == Fair.id)
        .outerjoin(
            count_subquery,
            count_subquery.c.fair_company_id == FairCompany.id,
        )
        .order_by(Fair.held_at.desc().nullslast(), Fair.id.desc())
    )
    if organization_id is not None:
        statement = statement.where(FairCompany.enterprise_organization_id == organization_id)
    else:
        assert token is not None
        statement = statement.where(FairCompany.company == token.company_name)
    rows = db.execute(statement).all()
    return [
        {
            "fair_id": fair.id,
            "fair_company_id": company.id,
            "name": fair.name,
            "fair_type": fair.fair_type,
            "held_at": fair.held_at,
            "location": fair.location,
            "booth": company.booth,
            "status": fair.status,
            "registration_status": company.status,
            "can_checkin": (
                fair.status == FairStatus.CHECKIN.value
                and company.status == FAIR_COMPANY_APPROVED
            ),
            "checked_in_representatives": int(checkin_count),
        }
        for fair, company, checkin_count in rows
    ]


def checkin_enterprise(
    db: Session,
    fair_id: int,
    *,
    access_token_id: int | None = None,
    organization_id: int | None = None,
    representative_name: str,
    gate_id: str | None,
    idempotency_key: str,
    actor: str,
) -> tuple[FairEnterpriseCheckin, bool]:
    fair = _get_fair(db, fair_id, lock=True)
    if (access_token_id is None) == (organization_id is None):
        raise ValueError("exactly one enterprise owner is required")
    statement = select(FairCompany).where(FairCompany.fair_id == fair_id)
    if organization_id is not None:
        statement = statement.where(FairCompany.enterprise_organization_id == organization_id)
    else:
        token = db.get(EnterpriseAccessToken, access_token_id)
        if token is None:
            raise ValueError("enterprise token not found")
        statement = statement.where(FairCompany.company == token.company_name)
    company = db.execute(statement).scalar_one_or_none()
    if company is None:
        raise ValueError("企业未报名该场次")
    if company.status != FAIR_COMPANY_APPROVED:
        raise FairConflictError("企业报名尚未审核通过")
    keyed = db.execute(
        select(FairEnterpriseCheckin).where(
            FairEnterpriseCheckin.idempotency_key == idempotency_key
        )
    ).scalar_one_or_none()
    if keyed is not None:
        if keyed.fair_company_id != company.id:
            raise FairConflictError("该操作标识已用于其他签到")
        db.commit()
        return keyed, True
    _require_checkin_open(fair)
    existing = db.execute(
        select(FairEnterpriseCheckin).where(
            FairEnterpriseCheckin.fair_company_id == company.id,
            FairEnterpriseCheckin.representative_name == representative_name.strip(),
        )
    ).scalar_one_or_none()
    if existing is not None:
        db.commit()
        return existing, True
    row = FairEnterpriseCheckin(
        fair_company_id=company.id,
        representative_name=representative_name.strip(),
        gate_id=gate_id,
        at=datetime.now(UTC),
        idempotency_key=idempotency_key,
        created_by=actor,
    )
    db.add(row)
    db.flush()
    record_db(
        db,
        actor=actor,
        action="fair_checkin.enterprise",
        object_type="fair_enterprise_checkin",
        object_id=row.id,
        change_summary={"fair_id": fair_id, "fair_company_id": company.id},
    )
    db.commit()
    db.refresh(row)
    return row, False


def correct_student_checkin(
    db: Session,
    fair_id: int,
    *,
    student_no: str,
    gate_id: str | None,
    at: datetime,
    reason: str,
    actor: str,
) -> FairCheckin:
    fair = _get_fair(db, fair_id, lock=True)
    _require_fair_status(
        fair,
        {FairStatus.ENDED.value, FairStatus.ARCHIVED.value},
        action="审计补录或更正",
    )
    student = db.execute(
        select(Student).where(Student.student_no == student_no)
    ).scalar_one_or_none()
    if student is None:
        raise ValueError("student not found")
    row = db.execute(
        select(FairCheckin)
        .where(
            FairCheckin.fair_id == fair_id,
            FairCheckin.student_id == student.id,
        )
        .with_for_update()
    ).scalar_one_or_none()
    now = datetime.now(UTC)
    normalized_reason = reason.strip()
    normalized_at = at.astimezone(UTC)
    if (
        row is not None
        and row.at.astimezone(UTC) == normalized_at
        and row.gate_id == gate_id
        and row.correction_reason == normalized_reason
        and row.source == "审计补录"
    ):
        db.commit()
        return row
    if row is None:
        row = FairCheckin(
            fair_id=fair_id,
            student_id=student.id,
            gate_id=gate_id,
            at=normalized_at,
            source="审计补录",
            corrected_by=actor,
            correction_reason=normalized_reason,
            corrected_at=now,
            created_by=actor,
        )
        db.add(row)
        db.flush()
    else:
        row.at = normalized_at
        row.gate_id = gate_id
        row.source = "审计补录"
        row.corrected_by = actor
        row.correction_reason = normalized_reason
        row.corrected_at = now
        db.flush()
    record_db(
        db,
        actor=actor,
        action="fair_checkin.audit_correction",
        object_type="fair_checkin",
        object_id=row.id,
        change_summary={
            "fair_id": fair_id,
            "student_id": student.id,
            "at": normalized_at.isoformat(),
            "gate_id": gate_id,
            "reason": normalized_reason,
        },
    )
    freeze_snapshot(db, fair, actor=actor)
    db.commit()
    db.refresh(row)
    return row


def get_student_checkin_for_correction(
    db: Session,
    fair_id: int,
    *,
    student_no: str,
) -> FairCheckin:
    fair = _get_fair(db, fair_id)
    _require_fair_status(
        fair,
        {FairStatus.ENDED.value, FairStatus.ARCHIVED.value},
        action="核对补录结果",
    )
    return _find_student_checkin(db, fair_id, student_no=student_no)


def get_student_checkin(
    db: Session,
    fair_id: int,
    *,
    student_no: str,
    college_id: int | None = None,
) -> FairCheckin:
    """Return one committed check-in fact for uncertain-write reconciliation."""
    _get_fair(db, fair_id)
    return _find_student_checkin(
        db,
        fair_id,
        student_no=student_no,
        college_id=college_id,
    )


def _find_student_checkin(
    db: Session,
    fair_id: int,
    *,
    student_no: str,
    college_id: int | None = None,
) -> FairCheckin:
    student = db.execute(
        select(Student).where(Student.student_no == student_no.strip())
    ).scalar_one_or_none()
    if student is None:
        raise ValueError("student not found")
    if college_id is not None and student.college_id != college_id:
        raise StudentOutsideCollegeError("只能核对本学院学生的签到事实。")
    row = db.execute(
        select(FairCheckin).where(
            FairCheckin.fair_id == fair_id,
            FairCheckin.student_id == student.id,
        )
    ).scalar_one_or_none()
    if row is None:
        raise ValueError("fair checkin not found")
    return row


def freeze_snapshot(db: Session, fair: Fair, *, actor: str) -> FairSnapshot:
    now = datetime.now(UTC)
    metrics = _snapshot_metrics(db, fair.id)
    snapshot = db.execute(
        select(FairSnapshot)
        .where(FairSnapshot.fair_id == fair.id)
        .with_for_update()
    ).scalar_one_or_none()
    if snapshot is None:
        snapshot = FairSnapshot(
            fair_id=fair.id,
            frozen_at=now,
            metrics_json=metrics,
            created_by=actor,
        )
        db.add(snapshot)
    else:
        snapshot.frozen_at = now
        snapshot.metrics_json = metrics
    db.flush()
    return snapshot


def get_snapshot(db: Session, fair_id: int) -> dict[str, Any]:
    fair = _get_fair(db, fair_id)
    snapshot = db.execute(
        select(FairSnapshot).where(FairSnapshot.fair_id == fair_id)
    ).scalar_one_or_none()
    if snapshot is None:
        raise FairConflictError("场次尚未结束，暂无冻结快照")
    return {
        "fair_id": fair.id,
        "status": fair.status,
        **snapshot.metrics_json,
        "frozen_at": snapshot.frozen_at,
    }


def _snapshot_metrics(db: Session, fair_id: int) -> dict[str, int]:
    registered = int(db.scalar(
        select(func.count(FairCompany.id)).where(FairCompany.fair_id == fair_id)
    ) or 0)
    approved = int(db.scalar(
        select(func.count(FairCompany.id)).where(
            FairCompany.fair_id == fair_id,
            FairCompany.status == FAIR_COMPANY_APPROVED,
        )
    ) or 0)
    student_checkins = int(db.scalar(
        select(func.count(FairCheckin.id)).where(
            FairCheckin.fair_id == fair_id,
            FairCheckin.student_id.is_not(None),
        )
    ) or 0)
    enterprise_checkins = int(db.scalar(
        select(func.count(FairEnterpriseCheckin.id))
        .join(FairCompany, FairCompany.id == FairEnterpriseCheckin.fair_company_id)
        .where(FairCompany.fair_id == fair_id)
    ) or 0)
    return {
        "student_checkins": student_checkins,
        "enterprise_checkins": enterprise_checkins,
        "registered_companies": registered,
        "approved_companies": approved,
    }


def _require_checkin_open(fair: Fair) -> None:
    if fair.status != FairStatus.CHECKIN.value:
        raise FairConflictError(f"场次当前为“{fair.status}”，未开放签到")


def get_funnel(db: Session, fair_id: int) -> dict[str, int]:
    _get_fair(db, fair_id)
    registered = db.execute(
        select(func.count()).select_from(FairCompany).where(FairCompany.fair_id == fair_id)
    ).scalar_one()
    approved = db.execute(
        select(func.count())
        .select_from(FairCompany)
        .where(
            FairCompany.fair_id == fair_id,
            FairCompany.status == FAIR_COMPANY_APPROVED,
        )
    ).scalar_one()
    named_checkins = db.execute(
        select(func.count(func.distinct(FairCheckin.student_id))).where(
            FairCheckin.fair_id == fair_id,
            FairCheckin.student_id.is_not(None),
        )
    ).scalar_one()
    anonymous = db.execute(
        select(func.coalesce(func.sum(FairCheckin.anonymous_count), 0)).where(
            FairCheckin.fair_id == fair_id,
            FairCheckin.anonymous_count.is_not(None),
        )
    ).scalar_one()
    enterprise_checkins = db.execute(
        select(func.count(FairEnterpriseCheckin.id))
        .join(FairCompany, FairCompany.id == FairEnterpriseCheckin.fair_company_id)
        .where(FairCompany.fair_id == fair_id)
    ).scalar_one()
    actions = _action_counts(db, fair_id)
    return {
        "fair_id": fair_id,
        "registered_companies": int(registered),
        "approved_companies": int(approved),
        "checked_in_students": int(named_checkins or 0) + int(anonymous or 0),
        "checked_in_enterprise_representatives": int(enterprise_checkins or 0),
        "applications": actions[JobActionType.APPLY.value],
        "interviews": actions[JobActionType.INTERVIEW.value],
        "offers": actions[JobActionType.OFFER.value],
        "acceptances": actions[JobActionType.ACCEPTED.value],
        "signings": actions[JobActionType.SIGN.value],
    }


def export_fair_materials_docx(db: Session, fair_id: int) -> bytes:
    fair = _get_fair(db, fair_id)
    companies = list_companies(db, fair_id)
    funnel = get_funnel(db, fair_id)
    doc = Document()
    doc.add_heading(f"{fair.name} 场次方案", level=1)
    doc.add_paragraph(
        f"时间：{fair.held_at.astimezone().strftime('%Y年%m月%d日 %H:%M') if fair.held_at else '待定'}"
    )
    doc.add_paragraph(f"地点：{fair.location or '待定'}")
    doc.add_heading("企业清单", level=2)
    for row in companies:
        doc.add_paragraph(
            f"{row.company} / {row.title or '-'} / {row.status} / 展位 {row.booth or '-'}",
            style=None,
        )
    doc.add_heading("岗位与漏斗统计", level=2)
    doc.add_paragraph(f"报名企业 {funnel['registered_companies']} 家")
    doc.add_paragraph(f"审核通过 {funnel['approved_companies']} 家")
    doc.add_paragraph(f"签到 {funnel['checked_in_students']} 人")
    doc.add_paragraph(f"投递 {funnel['applications']} 次")
    doc.add_paragraph(f"面试 {funnel['interviews']} 次")
    doc.add_paragraph(f"录用通知 {funnel['offers']} 次")
    doc.add_paragraph(f"实际录用 {funnel['acceptances']} 人")
    doc.add_paragraph(f"签约 {funnel['signings']} 人")
    doc.add_heading("安保要点", level=2)
    doc.add_paragraph("入口分流、消防通道保持畅通、展位用电提前巡检。")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _get_fair(
    db: Session,
    fair_id: int,
    *,
    lock: bool = False,
    shared_lock: bool = False,
) -> Fair:
    if lock and shared_lock:
        raise ValueError("fair lock mode is ambiguous")
    statement = select(Fair).where(Fair.id == fair_id)
    if lock:
        statement = statement.with_for_update().execution_options(populate_existing=True)
    elif shared_lock:
        statement = statement.with_for_update(read=True).execution_options(
            populate_existing=True
        )
    fair = db.execute(statement).scalar_one_or_none()
    if fair is None:
        raise ValueError("fair not found")
    return fair


def _require_fair_status(fair: Fair, allowed: set[str], *, action: str) -> None:
    if fair.status not in allowed:
        allowed_text = "、".join(sorted(allowed))
        raise FairConflictError(
            f"场次当前为“{fair.status}”，只能在“{allowed_text}”状态下{action}"
        )


def _operation_hash(payload: dict[str, Any]) -> str:
    encoded = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _acquire_operation_lock(db: Session, namespace: str, key: str) -> None:
    db.execute(
        text("SELECT pg_advisory_xact_lock(hashtextextended(:lock_key, 17))"),
        {"lock_key": f"{namespace}:{key}"},
    )


def _action_counts(db: Session, fair_id: int) -> dict[str, int]:
    result = {kind.value: 0 for kind in JobActionType}
    rows = db.execute(
        select(JobAction.type, func.count())
        .where(JobAction.source == f"fair:{fair_id}")
        .group_by(JobAction.type)
    ).all()
    for kind, count in rows:
        result[kind] = int(count)
    return result


def _fair_dict(fair: Fair) -> dict[str, Any]:
    return {
        "id": fair.id,
        "name": fair.name,
        "fair_type": fair.fair_type,
        "held_at": fair.held_at.isoformat() if fair.held_at else None,
        "location": fair.location,
        "status": fair.status,
    }


def _company_dict(row: FairCompany) -> dict[str, Any]:
    return {
        "id": row.id,
        "fair_id": row.fair_id,
        "company": row.company,
        "booth": row.booth,
        "status": row.status,
        "title": row.title,
        "job_id": row.job_id,
    }
